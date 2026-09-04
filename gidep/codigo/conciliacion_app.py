# -*- coding: utf-8 -*-
"""
PANEL CONCILIACIÓN BCE vs MEF
Deuda Externa Pública — cruce mensual por acreedor

Aplicación web local (Flask) que reutiliza la lógica probada de
Conciliacion.py. Permite cargar los dos reportes (.xls), ejecutar el
cruce de los 6 conceptos por acreedor, ver el resultado con semáforo,
exportar el Excel y guardar el historial en SQLite.

Uso:
    pip install flask xlrd openpyxl
    python conciliacion_app.py
    -> abre http://127.0.0.1:5001/
"""
import json
import hashlib
import os
import re
import shutil
import sqlite3
import uuid
import webbrowser
from datetime import datetime
from threading import Timer

from flask import (Flask, request, jsonify, render_template_string,
                   send_file, session, g, has_request_context)
from werkzeug.utils import secure_filename
from caf_condonados_module import register_caf_condonados_routes
from agenda_pagos_module import register_agenda_pagos_routes
from mensajes_firmados_module import register_mensajes_firmados_routes
from comprobantes_contables_module import register_comprobantes_contables_routes
from activaciones_cuentas_module import register_activaciones_cuentas_routes
from respuestas_correos_module import register_respuestas_correos_routes
from archivo_quipux_module import register_archivo_quipux_routes
from contratos_agencia_fiscal_module import register_contratos_agencia_fiscal_routes
from matriz_prestamos_module import register_matriz_prestamos_routes
from corresponsales_module import register_corresponsales_routes
from control_operativo_module import register_control_operativo_routes
from auth_module import register_auth

# Reutilizamos la lógica de conciliación ya validada.
# Debe estar Conciliacion.py en la MISMA carpeta que este archivo.
try:
    import Conciliacion as C
except ModuleNotFoundError:
    import sys
    print("=" * 60)
    print("  ERROR: falta 'Conciliacion.py' en esta carpeta.")
    print("  Copia conciliacion_app.py y Conciliacion.py JUNTOS,")
    print("  en la misma carpeta, y vuelve a ejecutar.")
    print("=" * 60)
    sys.exit(1)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "conciliacion.db")
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads_conciliacion")
EXPORT_DIR = os.path.join(BASE_DIR, "reportes_conciliacion")
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)

# =============================================================================
# BASE DE DATOS
# =============================================================================

def init_db():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS conciliaciones (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        periodo TEXT,
        acreedor TEXT,
        concepto TEXT,
        mef REAL,
        bce REAL,
        diferencia REAL,
        estado TEXT,
        fecha_corrida TEXT
    )''')
    # Migración: columna 'nota' (para pagos directos agregados manualmente)
    try:
        c.execute("ALTER TABLE conciliaciones ADD COLUMN nota TEXT DEFAULT ''")
    except Exception:
        pass
    c.execute('''CREATE TABLE IF NOT EXISTS conciliaciones_archivo (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        periodo TEXT NOT NULL,
        acreedor TEXT NOT NULL,
        concepto TEXT NOT NULL,
        mef REAL,
        bce REAL,
        diferencia REAL,
        estado TEXT,
        fecha_corrida TEXT,
        nota TEXT DEFAULT '',
        archivado_en TEXT NOT NULL,
        archivado_por TEXT DEFAULT '',
        motivo TEXT DEFAULT '',
        archivo_bce TEXT DEFAULT '',
        archivo_mef TEXT DEFAULT '',
        hash_bce TEXT DEFAULT '',
        hash_mef TEXT DEFAULT ''
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS periodo_archivos (
        periodo TEXT PRIMARY KEY,
        archivo_bce TEXT DEFAULT '',
        archivo_mef TEXT DEFAULT '',
        ajustes_json TEXT DEFAULT '[]',
        auto_ajustes_json TEXT DEFAULT '[]',
        actualizado_en TEXT NOT NULL,
        actualizado_por TEXT DEFAULT ''
    )''')
    c.execute("CREATE INDEX IF NOT EXISTS idx_conc_archivo_periodo ON conciliaciones_archivo(periodo, fecha_corrida)")
    conn.commit()
    conn.close()


def db_q(sql, params=(), fetch=False):
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute(sql, params)
    rows = [dict(r) for r in c.fetchall()] if fetch else None
    conn.commit()
    conn.close()
    return rows


# =============================================================================
# FLASK
# =============================================================================
app = Flask(__name__)
init_db()
register_auth(app, BASE_DIR)

# Archivos cargados por periodo (para regenerar el BCE ajustado con pagos directos)
PERIODO_FILES = {}


def _file_hash(path):
    if not path or not os.path.isfile(path):
        return ""
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _safe_json_list(raw):
    try:
        value = json.loads(raw or "[]")
        return value if isinstance(value, list) else []
    except (TypeError, ValueError):
        return []


def _load_period_files():
    rows = db_q("SELECT * FROM periodo_archivos", fetch=True) or []
    for row in rows:
        PERIODO_FILES[row["periodo"]] = {
            "bce": row.get("archivo_bce") or None,
            "mef": row.get("archivo_mef") or None,
            "ajustes": _safe_json_list(row.get("ajustes_json")),
            "auto_ajustes": _safe_json_list(row.get("auto_ajustes_json")),
        }


def _persist_period_info(periodo):
    info = PERIODO_FILES.get(periodo) or {}
    username = session.get("username", "") if has_request_context() else "sistema"
    db_q(
        """INSERT INTO periodo_archivos
           (periodo, archivo_bce, archivo_mef, ajustes_json, auto_ajustes_json,
            actualizado_en, actualizado_por)
           VALUES (?, ?, ?, ?, ?, ?, ?)
           ON CONFLICT(periodo) DO UPDATE SET
             archivo_bce=excluded.archivo_bce,
             archivo_mef=excluded.archivo_mef,
             ajustes_json=excluded.ajustes_json,
             auto_ajustes_json=excluded.auto_ajustes_json,
             actualizado_en=excluded.actualizado_en,
             actualizado_por=excluded.actualizado_por""",
        (
            periodo,
            info.get("bce") or "",
            info.get("mef") or "",
            json.dumps(info.get("ajustes") or [], ensure_ascii=False),
            json.dumps(info.get("auto_ajustes") or [], ensure_ascii=False),
            datetime.now().astimezone().isoformat(timespec="seconds"),
            username,
        ),
    )


def _remember_period(periodo, bce=None, mef=None, reset=False):
    if reset or periodo not in PERIODO_FILES:
        PERIODO_FILES[periodo] = {
            "bce": bce,
            "mef": mef,
            "ajustes": [],
            "auto_ajustes": [],
        }
    else:
        info = PERIODO_FILES[periodo]
        if bce:
            info["bce"] = bce
        if mef:
            info["mef"] = mef
    _persist_period_info(periodo)
    return PERIODO_FILES[periodo]


def _archive_current_reconciliation(periodo, reason="Nueva ejecución"):
    rows = db_q("SELECT * FROM conciliaciones WHERE periodo=? ORDER BY id", (periodo,), fetch=True) or []
    if not rows:
        return 0
    info = PERIODO_FILES.get(periodo) or {}
    archived_at = datetime.now().astimezone().isoformat(timespec="seconds")
    username = session.get("username", "") if has_request_context() else "sistema"
    bce_path = info.get("bce") or ""
    mef_path = info.get("mef") or ""
    bce_hash = _file_hash(bce_path)
    mef_hash = _file_hash(mef_path)
    conn = sqlite3.connect(DB_PATH)
    try:
        conn.executemany(
            """INSERT INTO conciliaciones_archivo
               (periodo, acreedor, concepto, mef, bce, diferencia, estado,
                fecha_corrida, nota, archivado_en, archivado_por, motivo,
                archivo_bce, archivo_mef, hash_bce, hash_mef)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (
                    row["periodo"], row["acreedor"], row["concepto"], row["mef"],
                    row["bce"], row["diferencia"], row["estado"], row["fecha_corrida"],
                    row.get("nota") or "", archived_at, username, reason,
                    bce_path, mef_path, bce_hash, mef_hash,
                )
                for row in rows
            ],
        )
        conn.execute("DELETE FROM conciliaciones WHERE periodo=?", (periodo,))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return len(rows)


def _replace_current_reconciliation(periodo, rows, run_at, reason="Nueva ejecución"):
    """Archiva la versión activa e inserta la nueva en una sola transacción."""
    info = PERIODO_FILES.get(periodo) or {}
    archived_at = datetime.now().astimezone().isoformat(timespec="seconds")
    username = session.get("username", "") if has_request_context() else "sistema"
    bce_path = info.get("bce") or ""
    mef_path = info.get("mef") or ""
    bce_hash = _file_hash(bce_path)
    mef_hash = _file_hash(mef_path)
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("BEGIN IMMEDIATE")
        previous = conn.execute(
            "SELECT * FROM conciliaciones WHERE periodo=? ORDER BY id", (periodo,)
        ).fetchall()
        if previous:
            conn.executemany(
                """INSERT INTO conciliaciones_archivo
                   (periodo, acreedor, concepto, mef, bce, diferencia, estado,
                    fecha_corrida, nota, archivado_en, archivado_por, motivo,
                    archivo_bce, archivo_mef, hash_bce, hash_mef)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                [
                    (
                        row["periodo"], row["acreedor"], row["concepto"], row["mef"],
                        row["bce"], row["diferencia"], row["estado"], row["fecha_corrida"],
                        row["nota"] or "", archived_at, username, reason,
                        bce_path, mef_path, bce_hash, mef_hash,
                    )
                    for row in previous
                ],
            )
        conn.execute("DELETE FROM conciliaciones WHERE periodo=?", (periodo,))
        conn.executemany(
            """INSERT INTO conciliaciones
               (periodo, acreedor, concepto, mef, bce, diferencia, estado, fecha_corrida, nota)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (periodo, ac, concepto, vm, vb, dif, estado, run_at, "")
                for ac, concepto, vm, vb, dif, estado in rows
            ],
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


_load_period_files()
register_caf_condonados_routes(app, BASE_DIR)
register_agenda_pagos_routes(app, BASE_DIR)
register_mensajes_firmados_routes(app, BASE_DIR)
register_comprobantes_contables_routes(app, BASE_DIR)
register_activaciones_cuentas_routes(app, BASE_DIR)
register_respuestas_correos_routes(app, BASE_DIR)
register_archivo_quipux_routes(app, BASE_DIR)
register_contratos_agencia_fiscal_routes(app, BASE_DIR)
register_corresponsales_routes(app, BASE_DIR)
register_matriz_prestamos_routes(app, BASE_DIR)
register_control_operativo_routes(app, BASE_DIR)


@app.errorhandler(Exception)
def handle_exception(e):
    from werkzeug.exceptions import HTTPException
    if isinstance(e, HTTPException):
        if request.path.startswith("/api/"):
            return jsonify({"ok": False, "error": e.description}), e.code
        return e
    request_id = getattr(g, "request_id", uuid.uuid4().hex[:12])
    app.logger.exception("Error no manejado. request_id=%s", request_id)
    if request.path.startswith("/api/"):
        return jsonify({
            "ok": False,
            "error": "Ocurrió un error interno. Consulte el identificador con el administrador.",
            "request_id": request_id,
        }), 500
    return (
        "<main style='max-width:720px;margin:60px auto;font:16px Segoe UI;color:#173a5e'>"
        "<h1>No se pudo completar la operación</h1>"
        f"<p>Identificador del incidente: <b>{request_id}</b></p>"
        "<p>Vuelva al gestor e intente nuevamente.</p></main>",
        500,
    )


def _guardar_subida(storage):
    """Guarda un FileStorage y devuelve su ruta, o None si no vino."""
    if not storage or not storage.filename:
        return None
    nombre = secure_filename(storage.filename) or "archivo.xls"
    lote = datetime.now().strftime("%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:8]
    carpeta = os.path.join(UPLOAD_DIR, lote)
    os.makedirs(carpeta, exist_ok=False)
    ruta = os.path.join(carpeta, nombre)
    storage.save(ruta)
    return ruta


def _origen_reporte(ruta):
    """Detecta si un .xls es del BCE o del MEF mirando sus hojas.
    BCE -> tiene hojas 'Giros del/al Exterior'; MEF -> tiene hoja 'Resumen'."""
    try:
        import xlrd
        nombres = [h.name.lower() for h in xlrd.open_workbook(ruta).sheets()]
    except Exception:
        return None
    if any("giros" in n for n in nombres):
        return "bce"
    if any("resumen" in n for n in nombres):
        return "mef"
    return None


def _clasificar_subidas():
    """Toma TODOS los archivos subidos (en cualquier campo) y los reparte en
    BCE/MEF por su contenido. Así el usuario puede arrastrarlos en cualquier
    orden o los dos juntos. Devuelve (ruta_bce, ruta_mef, advertencias)."""
    storages = []
    for campo in request.files:
        storages.extend(request.files.getlist(campo))
    ruta_bce = ruta_mef = None
    avisos = []
    for st in storages:
        ruta = _guardar_subida(st)
        if not ruta:
            continue
        origen = _origen_reporte(ruta)
        if origen == "bce" and not ruta_bce:
            ruta_bce = ruta
        elif origen == "mef" and not ruta_mef:
            ruta_mef = ruta
        elif origen is None:
            avisos.append(f"No reconocí '{os.path.basename(ruta)}' (sin hojas Giros/Resumen)")
        else:
            avisos.append(f"Dos reportes del mismo origen ({origen.upper()}): ignoré '{os.path.basename(ruta)}'")
    return ruta_bce, ruta_mef, avisos


@app.route("/api/conciliar", methods=["POST"])
def api_conciliar():
    """Recibe los dos .xls (multipart) o rutas (json), ejecuta el cruce,
    guarda el resultado en BD y lo devuelve."""
    avisos = []
    if request.files:
        # Detección automática por contenido: no importa en qué zona se soltó
        # cada archivo, ni el orden, ni si se arrastraron los dos juntos.
        ruta_bce, ruta_mef, avisos = _clasificar_subidas()
        periodo = request.form.get("periodo", "").strip()
    else:
        data = request.json or {}
        ruta_bce = data.get("ruta_bce")
        ruta_mef = data.get("ruta_mef")
        periodo = (data.get("periodo") or "").strip()

    if not ruta_bce or not os.path.exists(ruta_bce):
        return jsonify({"ok": False, "error": "No identifiqué el reporte del BCE "
                        "(debe tener hojas 'Giros del/al Exterior'). " + " ".join(avisos)})
    if not ruta_mef or not os.path.exists(ruta_mef):
        return jsonify({"ok": False, "error": "No identifiqué el reporte del MEF "
                        "(debe tener hoja 'Resumen'). " + " ".join(avisos)})

    filas = C.conciliar_archivos(ruta_bce, ruta_mef)
    # Periodo detectado automáticamente del archivo (el usuario no lo elige)
    info_periodo = C.detectar_periodo(ruta_mef, ruta_bce)
    periodo = info_periodo["periodo"]
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Recordar los archivos de este periodo para los ajustes posteriores
    _replace_current_reconciliation(periodo, filas, ahora)
    _remember_period(periodo, ruta_bce, ruta_mef, reset=True)

    res = _resultado_periodo(periodo)
    res.update({
        "fecha": ahora, "info_periodo": info_periodo,
        "archivo_bce": os.path.basename(ruta_bce),
        "archivo_mef": os.path.basename(ruta_mef), "avisos": avisos,
    })
    return jsonify(res)


@app.route("/api/exportar", methods=["POST"])
def api_exportar():
    """Genera el Excel con semáforo del periodo indicado y lo descarga."""
    data = request.json or {}
    periodo = (data.get("periodo") or "").strip()
    if not periodo:
        return jsonify({"ok": False, "error": "Falta el periodo"})
    rows = db_q("""SELECT acreedor, concepto, mef, bce, diferencia, estado
                   FROM conciliaciones WHERE periodo = ?
                   ORDER BY id""", (periodo,), fetch=True)
    if not rows:
        return jsonify({"ok": False, "error": "No hay datos para ese periodo"})
    filas = [(r["acreedor"], r["concepto"], r["mef"], r["bce"],
              r["diferencia"], r["estado"]) for r in rows]
    nombre = f"Conciliacion_BCE_MEF_{periodo.replace('/', '-')}.xlsx"
    ruta = os.path.join(EXPORT_DIR, nombre)
    if not C.reporte_xlsx(filas, ruta):
        return jsonify({"ok": False, "error": "openpyxl no disponible"})
    return send_file(ruta, as_attachment=True, download_name=nombre)


@app.route("/api/historial")
def api_historial():
    """Lista los periodos conciliados y su resumen."""
    rows = db_q("""SELECT periodo,
                          COUNT(*) AS total,
                          SUM(CASE WHEN estado='CONCILIADO' THEN 1 ELSE 0 END) AS conciliados,
                          SUM(CASE WHEN estado='DIFERENCIA' THEN 1 ELSE 0 END) AS diferencias,
                          MAX(fecha_corrida) AS ultima
                   FROM conciliaciones
                   GROUP BY periodo ORDER BY periodo DESC""", fetch=True)
    for row in rows or []:
        versions = db_q(
            "SELECT COUNT(DISTINCT fecha_corrida) AS total FROM conciliaciones_archivo WHERE periodo=?",
            (row["periodo"],), fetch=True,
        )
        row["versiones_anteriores"] = (versions or [{"total": 0}])[0]["total"] or 0
    return jsonify({"ok": True, "periodos": rows or []})


@app.get("/api/historial/versiones")
def api_historial_versiones():
    periodo = (request.args.get("periodo") or "").strip()
    if not periodo:
        return jsonify({"ok": False, "error": "Falta el periodo."}), 400
    rows = db_q(
        """SELECT fecha_corrida, archivado_en, archivado_por, motivo,
                  archivo_bce, archivo_mef, hash_bce, hash_mef,
                  COUNT(*) AS registros,
                  SUM(CASE WHEN estado='DIFERENCIA' THEN 1 ELSE 0 END) AS diferencias
           FROM conciliaciones_archivo WHERE periodo=?
           GROUP BY fecha_corrida, archivado_en, archivado_por, motivo,
                    archivo_bce, archivo_mef, hash_bce, hash_mef
           ORDER BY archivado_en DESC""",
        (periodo,), fetch=True,
    ) or []
    return jsonify({"ok": True, "periodo": periodo, "versiones": rows})


@app.route("/api/periodo/<periodo>")
def api_periodo(periodo):
    rows = db_q("""SELECT acreedor, concepto, mef, bce, diferencia, estado, nota
                   FROM conciliaciones WHERE periodo = ? ORDER BY id""",
                (periodo,), fetch=True)
    return jsonify({"ok": True, "periodo": periodo, "registros": rows or []})


# -------------------------------------------------------------------------
# Resultado consolidado de un periodo desde la BD (reutilizado por ajustes)
# -------------------------------------------------------------------------
def _resultado_periodo(periodo):
    rows = db_q("""SELECT acreedor, concepto, mef, bce, diferencia, estado, nota
                   FROM conciliaciones WHERE periodo = ? ORDER BY id""",
                (periodo,), fetch=True) or []
    filas = [(r["acreedor"], r["concepto"], r["mef"], r["bce"],
              r["diferencia"], r["estado"]) for r in rows]
    conciliados = sum(1 for r in rows if r["estado"] == "CONCILIADO")
    diferencias = len(rows) - conciliados
    total_dif = round(sum(abs(r["diferencia"]) for r in rows if r["estado"] == "DIFERENCIA"), 2)
    totales = [{"acreedor": ac, **t} for ac, t in C.totales_por_cartera(filas).items()]
    # Gran total (todas las carteras)
    g_mef = round(sum(t["mef"] for t in totales), 2)
    g_bce = round(sum(t["bce"] for t in totales), 2)
    g_dif = round(g_mef - g_bce, 2) or 0.0
    # ¿Todo concilia? (gobierna habilitación de Quipux)
    conciliado_total = diferencias == 0 and len(rows) > 0
    # Diagnóstico SOLO de lo que no concilia
    diag = []
    prestamos = {}
    info = PERIODO_FILES.get(periodo) or {}
    mef_path = info.get("mef")
    bce_path = info.get("bce")
    pares = [(r["acreedor"], r["concepto"]) for r in rows if r["estado"] == "DIFERENCIA"]
    if mef_path and os.path.exists(mef_path):
        try:
            diag = C.diagnostico(mef_path, filas)
        except Exception:
            diag = []
        # Diagnóstico por préstamo: qué crédito exacto no cuadra
        if pares and bce_path and os.path.exists(bce_path):
            try:
                prestamos = C.diagnostico_prestamos(mef_path, bce_path, pares,
                                                    info.get("ajustes"))
            except Exception:
                prestamos = {}
    bce_aj = bool(info.get("ajustes")) or any(
        "Pago directo agregado al BCE" in (r.get("nota") or "") for r in rows
    )
    return {
        "ok": True, "periodo": periodo, "registros": rows,
        "total": len(rows), "conciliados": conciliados, "diferencias": diferencias,
        "total_diferencia": total_dif, "totales": totales,
        "gran_total": {"mef": g_mef, "bce": g_bce, "dif": g_dif},
        "conciliado_total": conciliado_total, "diagnostico": diag,
        "prestamos": prestamos, "bce_ajustado": bce_aj,
    }


@app.route("/api/pago_directo", methods=["POST"])
def api_pago_directo():
    """Previsualiza los respaldos de pago directo del MEF (sin aplicar)."""
    previos = []
    storages = []
    for campo in request.files:
        storages.extend(request.files.getlist(campo))
    for st in storages:
        ruta = _guardar_subida(st)
        if not ruta:
            continue
        try:
            pd = C.leer_pago_directo(ruta)
            pd["archivo"] = os.path.basename(ruta)
            previos.append(pd)
        except Exception as e:
            previos.append({"archivo": st.filename, "error": str(e)})
    return jsonify({"ok": True, "previos": previos})


@app.route("/api/aplicar_pago_directo", methods=["POST"])
def api_aplicar_pago_directo():
    """Agrega un pago directo al lado del BCE (Desembolsos) y recalcula."""
    data = request.json or {}
    periodo = (data.get("periodo") or "").strip()
    acreedor = (data.get("acreedor") or "").strip().upper()
    concepto = (data.get("concepto") or "Desembolsos").strip()
    valor = float(data.get("valor") or 0)
    nota = (data.get("nota") or "").strip()
    if not periodo or not acreedor or valor == 0:
        return jsonify({"ok": False, "error": "Faltan periodo, acreedor o valor"})

    fila = db_q("""SELECT * FROM conciliaciones
                   WHERE periodo=? AND acreedor=? AND concepto=?""",
                (periodo, acreedor, concepto), fetch=True)
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    if fila:
        r = fila[0]
        nuevo_bce = round((r["bce"] or 0) + valor, 2)
        dif = round((r["mef"] or 0) - nuevo_bce, 2) or 0.0
        estado = "CONCILIADO" if abs(dif) <= C.TOLERANCIA else "DIFERENCIA"
        nota_final = (r["nota"] + " | " + nota).strip(" |") if r["nota"] else nota
        db_q("""UPDATE conciliaciones SET bce=?, diferencia=?, estado=?, nota=?
                WHERE id=?""", (nuevo_bce, dif, estado, nota_final, r["id"]))
    else:
        # No existía fila de Desembolsos para esa cartera: se crea
        dif = round(0 - valor, 2)
        estado = "CONCILIADO" if abs(dif) <= C.TOLERANCIA else "DIFERENCIA"
        db_q("""INSERT INTO conciliaciones
                (periodo, acreedor, concepto, mef, bce, diferencia, estado, fecha_corrida, nota)
                VALUES (?,?,?,?,?,?,?,?,?)""",
             (periodo, acreedor, concepto, 0.0, valor, dif, estado, ahora, nota))

    # Agregar la fila al MISMO archivo .xls del BCE que el usuario subió
    info = PERIODO_FILES.setdefault(periodo, {"bce": None, "mef": None, "ajustes": [], "auto_ajustes": []})
    ajuste = {"acreedor": acreedor, "concepto": concepto,
              "referencia": data.get("referencia", ""), "valor": valor, "nota": nota,
              "prestamista": data.get("prestamista", acreedor)}
    info["ajustes"].append(ajuste)
    _persist_period_info(periodo)
    modificado = False
    motivo = ""
    if info.get("bce") and os.path.exists(info["bce"]):
        try:
            modificado = C.modificar_bce_xls(info["bce"], [ajuste])
            if not modificado:
                motivo = ("No se pudo modificar el .xls (falta xlutils/xlwt). "
                          "Ejecute: pip install xlutils xlwt")
        except Exception as e:
            motivo = f"No se pudo modificar el archivo: {e}"
    else:
        motivo = "No tengo el archivo del BCE de este periodo en memoria; vuelva a conciliar."

    res = _resultado_periodo(periodo)
    res["archivo_modificado"] = modificado
    res["archivo_motivo"] = motivo
    return jsonify(res)


@app.route("/api/descargar_bce_ajustado")
def api_descargar_bce_ajustado():
    """Descarga el MISMO archivo del BCE ya modificado con los pagos directos."""
    periodo = (request.args.get("periodo") or "").strip()
    if "_asegurar_archivos_periodo" in globals():
        info = _asegurar_archivos_periodo(periodo)
    else:
        info = PERIODO_FILES.get(periodo) or {}
    ruta = info.get("bce")
    if not ruta or not os.path.exists(ruta):
        return jsonify({"ok": False, "error": "No hay archivo BCE para este periodo"})
    return send_file(ruta, as_attachment=True, download_name=os.path.basename(ruta))


def _ultimos_reportes_subidos():
    """Devuelve el ultimo par BCE/MEF cargado en uploads_conciliacion."""
    encontrados = {"bce": [], "mef": []}
    for root, _dirs, files in os.walk(UPLOAD_DIR):
        for nombre in files:
            ruta = os.path.join(root, nombre)
            if not nombre.lower().endswith((".xls", ".xlsx")):
                continue
            origen = _origen_reporte(ruta)
            if origen in encontrados:
                encontrados[origen].append((os.path.getmtime(ruta), ruta))
    if not encontrados["bce"] or not encontrados["mef"]:
        return None, None
    return max(encontrados["bce"])[1], max(encontrados["mef"])[1]


def _guardar_conciliacion_desde_rutas(ruta_bce, ruta_mef, avisos=None):
    avisos = avisos or []
    filas = C.conciliar_archivos(ruta_bce, ruta_mef)
    info_periodo = C.detectar_periodo(ruta_mef, ruta_bce)
    periodo = info_periodo["periodo"]
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    _replace_current_reconciliation(periodo, filas, ahora, "Nueva ejecución automática")
    _remember_period(periodo, ruta_bce, ruta_mef, reset=True)

    res = _resultado_periodo(periodo)
    res.update({
        "fecha": ahora,
        "info_periodo": info_periodo,
        "archivo_bce": os.path.basename(ruta_bce),
        "archivo_mef": os.path.basename(ruta_mef),
        "avisos": avisos,
    })
    return res


def _cerrar_fila_automaticamente(row, nota):
    """Ajusta solo la matriz de conciliacion: BCE pasa a MEF y la nota queda auditada."""
    ajuste = round((row["mef"] or 0) - (row["bce"] or 0), 2)
    nota_actual = row.get("nota") or ""
    nota_final = (nota_actual + " | " + nota).strip(" |")
    db_q("""UPDATE conciliaciones
            SET bce=?, diferencia=?, estado=?, nota=?
            WHERE id=?""",
         (round(row["mef"] or 0, 2), 0.0, "CONCILIADO", nota_final, row["id"]))
    return ajuste



# -------------------------------------------------------------------------
# Gestion de observaciones con anexos del mes
# -------------------------------------------------------------------------
MESES_OBS = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]
MESES_OBS_ABR = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
EXT_ANEXOS = (".xls", ".xlsx", ".xlsm", ".pdf")
LENDER_OBS = {
    "IBRD": "BIRF", "BIRF": "BIRF", "WORLD BANK": "BIRF",
    "KFW": "GOBIERNOS", "EXIMBANK KOREA": "GOBIERNOS", "EXIMBANK": "GOBIERNOS",
    "CAF": "CAF", "CFA": "CAF", "BID": "BID", "IDB": "BID",
    "FIDA": "FIDA", "IFAD": "FIDA",
}


def _normal_obs(texto):
    return str(texto or "").upper().replace("\xa0", " ")


def _periodo_mes_obs(periodo):
    try:
        anio, mes = str(periodo).split("-")[:2]
        idx = int(mes) - 1
        if idx < 0 or idx > 11:
            raise ValueError
        return int(anio), MESES_OBS[idx], MESES_OBS_ABR[idx]
    except Exception:
        ahora = datetime.now()
        return ahora.year, MESES_OBS[ahora.month - 1], MESES_OBS_ABR[ahora.month - 1]


def _leer_config_gestor_obs():
    ruta = os.path.join(BASE_DIR, "gestor_pagos_config.json")
    base = {
        "ruta_raiz_destino": r"Z:\GISI\SSFI\GESTIÓN PAGOS INTERNACIONALES",
        "carpeta_puente": os.path.join(os.path.expanduser("~"), "Desktop", "Puente Pagos Quipux"),
    }
    try:
        with open(ruta, "r", encoding="utf-8") as fh:
            data = json.load(fh) or {}
        base.update({k: v for k, v in data.items() if v})
    except Exception:
        pass
    return base


def _buscar_reportes_periodo_uploads(periodo):
    anio, mes, abr = _periodo_mes_obs(periodo)
    tokens = [
        f"{abr}{anio}".lower(), f"{mes}{anio}".lower(),
        f"{periodo}".lower(), f"{mes.lower()} {anio}", f"{abr.lower()} {anio}",
    ]
    encontrados = {"bce": [], "mef": []}
    for root, _dirs, files in os.walk(UPLOAD_DIR):
        for nombre in files:
            ruta = os.path.join(root, nombre)
            if not nombre.lower().endswith((".xls", ".xlsx")):
                continue
            nlow = nombre.lower().replace("_", "").replace("-", "")
            if not any(tok.replace("-", "").replace(" ", "") in nlow for tok in tokens):
                continue
            origen = _origen_reporte(ruta)
            if origen in encontrados:
                encontrados[origen].append((os.path.getmtime(ruta), ruta))
    bce = max(encontrados["bce"])[1] if encontrados["bce"] else None
    mef = max(encontrados["mef"])[1] if encontrados["mef"] else None
    return bce, mef


def _asegurar_archivos_periodo(periodo):
    info = PERIODO_FILES.setdefault(periodo, {"bce": None, "mef": None, "ajustes": [], "auto_ajustes": []})
    if info.get("bce") and info.get("mef") and os.path.exists(info["bce"]) and os.path.exists(info["mef"]):
        return info
    bce, mef = _buscar_reportes_periodo_uploads(periodo)
    if not bce or not mef:
        ult_bce, ult_mef = _ultimos_reportes_subidos()
        bce = bce or ult_bce
        mef = mef or ult_mef
    if bce:
        info["bce"] = bce
    if mef:
        info["mef"] = mef
    info.setdefault("ajustes", [])
    info.setdefault("auto_ajustes", [])
    _persist_period_info(periodo)
    return info


def _carpetas_anexos_periodo(periodo, acreedor=None):
    anio, mes, _abr = _periodo_mes_obs(periodo)
    config = _leer_config_gestor_obs()
    raiz = config.get("ruta_raiz_destino") or ""
    deuda = os.path.join(raiz, str(anio), "DEUDA EXTERNA PÚBLICA") if raiz else ""
    carpetas = []

    # Carpeta que se ha usado en conciliaciones anteriores: primero por ser la mas confiable.
    if deuda:
        carpetas.append(os.path.join(deuda, "Conciliaciones", mes, "Anexos"))
        base_acreedores = os.path.join(deuda, "Acreedores Internacionales")
        carteras = [acreedor] if acreedor else []
        if acreedor == "GOBIERNOS":
            carteras.append("GOBIERNOS")
        for cart in [c for c in dict.fromkeys(carteras) if c]:
            for tipo in ("Desembolsos", "Pagos"):
                carpetas.append(os.path.join(base_acreedores, cart, tipo, mes))

    puente = config.get("carpeta_puente")
    if puente:
        carpetas.append(puente)
    carpetas.append(UPLOAD_DIR)
    return [c for c in dict.fromkeys(carpetas) if c and os.path.exists(c)]


def _parse_monto_obs(valor):
    if isinstance(valor, (int, float)):
        return round(float(valor), 2)
    texto = str(valor or "").strip().replace("\xa0", " ")
    if not texto:
        return None
    neg = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()")
    texto = re.sub(r"[^0-9,\.\-]", "", texto)
    if not texto:
        return None
    if texto.startswith("-"):
        neg = True
        texto = texto[1:]
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        partes = texto.split(",")
        texto = texto.replace(".", "")
        if len(partes[-1]) == 2:
            texto = texto.replace(",", ".")
        else:
            texto = texto.replace(",", "")
    else:
        partes = texto.split(".")
        if len(partes) > 2:
            texto = "".join(partes[:-1]) + "." + partes[-1]
    try:
        valor = float(texto)
    except ValueError:
        return None
    return round(-valor if neg else valor, 2)


def _montos_en_texto_obs(texto):
    patron = r"\(?\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{2})\)?|\(?\d+(?:[.,]\d{2})\)?"
    montos = []
    for raw in re.findall(patron, texto or ""):
        val = _parse_monto_obs(raw)
        if val is not None and abs(val) >= 1000:
            montos.append(abs(round(val, 2)))
    return sorted(set(montos))


def _refs_obs(texto):
    refs = set()
    for raw in re.findall(r"\d{3,8}", str(texto or "")):
        refs.add(raw)
        if len(raw) > 4 and raw.endswith("0"):
            refs.add(raw.rstrip("0") or raw)
            refs.add(raw[:-1])
    return refs


def _acreedor_desde_texto_obs(texto):
    t = _normal_obs(texto)
    for clave, dest in LENDER_OBS.items():
        if clave in t:
            return dest
    return ""


def _prestamista_desde_texto_obs(texto, acreedor):
    t = _normal_obs(texto)
    for clave in ("KFW", "EXIMBANK KOREA", "EXIMBANK", "IBRD", "BIRF", "WORLD BANK", "CAF", "BID"):
        if clave in t:
            return "BIRF" if clave in ("IBRD", "WORLD BANK") else clave
    return acreedor


def _leer_excel_anexo_obs(ruta):
    ext = os.path.splitext(ruta)[1].lower()
    filas = []
    if ext == ".xls":
        import xlrd
        libro = xlrd.open_workbook(ruta)
        hoja = libro.sheet_by_index(0)
        for r in range(hoja.nrows):
            filas.append([hoja.cell_value(r, c) for c in range(hoja.ncols)])
    else:
        from openpyxl import load_workbook
        wb = load_workbook(ruta, data_only=True, read_only=True)
        ws = wb.active
        filas = [list(row) for row in ws.iter_rows(values_only=True)]

    texto = os.path.basename(ruta) + " " + " ".join(
        str(c) for fila in filas[:80] for c in fila if c not in ("", None)
    )
    montos = []
    etiquetas = ("MONTO", "TOTAL", "AMOUNT", "PYMT", "PAID", "PAGO", "VALOR", "DIRECT PAYMENT", "DESEMBOLSO")
    for i, fila in enumerate(filas):
        fila_txt = _normal_obs(" ".join(str(c) for c in fila if c not in ("", None)))
        prev_txt = _normal_obs(" ".join(str(c) for c in filas[i - 1] if c not in ("", None))) if i else ""
        if not any(e in fila_txt or e in prev_txt for e in etiquetas):
            continue
        for celda in fila:
            val = _parse_monto_obs(celda)
            if val is not None and 1000 <= abs(val) < 100000000:
                montos.append(abs(round(val, 2)))
    return {
        "archivo": ruta,
        "nombre": os.path.basename(ruta),
        "acreedor": _acreedor_desde_texto_obs(texto),
        "prestamista": _prestamista_desde_texto_obs(texto, ""),
        "refs": _refs_obs(texto),
        "montos": sorted(set(montos)),
        "texto": texto[:1200],
        "tipo": "excel",
    }


def _leer_pdf_anexo_obs(ruta):
    texto = ""
    try:
        import PyPDF2
        with open(ruta, "rb") as fh:
            lector = PyPDF2.PdfReader(fh)
            for pagina in lector.pages[:3]:
                texto += "\n" + (pagina.extract_text() or "")
    except Exception:
        texto = ""

    nombre = os.path.basename(ruta)
    # OCR de respaldo para anexos escaneados. Se intenta solo en archivos que parecen relevantes.
    if not texto.strip() and ("PAGO" in _normal_obs(nombre) or "DIRECT" in _normal_obs(nombre)):
        try:
            from caf_condonados_module import extraer_texto_ocr_pdf
            texto = extraer_texto_ocr_pdf(ruta) or ""
        except Exception:
            texto = ""

    total_texto = nombre + " " + texto
    return {
        "archivo": ruta,
        "nombre": nombre,
        "acreedor": _acreedor_desde_texto_obs(total_texto),
        "prestamista": _prestamista_desde_texto_obs(total_texto, ""),
        "refs": _refs_obs(total_texto),
        "montos": _montos_en_texto_obs(texto),
        "texto": total_texto[:1200],
        "tipo": "pdf",
    }


def _iterar_anexos_obs(periodo, acreedor):
    vistos = set()
    for carpeta in _carpetas_anexos_periodo(periodo, acreedor):
        for root, _dirs, files in os.walk(carpeta):
            for nombre in files:
                if not nombre.lower().endswith(EXT_ANEXOS):
                    continue
                nlow = nombre.lower()
                if nlow.startswith("reporte concili") or nlow.startswith("reporte sigade"):
                    continue
                ruta = os.path.join(root, nombre)
                if ruta in vistos:
                    continue
                vistos.add(ruta)
                yield ruta


def _candidato_relevante_obs(cand, acreedor, refs_objetivo):
    texto = _normal_obs(cand.get("nombre", "") + " " + cand.get("texto", "") + " " + cand.get("archivo", ""))
    acreedor_cand = cand.get("acreedor") or _acreedor_desde_texto_obs(texto)
    if acreedor_cand and acreedor_cand != acreedor:
        return False
    if refs_objetivo and cand.get("refs") and cand["refs"].isdisjoint(refs_objetivo):
        # Si el archivo se llama claramente pago directo de esa cartera, no se descarta:
        # algunos anexos escaneados no dejan leer el numero de credito.
        if "PAGO DIRECTO" not in texto and "DIRECT PAYMENT" not in texto:
            return False
    if acreedor == "GOBIERNOS" and any(k in texto for k in ("KFW", "EXIMBANK", "GOBIERNOS", "PAGO DIRECTO")):
        return True
    if acreedor in texto or "PAGO DIRECTO" in texto or "DIRECT PAYMENT" in texto:
        return True
    return bool(refs_objetivo and cand.get("refs") and not cand["refs"].isdisjoint(refs_objetivo))


def _prestamo_objetivo_obs(row, prestamos):
    clave = f"{row['acreedor']}|{row['concepto']}"
    diff = abs(round(row.get("diferencia") or 0, 2))
    for item in prestamos.get(clave, []):
        if item.get("dif", 0) > 0 and abs(abs(item.get("dif", 0)) - diff) <= C.TOLERANCIA:
            return item
    positivos = [i for i in prestamos.get(clave, []) if i.get("dif", 0) > 0]
    return positivos[0] if len(positivos) == 1 else {}


def _combinar_montos_obs(piezas, objetivo):
    limite = min(len(piezas), 8)
    piezas = piezas[:limite]
    for mascara in range(1, 1 << len(piezas)):
        combo = [piezas[i] for i in range(len(piezas)) if mascara & (1 << i)]
        total = round(sum(p["valor"] for p in combo), 2)
        if abs(total - objetivo) <= C.TOLERANCIA:
            return combo
    return []


def _candidatos_anexo_obs(periodo, row, refs_objetivo):
    candidatos = []
    for ruta in _iterar_anexos_obs(periodo, row["acreedor"]):
        try:
            if ruta.lower().endswith(".pdf"):
                cand = _leer_pdf_anexo_obs(ruta)
            else:
                cand = _leer_excel_anexo_obs(ruta)
        except Exception:
            continue
        if _candidato_relevante_obs(cand, row["acreedor"], refs_objetivo):
            candidatos.append(cand)
    return candidatos


def _respaldar_bce_obs(info):
    ruta = info.get("bce")
    if not ruta or not os.path.exists(ruta) or info.get("bce_backup_anexos"):
        return
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = f"{ruta}.backup-anexos-{stamp}.xls"
    try:
        shutil.copy2(ruta, backup)
        info["bce_backup_anexos"] = backup
    except Exception:
        pass


def _resolver_pagos_directos_con_anexos(periodo):
    info = _asegurar_archivos_periodo(periodo)
    if not info.get("bce") or not os.path.exists(info["bce"]):
        return {"aplicados": [], "avisos": ["No hay reporte BCE disponible para modificar."]}

    res = _resultado_periodo(periodo)
    diag_map = {}
    for item in res.get("diagnostico", []):
        for concepto in (item.get("rubros") or {}):
            diag_map[(item.get("acreedor"), concepto)] = item
    prestamos = res.get("prestamos") or {}

    rows = db_q("""SELECT * FROM conciliaciones
                   WHERE periodo=? AND estado='DIFERENCIA' AND concepto='Desembolsos'
                   ORDER BY id""", (periodo,), fetch=True) or []
    aplicados, avisos = [], []
    for row in rows:
        if (row.get("diferencia") or 0) <= 0:
            continue
        diag = diag_map.get((row["acreedor"], row["concepto"])) or {}
        if diag.get("tipo") != "Pago Directo":
            continue

        objetivo = abs(round(row["diferencia"] or 0, 2))
        prestamo = _prestamo_objetivo_obs(row, prestamos)
        refs_obj = _refs_obs((prestamo.get("referencia") or "") + " " + str(prestamo.get("credito") or ""))
        candidatos = _candidatos_anexo_obs(periodo, row, refs_obj)
        piezas = []
        for cand in candidatos:
            for monto in cand.get("montos") or []:
                if 0 < monto <= objetivo + C.TOLERANCIA:
                    piezas.append({"valor": round(monto, 2), "cand": cand})
        combo = _combinar_montos_obs(piezas, objetivo)
        validado_por_monto = bool(combo)

        if not combo:
            # Respaldo conservador: si existe un PDF/Excel explicitamente nombrado
            # "Pago Directo" para esa cartera pero es escaneado, se usa la diferencia
            # MEF-BCE como monto y se deja auditado en la nota.
            directos = [
                c for c in candidatos
                if "PAGO DIRECTO" in _normal_obs(c.get("nombre", "")) or "DIRECT PAYMENT" in _normal_obs(c.get("texto", ""))
            ]
            if directos:
                combo = [{"valor": objetivo, "cand": directos[0]}]
            else:
                avisos.append(f"{row['acreedor']}: no se encontro anexo de pago directo por {objetivo:,.2f}.")
                continue

        referencia = prestamo.get("referencia") or (row["acreedor"] + " pago directo")
        ajustes_archivo = []
        anexos_usados = []
        for pieza in combo:
            cand = pieza["cand"]
            anexos_usados.append(cand["archivo"])
            nota_monto = "monto validado en anexo" if validado_por_monto else "monto tomado de la diferencia MEF-BCE; anexo sin texto/monto extraible"
            nota = (
                f"Pago directo respaldado por anexo {cand['nombre']} ({nota_monto}). "
                f"Observacion MEF: {diag.get('observacion', '')}"
            )
            ajustes_archivo.append({
                "acreedor": row["acreedor"],
                "concepto": "Desembolsos",
                "referencia": referencia,
                "valor": round(pieza["valor"], 2),
                "nota": nota,
                "prestamista": cand.get("prestamista") or row["acreedor"],
            })

        _respaldar_bce_obs(info)
        try:
            modificado = C.modificar_bce_xls(info["bce"], ajustes_archivo)
        except Exception as exc:
            modificado = False
            avisos.append(f"{row['acreedor']}: no se pudo modificar BCE ({exc}).")
        if not modificado:
            avisos.append(f"{row['acreedor']}: se encontro anexo, pero no se pudo escribir en el .xls BCE.")
            continue

        total = round(sum(a["valor"] for a in ajustes_archivo), 2)
        nota_db = (
            f"Pago directo agregado al BCE desde {len(set(anexos_usados))} anexo(s): "
            + ", ".join(sorted({os.path.basename(a) for a in anexos_usados}))
        )
        ajuste_db = _cerrar_fila_automaticamente(row, nota_db)
        info.setdefault("ajustes", []).extend(ajustes_archivo)
        aplicado = {
            "tipo": "Pago Directo",
            "acreedor": row["acreedor"],
            "concepto": "Desembolsos",
            "valor": total,
            "valor_matriz": ajuste_db,
            "referencia": referencia,
            "nota": nota_db,
            "anexos": sorted({os.path.basename(a) for a in anexos_usados}),
            "validado_por_monto": validado_por_monto,
            "archivo_modificado": True,
        }
        aplicados.append(aplicado)

    return {"aplicados": aplicados, "avisos": avisos}

def _aplicar_ajustes_automaticos(periodo):
    ajustes = []
    res = _resultado_periodo(periodo)
    diag_map = {}
    for item in res.get("diagnostico", []):
        for concepto in (item.get("rubros") or {}):
            diag_map[(item.get("acreedor"), concepto)] = item

    rows = db_q("""SELECT * FROM conciliaciones
                   WHERE periodo=? AND estado='DIFERENCIA'
                   ORDER BY id""", (periodo,), fetch=True) or []

    # 1) Pagos directos respaldados por anexos del mes.
    anexos_resultado = _resolver_pagos_directos_con_anexos(periodo)
    ajustes.extend(anexos_resultado.get("aplicados", []))

    rows = db_q("""SELECT * FROM conciliaciones
                   WHERE periodo=? AND estado='DIFERENCIA'
                   ORDER BY concepto, id""", (periodo,), fetch=True) or []

    # 2) Reclasificaciones compensadas: un rubro sobra en una cartera y falta en otra.
    usados = set()
    por_concepto = {}
    for row in rows:
        por_concepto.setdefault(row["concepto"], []).append(row)

    for concepto, grupo in por_concepto.items():
        for i, a in enumerate(grupo):
            if a["id"] in usados:
                continue
            for b in grupo[i + 1:]:
                if b["id"] in usados:
                    continue
                if abs((a["diferencia"] or 0) + (b["diferencia"] or 0)) <= C.TOLERANCIA:
                    nota_a = (
                        f"Reclasificacion automatica compensada con {b['acreedor']} "
                        f"en {concepto}. Ajuste matriz BCE: {(a['mef'] or 0) - (a['bce'] or 0):,.2f}."
                    )
                    nota_b = (
                        f"Reclasificacion automatica compensada con {a['acreedor']} "
                        f"en {concepto}. Ajuste matriz BCE: {(b['mef'] or 0) - (b['bce'] or 0):,.2f}."
                    )
                    ajuste_a = _cerrar_fila_automaticamente(a, nota_a)
                    ajuste_b = _cerrar_fila_automaticamente(b, nota_b)
                    ajustes.extend([
                        {"tipo": "Reclasificacion", "acreedor": a["acreedor"], "concepto": concepto, "valor": ajuste_a, "nota": nota_a},
                        {"tipo": "Reclasificacion", "acreedor": b["acreedor"], "concepto": concepto, "valor": ajuste_b, "nota": nota_b},
                    ])
                    usados.add(a["id"])
                    usados.add(b["id"])
                    break

    info = PERIODO_FILES.setdefault(periodo, {"bce": None, "mef": None, "ajustes": [], "auto_ajustes": []})
    info.setdefault("auto_ajustes", []).extend(ajustes)
    _persist_period_info(periodo)
    res = _resultado_periodo(periodo)
    res["auto_ajustes"] = ajustes
    res["reporte_url"] = f"/api/exportar_archivo?periodo={periodo}"
    return res



@app.route("/api/resolver_observaciones_anexos", methods=["POST"])
def api_resolver_observaciones_anexos():
    """Busca anexos del mes y corrige pagos directos en el reporte BCE .xls."""
    data = request.get_json(silent=True) or {}
    periodo = (data.get("periodo") or "").strip()
    if not periodo:
        return jsonify({"ok": False, "error": "Falta el periodo."})
    if not db_q("SELECT id FROM conciliaciones WHERE periodo=? LIMIT 1", (periodo,), fetch=True):
        return jsonify({"ok": False, "error": "Primero ejecute la conciliacion del periodo."})

    resultado = _resolver_pagos_directos_con_anexos(periodo)
    info = PERIODO_FILES.setdefault(periodo, {"bce": None, "mef": None, "ajustes": [], "auto_ajustes": []})
    info.setdefault("auto_ajustes", []).extend(resultado.get("aplicados", []))
    _persist_period_info(periodo)
    res = _resultado_periodo(periodo)
    res["anexos_aplicados"] = resultado.get("aplicados", [])
    res["anexos_avisos"] = resultado.get("avisos", [])
    res["auto_ajustes"] = info.get("auto_ajustes", [])
    res["reporte_url"] = f"/api/exportar_archivo?periodo={periodo}"
    return jsonify(res)

@app.route("/api/conciliar_automatico", methods=["POST"])
def api_conciliar_automatico():
    """Conciliacion asistida: ejecuta el cruce y aplica ajustes auditables."""
    avisos = []
    if request.files:
        ruta_bce, ruta_mef, avisos = _clasificar_subidas()
    else:
        data = request.get_json(silent=True) or {}
        periodo = (data.get("periodo") or "").strip()
        if periodo:
            rows = db_q("SELECT id FROM conciliaciones WHERE periodo=? LIMIT 1", (periodo,), fetch=True)
            if rows:
                return jsonify(_aplicar_ajustes_automaticos(periodo))
        ruta_bce = data.get("ruta_bce")
        ruta_mef = data.get("ruta_mef")
        if not ruta_bce or not ruta_mef:
            ruta_bce, ruta_mef = _ultimos_reportes_subidos()

    if not ruta_bce or not os.path.exists(ruta_bce):
        return jsonify({"ok": False, "error": "No encontre un reporte BCE valido para conciliar automaticamente."})
    if not ruta_mef or not os.path.exists(ruta_mef):
        return jsonify({"ok": False, "error": "No encontre un reporte MEF valido para conciliar automaticamente."})

    base = _guardar_conciliacion_desde_rutas(ruta_bce, ruta_mef, avisos)
    res = _aplicar_ajustes_automaticos(base["periodo"])
    res.update({
        "fecha": base.get("fecha"),
        "info_periodo": base.get("info_periodo"),
        "archivo_bce": base.get("archivo_bce"),
        "archivo_mef": base.get("archivo_mef"),
        "avisos": avisos,
    })
    return jsonify(res)


@app.route("/api/exportar_archivo")
def api_exportar_archivo():
    """Descarga el reporte final de conciliacion con notas de ajustes."""
    periodo = (request.args.get("periodo") or "").strip()
    if not periodo:
        return jsonify({"ok": False, "error": "Falta el periodo"})
    rows = db_q("""SELECT acreedor, concepto, mef, bce, diferencia, estado, nota
                   FROM conciliaciones WHERE periodo = ?
                   ORDER BY id""", (periodo,), fetch=True)
    if not rows:
        return jsonify({"ok": False, "error": "No hay datos para ese periodo"})
    filas = [(r["acreedor"], r["concepto"], r["mef"], r["bce"], r["diferencia"], r["estado"], r.get("nota") or "")
             for r in rows]
    nombre = f"Conciliacion_BCE_MEF_{periodo.replace('/', '-')}_final.xlsx"
    ruta = os.path.join(EXPORT_DIR, nombre)
    if not C.reporte_xlsx(filas, ruta):
        return jsonify({"ok": False, "error": "openpyxl no disponible"})
    return send_file(ruta, as_attachment=True, download_name=nombre)



def _info_periodo_desde_codigo(periodo):
    try:
        anio, mes = str(periodo).split("-")[:2]
        idx = int(mes) - 1
        meses = globals().get("MESES_OBS") or [
            "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
            "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
        ]
        nombre = meses[idx].capitalize()
        return {"periodo": periodo, "mes": nombre, "anio": anio, "texto": f"{nombre} {anio}"}
    except Exception:
        return {"periodo": periodo, "mes": "", "anio": "", "texto": periodo}


@app.route("/api/ultimo_periodo_conciliado")
def api_ultimo_periodo_conciliado():
    """Recupera el ultimo periodo completamente conciliado para habilitar el oficio tras refrescar."""
    rows = db_q("""SELECT periodo, MAX(fecha_corrida) AS ultima
                   FROM conciliaciones
                   GROUP BY periodo
                   HAVING SUM(CASE WHEN estado='DIFERENCIA' THEN 1 ELSE 0 END)=0
                      AND COUNT(*) > 0
                   ORDER BY periodo DESC
                   LIMIT 1""", fetch=True) or []
    if not rows:
        return jsonify({"ok": False, "error": "No hay periodos conciliados para emitir oficio."})
    periodo = rows[0]["periodo"]
    res = _resultado_periodo(periodo)
    res["info_periodo"] = _info_periodo_desde_codigo(periodo)
    res["reporte_url"] = f"/api/exportar_archivo?periodo={periodo}"
    return jsonify(res)

# =============================================================================
# FRONT-END (una sola página, estilo BCE oscuro)
# =============================================================================
PANEL_HTML = r"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Gestor Integral de Deuda Externa Pública - Banco Central del Ecuador</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Archivo:wght@400;500;600;700;800&family=Libre+Franklin:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>
  :root{--bg:#0a0f1f;--bg2:#0e1530;--panel:#121b36cc;--panel2:#0f1730;--line:#24314f;
        --txt:#e9eefb;--muted:#94a3c4;--gold:#d9b572;--gold2:#f0d49a;--cyan:#5ad1e6;
        --ok:#34d399;--okbg:#0e2b22;--bad:#f87171;--badbg:#2c1620;--blue:#3b82f6;}
  *{box-sizing:border-box}html,body{margin:0;height:100%}
  body{display:flex;min-height:100vh;color:var(--txt);font-family:"Libre Franklin","Segoe UI",system-ui,sans-serif;font-size:14px;
       background:radial-gradient(1200px 700px at 0% -10%,#16224a 0%,transparent 55%),
                  radial-gradient(1000px 600px at 100% 0%,#0c2a3a 0%,transparent 50%),
                  linear-gradient(180deg,var(--bg),var(--bg2));}
  .side{width:266px;flex:none;background:linear-gradient(180deg,#0c1430,#0a0f22);border-right:1px solid var(--line);display:flex;flex-direction:column;min-height:100vh}
  .brand{padding:24px 22px 20px;border-bottom:1px solid var(--line);text-align:center}
  .brand .crest{width:50px;height:50px;margin:0 auto 10px;border-radius:14px;display:grid;place-items:center;font-size:26px;
        background:linear-gradient(135deg,#1a2murl);background:linear-gradient(135deg,#26345e,#16203f);border:1px solid var(--gold);box-shadow:0 6px 20px #0006}
  .brand h2{margin:0;font-size:14px;letter-spacing:1.5px;color:var(--gold2);font-weight:700}
  .brand .ln{height:2px;width:46px;margin:8px auto;background:linear-gradient(90deg,transparent,var(--gold),transparent)}
  .brand small{color:var(--muted);font-size:11px;letter-spacing:.3px}
  .nav{padding:16px 12px;flex:1}
  .nav .it{display:flex;gap:13px;align-items:center;padding:13px 13px;border-radius:12px;cursor:pointer;color:#b9c4e0;margin-bottom:6px;transition:.15s;border:1px solid transparent}
  .nav .it:hover{background:#ffffff0a}
  .nav .it.act{background:linear-gradient(90deg,#1b2banchor,#16203f00);background:#16203f;border-color:var(--line);color:#fff}
  .nav .it.act{box-shadow:inset 3px 0 0 var(--gold)}
  .nav .it.dis{opacity:.4;cursor:not-allowed}
  .nav .it .num{width:30px;height:30px;border-radius:9px;flex:none;display:grid;place-items:center;font-size:13px;font-weight:800;background:#1c2950;color:var(--gold2);border:1px solid var(--line)}
  .nav .it.act .num{background:var(--gold);color:#10182f;border-color:transparent}
  .nav .it .tt{font-weight:600;font-size:13.5px}.nav .it .ss{font-size:10.5px;color:var(--muted)}
  .icon-sprite{position:absolute;width:0;height:0;overflow:hidden}.nav-symbol svg{width:17px;height:17px;fill:none;stroke:currentColor;stroke-width:1.8;stroke-linecap:round;stroke-linejoin:round}
  .nav-group{margin-bottom:7px}.nav-group-head{width:100%;display:flex;align-items:center;gap:10px;padding:9px 11px;border:1px solid var(--line);border-radius:11px;background:#101a36;color:var(--gold2);cursor:pointer;text-align:left}.nav-group-head .group-icon{width:29px;height:29px;display:grid;place-items:center;border-radius:8px;background:#1c2950}.nav-group-head .group-icon svg{width:17px;height:17px;fill:none;stroke:currentColor;stroke-width:1.8}.nav-group-head strong,.nav-group-head small{display:block}.nav-group-head strong{font-size:11.5px}.nav-group-head small{margin-top:2px;color:var(--muted);font-size:9.5px}.nav-group-head .chevron{margin-left:auto;transition:.18s}.nav-group.closed .chevron{transform:rotate(-90deg)}.nav-group.closed .nav-group-items{display:none}.nav-group-items{padding:6px 0 0 11px;border-left:1px solid var(--line);margin-left:15px}.nav-group-items .it{padding:10px 9px;margin-bottom:3px}.nav-group-items .it .num{width:27px;height:27px;border-radius:8px}.nav-group-items .it .tt{font-size:12px}.nav-group-items .it .ss{font-size:9.5px}
  .conc-module-head{margin-bottom:18px}.conc-module-title{display:flex;align-items:flex-start;justify-content:space-between;gap:14px;margin-bottom:12px}.conc-module-title p{margin:4px 0 0;color:var(--muted);font-size:12px}.conc-flow-tabs{display:grid;grid-template-columns:repeat(auto-fit,minmax(145px,1fr));gap:9px;margin-bottom:12px}.conc-flow-tabs button{display:flex;align-items:center;gap:9px;min-height:60px;padding:10px 12px;border:1px solid var(--line);border-radius:12px;background:#101a36;color:#c7d2ea;text-align:left;cursor:pointer;transition:.18s}.conc-flow-tabs button:hover{border-color:var(--gold)}.conc-flow-tabs button.act{border-color:var(--gold);background:#2a220f}.conc-tab-icon{width:30px;height:30px;display:grid;place-items:center;flex:none;border-radius:8px;background:#1c2950}.conc-tab-icon svg{width:17px;height:17px;fill:none;stroke:currentColor;stroke-width:1.8;stroke-linecap:round;stroke-linejoin:round}.conc-flow-tabs button span:last-child{display:grid;gap:2px}.conc-flow-tabs button strong{font-size:11px}.conc-flow-tabs button small{color:var(--muted);font-size:9px}.conc-portfolio-card{display:flex;align-items:flex-start;gap:16px;padding:14px 16px;margin-bottom:0}.conc-portfolio-label{min-width:150px}.conc-portfolio-label strong,.conc-portfolio-label span{display:block}.conc-portfolio-label strong{color:var(--gold2);font-size:12px}.conc-portfolio-label span{margin-top:3px;color:var(--muted);font-size:10px}.conc-portfolio-list{display:flex;align-items:center;gap:7px;flex-wrap:wrap;color:var(--muted);font-size:11px}.conc-portfolio-list .e{display:inline-flex;align-items:center;padding:6px 9px;border:1px solid var(--line);border-radius:999px;background:#0c1428}.conc-portfolio-list .dot{margin-right:6px}
  .perbox{margin:8px 16px;padding:11px;border:1px solid var(--gold);border-radius:11px;text-align:center;color:var(--gold2);font-weight:700;font-size:13px;background:#1a223f;display:none}
  .estado{padding:16px 20px;border-top:1px solid var(--line);font-size:12px}
  .estado .h{color:var(--gold2);text-transform:uppercase;letter-spacing:1px;margin-bottom:9px;font-size:10px}
  .dot{display:inline-block;width:8px;height:8px;border-radius:50%;margin-right:8px}.dot.ok{background:var(--ok);box-shadow:0 0 6px var(--ok)}.dot.bad{background:var(--bad);box-shadow:0 0 6px var(--bad)}
  .estado .e{margin-bottom:6px;color:#c3cde6}
  .main{flex:1;min-width:0;display:flex;flex-direction:column}
  .top{padding:18px 30px;display:flex;align-items:center;justify-content:space-between;border-bottom:1px solid var(--line);background:#0b122699}
  .top h1{margin:0;font-size:16px;font-weight:600;letter-spacing:.3px}.top h1 b{color:var(--gold2)}
  .top .right{display:flex;gap:14px;align-items:center}
  .badge{padding:6px 14px;border-radius:999px;font-size:12px;font-weight:700;border:1px solid var(--line);color:var(--muted)}
  .badge.set{border-color:var(--gold);color:var(--gold2);background:#1a223f}.badge.ok{border-color:var(--ok);color:var(--ok);background:var(--okbg)}
  .uchip{display:flex;gap:9px;align-items:center;color:var(--muted)}.uav{width:32px;height:32px;border-radius:50%;background:linear-gradient(135deg,var(--gold),#b8924a);color:#10182f;display:grid;place-items:center;font-weight:800}.user-meta{display:grid;gap:1px}.user-meta strong{color:var(--txt);font-size:12px}.user-meta small{font-size:9.5px}.logout-link{padding-left:9px;border-left:1px solid var(--line);color:var(--blue);font-size:10.5px;font-weight:800;text-decoration:none}.logout-link:hover{text-decoration:underline}
  .content{padding:28px 30px;overflow:auto}
  h3.sec{display:flex;align-items:center;gap:11px;margin:0 0 5px;font-size:19px;font-weight:600}
  h3.sec .bar{width:5px;height:22px;border-radius:3px;background:linear-gradient(180deg,var(--gold),#9c7b3e)}
  .sub{color:var(--muted);margin:0 0 22px;font-size:13px}
  .card{background:var(--panel);border:1px solid var(--line);border-radius:16px;padding:24px;margin-bottom:20px;backdrop-filter:blur(8px);box-shadow:0 12px 30px #00000035}
  .drop{border:2px dashed #35466e;border-radius:16px;padding:46px 22px;text-align:center;cursor:pointer;transition:.18s;background:#0e1730}
  .drop:hover{border-color:var(--gold);background:#121c3a}.drop.over{border-color:var(--cyan);background:#10243a}
  .drop.set{border-style:solid;border-color:var(--ok)}
  .drop .ic{font-size:42px;margin-bottom:6px}.drop .ti{font-size:17px;font-weight:600;color:var(--gold2)}.drop .de{color:var(--muted);font-size:13px;margin-top:6px}
  .drop .fl{margin-top:14px;display:flex;flex-direction:column;gap:6px;align-items:center}
  .chipf{background:#16203f;border:1px solid var(--line);border-radius:999px;padding:6px 14px;font-size:12.5px;color:#cdd8f0}
  .chipf.mef{border-color:#5aa9e6}.chipf.bce{border-color:var(--gold)}
  input[type=file]{display:none}
  .footbar{display:flex;align-items:center;justify-content:space-between;gap:16px}
  .btn{background:linear-gradient(135deg,var(--gold),#c69a4f);color:#10182f;border:0;border-radius:11px;padding:13px 26px;font-size:14px;font-weight:800;cursor:pointer;box-shadow:0 8px 22px #d9b57240;letter-spacing:.3px}
  .btn:hover{filter:brightness(1.07)}.btn:disabled{opacity:.4;cursor:not-allowed;box-shadow:none}
  .btn.alt{background:transparent;color:var(--gold2);border:1px solid var(--gold)}
  .btn.gh{background:transparent;color:var(--txt);border:1px solid var(--line)}
  .info{border-left:3px solid var(--gold);background:#121c3a;border-radius:10px;padding:12px 16px;margin-bottom:20px;font-size:12.5px;color:#c7d2ea}
  .frow{display:flex;gap:16px;flex-wrap:wrap;align-items:flex-end}.fld label{display:block;font-size:11px;color:var(--muted);font-weight:700;margin-bottom:6px;text-transform:uppercase;letter-spacing:.4px}
  input[type=text],input[type=number]{background:#0c1428;border:1px solid var(--line);color:var(--txt);border-radius:9px;padding:10px 12px;font-size:14px}
  /* Matriz */
  .scrollx{overflow:auto;border:1px solid var(--line);border-radius:14px;max-height:620px;background:#0c1326}
  .mtable{width:100%;border-collapse:separate;border-spacing:0;font-size:12.5px;min-width:980px}
  .mtable th,.mtable td{padding:11px 12px;white-space:nowrap;border-bottom:1px solid #1c2746}
  .mtable thead th{position:sticky;top:0;background:#101a36;color:var(--gold2);text-align:center;font-size:11px;text-transform:uppercase;letter-spacing:.6px;z-index:2}
  .mtable thead tr.sub th{top:38px;background:#0e1730;color:#9fb0d6;font-size:10px;font-weight:600;letter-spacing:.3px}
  .mtable thead th.acr,.mtable thead th.est{left:0;text-align:left;vertical-align:middle;background:#101a36}
  .mtable th.acr{left:0;z-index:3}
  .mtable td.acr{font-weight:700;color:#fff;position:sticky;left:0;background:#0e1730;z-index:1}
  .mtable td.num{text-align:right;font-variant-numeric:tabular-nums;color:#dbe3f5}
  .mtable td.d{text-align:right;font-weight:700;color:var(--bad)}.mtable td.d.ok{color:var(--ok)}
  .mtable td.sep{border-left:1px solid #1c2746}
  .mtable tbody tr:hover td{background:#13203f}
  .mtable tbody tr:hover td.acr{background:#16213f}
  .mtable tr.bad td.acr{box-shadow:inset 3px 0 0 var(--bad)}.mtable tr.ok td.acr{box-shadow:inset 3px 0 0 var(--ok)}
  .mtable tr.tot td{background:#101a36;color:var(--gold2);font-weight:800;border-top:2px solid var(--gold);position:sticky;bottom:0}
  .mtable tr.tot td.acr{background:#101a36}.mtable tr.tot td.d{color:#ffd9d9}.mtable tr.tot td.d.ok{color:var(--ok)}
  .est-pill{padding:4px 11px;border-radius:999px;font-size:10.5px;font-weight:800;letter-spacing:.3px}
  .est-pill.ok{background:var(--okbg);color:var(--ok);border:1px solid #1c5b44}.est-pill.bad{background:var(--badbg);color:var(--bad);border:1px solid #5b2230}
  .stats{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin-bottom:18px}@media(max-width:860px){.stats{grid-template-columns:repeat(2,1fr)}}
  .stat{background:#101a36;border:1px solid var(--line);border-radius:13px;padding:15px 17px}.stat .n{font-size:24px;font-weight:800}.stat .l{font-size:11px;color:var(--muted);text-transform:uppercase;letter-spacing:.4px}
  .stat.ok .n{color:var(--ok)}.stat.bad .n{color:var(--bad)}.stat.gold .n{color:var(--gold2)}
  .note-ok{background:linear-gradient(135deg,#0e2b22,#10243a);border:1px solid #1c5b44;color:#8ef0c4;border-radius:14px;padding:24px;text-align:center;font-size:15px;font-weight:600}
  .note-warn{background:#241a10;border:1px solid #6b4e1f;color:#f0cd8f;border-radius:14px;padding:18px;font-size:13.5px}
  .hidden{display:none}#msg{font-size:13px;margin-top:10px}
  .diaghead{background:#101a36;border:1px solid var(--line);border-radius:13px;padding:14px 17px;margin-bottom:14px;font-weight:700;color:var(--gold2)}
  .diaghead small{display:block;color:var(--muted);font-weight:400;font-size:11.5px;margin-top:2px}
  .dcard{background:#101a36;border:1px solid var(--line);border-radius:13px;padding:16px;margin-bottom:12px}
  .dcard.pd{border-color:#6b4e1f}.dcard.dc{border-color:#28507a}
  .dcard .dh{display:flex;align-items:center;gap:10px;flex-wrap:wrap;margin-bottom:9px}.dcard .dname{font-weight:800;color:#fff}
  .dtag{padding:3px 10px;border-radius:999px;font-size:10.5px;font-weight:700}.dtag.pd{background:#3a2c10;color:#f0cd8f}.dtag.dc{background:#10283f;color:#7fc2ec}.dtag.ot{background:#1c2746;color:#a7b4d4}
  .dmonto{margin-left:auto;color:var(--bad);font-weight:800;font-size:12.5px}
  .dobs{background:#0c1428;border:1px solid var(--line);border-radius:9px;padding:9px 12px;font-size:12px;margin-bottom:10px;color:#c7d2ea}
  .dgrid{display:grid;grid-template-columns:1fr 1fr;gap:11px}@media(max-width:760px){.dgrid{grid-template-columns:1fr}}
  .dbox{background:#0c1428;border:1px solid var(--line);border-radius:9px;padding:10px 12px}.dbox .bt{font-size:10px;font-weight:800;text-transform:uppercase;margin-bottom:4px;letter-spacing:.4px}.dbox.why .bt{color:#f0cd8f}.dbox.act .bt{color:var(--ok)}.dbox .bd{font-size:12px;color:#c7d2ea}
  textarea{width:100%;height:440px;background:#0c1428;border:1px solid var(--line);border-radius:11px;padding:14px;font-family:Consolas,monospace;font-size:12.5px;color:#dbe3f5}
  h1,h2,h3,.brand h2,.nav .it .tt,.btn,.est-pill,.dtag{font-family:"Archivo","Segoe UI",sans-serif}
  .bce-sun{width:62px;height:62px;display:block;margin:0 auto;filter:drop-shadow(0 3px 8px #0007)}
  .wordmark{font-family:"Archivo",sans-serif;font-weight:800;letter-spacing:.5px;line-height:1}
  .wordmark .b1{color:#fff;font-size:14px}.wordmark .b2{color:var(--gold2);font-size:11px;letter-spacing:2px}
  .topbrand{display:flex;align-items:center;gap:12px}.topbrand .bce-sun{width:34px;height:34px;margin:0}
  .bce-logo{max-width:170px;max-height:84px;display:block;margin:0 auto;filter:drop-shadow(0 3px 8px #0007)}
  .topbrand .bce-logo{max-height:36px;margin:0}
  .flagbar{display:flex;height:6px;width:96px;border-radius:3px;overflow:hidden;margin:9px auto 0;box-shadow:0 1px 4px #0006}
  .flagbar i{flex:1}.flagbar .y{background:#ffd200}.flagbar .b{background:#0033a0}.flagbar .r{background:#ed1c24}
  .logo-slot{position:relative}
  .loan{background:#0c1428;border:1px solid var(--line);border-radius:9px;padding:11px 13px;margin-top:10px}
  .loan .lh{font-size:13px;color:#e9eefb}.loan .lm{display:block;color:var(--muted);font-size:11.5px;margin-top:3px}

  select{background:#0c1428;border:1px solid var(--line);color:var(--txt);border-radius:9px;padding:10px 12px;font-size:14px}
  input[type=checkbox]{accent-color:var(--gold)}
  .caf-panel-grid{display:grid;grid-template-columns:minmax(280px,420px) minmax(0,1fr);gap:18px;align-items:start}
  .caf-path{width:100%;height:72px;min-height:72px;resize:vertical;font-family:Consolas,monospace;font-size:11.5px}
  .caf-wide{flex:1 1 100%}.caf-wide textarea{width:100%}
  .caf-actions{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin-top:14px}
  .caf-check{display:flex;gap:9px;align-items:center;color:#c7d2ea;font-weight:700}
  .caf-check input{width:16px;height:16px}
  .caf-stats{grid-template-columns:160px 220px minmax(0,1fr)}
  .caf-status{color:var(--muted);font-size:12.5px}
  .caf-table{min-width:1040px}.caf-table td code{display:block;max-width:360px;white-space:normal;overflow-wrap:anywhere;color:#cdd8f0;font-size:11px}
  .caf-amount{width:130px;text-align:right;font-family:Consolas,monospace}
  .source-pill{display:inline-block;border:1px solid var(--line);border-radius:999px;background:#101a36;color:#c7d2ea;padding:4px 10px;font-size:10.5px;font-weight:800}
  .ok-text{color:var(--ok);font-weight:800}.warn-text{color:var(--gold2);font-weight:800}
  .caf-log{height:160px;overflow:auto;background:#0c1428;border:1px solid var(--line);border-radius:11px;padding:12px;color:#c7d2ea;font-family:Consolas,monospace;font-size:12px;white-space:pre-wrap}
  @media(max-width:980px){.caf-panel-grid,.caf-stats{grid-template-columns:1fr}}


  .firm-panel-grid{display:grid;grid-template-columns:minmax(300px,390px) minmax(0,1fr);gap:18px;align-items:start}
  .firm-actions{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin-top:14px}
  .firm-date-row{display:flex;gap:10px;align-items:end;flex-wrap:wrap}
  .firm-stats{grid-template-columns:repeat(4,minmax(0,1fr))}
  .firm-paths{display:grid;grid-template-columns:120px minmax(0,1fr);gap:8px 12px;font-size:12px;color:#c7d2ea}
  .firm-paths dt{color:var(--muted);font-weight:800}.firm-paths dd{margin:0;overflow-wrap:anywhere}
  .firm-check{display:flex;gap:9px;align-items:center;color:#c7d2ea;font-weight:700;margin-top:10px}
  .firm-check input{width:16px;height:16px}
  .firm-slicer{display:flex;gap:8px;flex-wrap:wrap;max-height:105px;overflow:auto;margin-top:10px}
  .firm-slicer button{border:1px solid var(--line);border-radius:999px;background:#101a36;color:#c7d2ea;padding:6px 10px;font-size:11px}
  .firm-slicer button.act{background:var(--gold);border-color:var(--gold);color:#0b1020}
  .firm-table{min-width:980px}.firm-table td code{display:block;max-width:460px;white-space:normal;overflow-wrap:anywhere;color:#cdd8f0;font-size:11px}
  .firm-tag{display:inline-flex;border-radius:999px;padding:4px 9px;font-weight:800;font-size:11px;background:#1d2946;color:#c7d2ea;white-space:nowrap}
  .firm-tag.move{background:#12361f;color:#86efac}.firm-tag.clean{background:#3b2512;color:#fbbf24}.firm-tag.warn{background:#3b1620;color:#fca5a5}.firm-tag.done{background:#0f2f3d;color:#7dd3fc}
  .firm-log{height:130px;overflow:auto;background:#0c1428;border:1px solid var(--line);border-radius:11px;padding:12px;color:#c7d2ea;font-family:Consolas,monospace;font-size:12px;white-space:pre-wrap}
  @media(max-width:980px){.firm-panel-grid,.firm-stats{grid-template-columns:1fr}}

  .comp-panel-grid{display:grid;grid-template-columns:minmax(300px,390px) minmax(0,1fr);gap:18px;align-items:start}
  .comp-actions{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin-top:14px}
  .comp-date-row{display:flex;gap:10px;align-items:end;flex-wrap:wrap}
  .comp-stats{grid-template-columns:repeat(5,minmax(0,1fr))}
  .comp-paths{display:grid;grid-template-columns:125px minmax(0,1fr);gap:8px 12px;font-size:12px;color:#c7d2ea}
  .comp-paths dt{color:var(--muted);font-weight:800}.comp-paths dd{margin:0;overflow-wrap:anywhere}
  .comp-check{display:flex;gap:9px;align-items:center;color:#c7d2ea;font-weight:700;margin-top:10px}
  .comp-check input{width:16px;height:16px}
  .comp-table{min-width:1240px}.comp-table td code{display:block;max-width:360px;white-space:normal;overflow-wrap:anywhere;color:#cdd8f0;font-size:11px}
  .comp-tag{display:inline-flex;border-radius:999px;padding:4px 9px;font-weight:800;font-size:11px;background:#1d2946;color:#c7d2ea;white-space:nowrap}
  .comp-tag.ready{background:#12361f;color:#86efac}.comp-tag.ready_no_ack{background:#3b2512;color:#fbbf24}
  .comp-tag.existing,.comp-tag.existing_no_ack{background:#0f2f3d;color:#7dd3fc}.comp-tag.missing_destination,.comp-tag.missing_reference,.comp-tag.no_text,.comp-tag.error{background:#3b1620;color:#fca5a5}
  .comp-log{height:130px;overflow:auto;background:#0c1428;border:1px solid var(--line);border-radius:11px;padding:12px;color:#c7d2ea;font-family:Consolas,monospace;font-size:12px;white-space:pre-wrap}
  @media(max-width:1100px){.comp-panel-grid,.comp-stats{grid-template-columns:1fr}}

  .comp-panel-grid{display:grid;grid-template-columns:minmax(300px,390px) minmax(0,1fr);gap:18px;align-items:start}
  .comp-actions{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin-top:14px}
  .comp-date-row{display:flex;gap:10px;align-items:end;flex-wrap:wrap}
  .comp-stats{grid-template-columns:repeat(6,minmax(0,1fr))}
  .comp-paths{display:grid;grid-template-columns:125px minmax(0,1fr);gap:8px 12px;font-size:12px;color:#c7d2ea}
  .comp-paths dt{color:var(--muted);font-weight:800}.comp-paths dd{margin:0;overflow-wrap:anywhere}
  .comp-check{display:flex;gap:9px;align-items:center;color:#c7d2ea;font-weight:700;margin-top:10px}
  .comp-check input{width:16px;height:16px}
  .comp-table{min-width:1240px}.comp-table td code{display:block;max-width:360px;white-space:normal;overflow-wrap:anywhere;color:#cdd8f0;font-size:11px}
  .comp-tag{display:inline-flex;border-radius:999px;padding:4px 9px;font-weight:800;font-size:11px;background:#1d2946;color:#c7d2ea;white-space:nowrap}
  .comp-tag.ready{background:#12361f;color:#86efac}.comp-tag.ready_no_ack{background:#3b2512;color:#fbbf24}
  .comp-tag.existing,.comp-tag.existing_no_ack{background:#0f2f3d;color:#7dd3fc}.comp-tag.fallback_no_text,.comp-tag.fallback_missing_reference,.comp-tag.fallback_missing_destination,.comp-tag.fallback_error{background:#3b2512;color:#fbbf24}
  .comp-tag.missing_destination,.comp-tag.missing_reference,.comp-tag.no_text,.comp-tag.error{background:#3b1620;color:#fca5a5}
  .comp-log{height:130px;overflow:auto;background:#0c1428;border:1px solid var(--line);border-radius:11px;padding:12px;color:#c7d2ea;font-family:Consolas,monospace;font-size:12px;white-space:pre-wrap}
  @media(max-width:1100px){.comp-panel-grid,.comp-stats{grid-template-columns:1fr}}

  .firm-panel-grid{display:grid;grid-template-columns:minmax(300px,390px) minmax(0,1fr);gap:18px;align-items:start}
  .firm-actions{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin-top:14px}
  .firm-folder-actions{display:flex;gap:10px;flex-wrap:wrap;align-items:center;margin-top:12px}
  .firm-date-row{display:flex;gap:10px;align-items:end;flex-wrap:wrap}
  .firm-stats{grid-template-columns:repeat(4,minmax(0,1fr))}
  .firm-paths{display:grid;grid-template-columns:120px minmax(0,1fr);gap:8px 12px;font-size:12px;color:#c7d2ea}
  .firm-paths dt{color:var(--muted);font-weight:800}.firm-paths dd{margin:0;overflow-wrap:anywhere}
  .firm-check{display:flex;gap:9px;align-items:center;color:#c7d2ea;font-weight:700;margin-top:10px}
  .firm-check input{width:16px;height:16px}
  .firm-monitor-row{display:flex;gap:10px;align-items:end;flex-wrap:wrap;margin-top:12px}
  .firm-monitor-row .mini{max-width:115px}
  .firm-alert{border:1px solid var(--line);border-radius:11px;background:#0c1428;color:#c7d2ea;padding:12px;margin-bottom:14px;font-weight:800}
  .firm-alert.on{border-color:#f59e0b;background:#2a1c08;color:#fde68a}
  .firm-alert.ready{border-color:#22c55e;background:#082512;color:#bbf7d0}
  .firm-slicer{display:flex;gap:8px;flex-wrap:wrap;max-height:105px;overflow:auto;margin-top:10px}
  .firm-slicer button{border:1px solid var(--line);border-radius:999px;background:#101a36;color:#c7d2ea;padding:6px 10px;font-size:11px}
  .firm-slicer button.act{background:var(--gold);border-color:var(--gold);color:#0b1020}
  .firm-table{min-width:980px}.firm-table td code{display:block;max-width:460px;white-space:normal;overflow-wrap:anywhere;color:#cdd8f0;font-size:11px}
  .firm-tag{display:inline-flex;border-radius:999px;padding:4px 9px;font-weight:800;font-size:11px;background:#1d2946;color:#c7d2ea;white-space:nowrap}
  .firm-tag.move{background:#12361f;color:#86efac}.firm-tag.clean{background:#3b2512;color:#fbbf24}.firm-tag.warn{background:#3b1620;color:#fca5a5}.firm-tag.done{background:#0f2f3d;color:#7dd3fc}
  .firm-log{height:130px;overflow:auto;background:#0c1428;border:1px solid var(--line);border-radius:11px;padding:12px;color:#c7d2ea;font-family:Consolas,monospace;font-size:12px;white-space:pre-wrap}
  @media(max-width:980px){.firm-panel-grid,.firm-stats{grid-template-columns:1fr}}

  /* Comprobantes Contables v2: motor BRYAN + STEVEN del Gestor 5088. */
  .comp-head-row{display:flex;justify-content:space-between;align-items:flex-start;gap:18px;flex-wrap:wrap}
  .comp-head-row .sub{margin-bottom:0}.comp-last-run{color:var(--muted);font-size:12px;max-width:290px}
  .comp-notice{display:flex;gap:8px 16px;align-items:center;flex-wrap:wrap;margin:18px 0;padding:14px 16px;border:1px solid #166534;border-radius:12px;background:#0b2b1b;color:#bbf7d0}
  .comp-notice strong{color:#86efac}.comp-notice span{font-size:12px}.comp-notice.warn{border-color:#a16207;background:#2a1c08;color:#fde68a}.comp-notice.bad{border-color:#991b1b;background:#32131a;color:#fecaca}
  .comp-stats-v2{grid-template-columns:repeat(5,minmax(0,1fr));margin-bottom:18px}
  .comp-config-card details>summary{display:flex;justify-content:space-between;align-items:center;gap:12px;cursor:pointer;color:#fff;list-style:none}.comp-config-card details>summary::-webkit-details-marker{display:none}.comp-config-card details>summary span{color:var(--gold2);font-size:12px}
  .comp-config-grid{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin-top:18px;padding-top:16px;border-top:1px solid var(--line)}
  .comp-config-grid .fld input{width:100%;font-family:Consolas,monospace;font-size:11px}.comp-config-grid .comp-actions{grid-column:1/-1}
  .comp-table-actions{display:flex;align-items:center;justify-content:flex-end;gap:14px;flex-wrap:wrap;margin:-4px 0 14px}.comp-table-actions .comp-check{margin:0 auto 0 0}
  .comp-table-v2{min-width:1460px}.comp-table-v2 td{vertical-align:top}.comp-table-v2 input[type=checkbox]{width:16px;height:16px}.comp-table-v2 small{display:block;color:var(--muted);font-size:10.5px;margin-top:3px}
  .comp-owner{display:inline-flex;border-radius:7px;padding:4px 8px;background:#102f48;color:#7dd3fc;font-size:10px;font-weight:900;letter-spacing:.05em}.comp-owner.steven{background:#3b2b10;color:#fde68a}
  .comp-code{display:block;color:#fff;font-weight:800;white-space:nowrap}.comp-ref{color:#7dd3fc;font-weight:750;white-space:nowrap}.comp-dest{display:block;max-width:245px;color:#e9eefb;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.comp-copy{border:0;background:none;color:#60a5fa;padding:3px 0;font-size:10.5px;cursor:pointer}
  .comp-match{display:block;max-width:175px;color:#fff;font-size:11px;font-weight:750}.comp-score{color:var(--muted);font-size:10.5px}.comp-status{display:inline-flex;border-radius:999px;padding:5px 9px;font-weight:850;font-size:10.5px;white-space:nowrap;background:#1d2946;color:#c7d2ea}.comp-status.ready,.comp-status.partial{background:#12361f;color:#86efac}.comp-status.pending{background:#3b2512;color:#fbbf24}.comp-status.archived{background:#0f2f3d;color:#7dd3fc}.comp-status.probable,.comp-status.review{background:#3b2512;color:#fde68a}.comp-status.error,.comp-status.conflict{background:#3b1620;color:#fca5a5}
  .comp-detail summary{cursor:pointer;color:#7dd3fc;font-weight:750}.comp-detail-box{width:min(430px,70vw);margin-top:8px;padding:11px;border:1px solid var(--line);border-radius:9px;background:#0c1428;color:#c7d2ea;line-height:1.45;overflow-wrap:anywhere}.comp-detail-box p{margin:5px 0}.comp-detail-box strong{color:#fff}

  /* Tema claro solicitado: se limita al módulo de Comprobantes y ACK. */
  #v-comprobantes{min-height:calc(100vh - 76px);margin:-28px -30px;padding:28px 30px;background:#f4f7fa;color:#24364a}
  #v-comprobantes h3.sec{color:#0d2743}#v-comprobantes .sub{color:#607487}#v-comprobantes .comp-last-run{color:#6e8091}
  #v-comprobantes .card{background:#fff;border-color:#dce5ed;box-shadow:0 10px 28px rgba(17,42,68,.08);backdrop-filter:none}
  #v-comprobantes .stat{background:#fff;border-color:#dce5ed;box-shadow:0 7px 20px rgba(17,42,68,.06)}
  #v-comprobantes .stat .n{color:#173a5e}#v-comprobantes .stat.ok .n{color:#167553}#v-comprobantes .stat.gold .n{color:#9a6211}#v-comprobantes .stat.bad .n{color:#a63e48}#v-comprobantes .stat .l{color:#708293}
  #v-comprobantes .comp-notice{border-color:#bfe6d5;background:#e8f7f0;color:#3e6c5d}#v-comprobantes .comp-notice strong{color:#135c45}
  #v-comprobantes .comp-notice.warn{border-color:#eed7a2;background:#fff5dd;color:#765313}#v-comprobantes .comp-notice.warn strong{color:#765313}
  #v-comprobantes .comp-notice.bad{border-color:#f0c5ca;background:#ffedf0;color:#8f3540}#v-comprobantes .comp-notice.bad strong{color:#8f3540}
  #v-comprobantes .comp-config-card details>summary{color:#173a5e}#v-comprobantes .comp-config-card details>summary span{color:#145ca8}
  #v-comprobantes .comp-config-grid{border-top-color:#e3e9ef}#v-comprobantes .fld label{color:#607487}
  #v-comprobantes input[type=text]{background:#fbfcfd;border-color:#cbd5de;color:#173a5e}#v-comprobantes input[type=text]:focus{border-color:#1871c9;outline:3px solid rgba(24,113,201,.12)}
  #v-comprobantes .comp-check{color:#4b6073}#v-comprobantes input[type=checkbox]{accent-color:#145ca8}
  #v-comprobantes .btn.gh{background:#fff;border-color:#cbd5de;color:#173a5e;box-shadow:0 4px 12px rgba(15,43,70,.07)}
  #v-comprobantes .btn:not(.gh){background:#145ca8;color:#fff;box-shadow:0 7px 18px rgba(20,92,168,.2)}
  #v-comprobantes .diaghead{background:#f7f9fb;border-color:#dce5ed;color:#173a5e}#v-comprobantes .diaghead small{color:#708293}
  #v-comprobantes .scrollx{background:#fff;border-color:#dce5ed}
  #v-comprobantes .mtable th,#v-comprobantes .mtable td{border-bottom-color:#e8edf2;color:#4b6073}
  #v-comprobantes .mtable thead th{background:#f7f9fb;color:#607487}
  #v-comprobantes .mtable tbody tr:hover td{background:#f7fbff}
  #v-comprobantes .comp-code,#v-comprobantes .comp-match{color:#173a5e}#v-comprobantes .comp-ref,#v-comprobantes .comp-copy,#v-comprobantes .comp-detail summary{color:#145ca8}#v-comprobantes .comp-dest{color:#24364a}#v-comprobantes .comp-score,#v-comprobantes .comp-table-v2 small,#v-comprobantes .muted{color:#7d8d9c}
  #v-comprobantes .comp-owner{background:#e8f3fb;color:#145c89}#v-comprobantes .comp-owner.steven{background:#fff2cf;color:#6b4b12}
  #v-comprobantes .comp-status{background:#edf2f6;color:#4b6073}#v-comprobantes .comp-status.ready,#v-comprobantes .comp-status.partial{background:#e8f7f0;color:#167553}#v-comprobantes .comp-status.pending{background:#f2edfb;color:#7557a8}#v-comprobantes .comp-status.archived{background:#eaf4ff;color:#145ca8}#v-comprobantes .comp-status.probable,#v-comprobantes .comp-status.review{background:#fff5dd;color:#9a6211}#v-comprobantes .comp-status.error,#v-comprobantes .comp-status.conflict{background:#ffedf0;color:#a63e48}
  #v-comprobantes .comp-detail-box{border-color:#dce5ed;background:#fff;color:#4b6073;box-shadow:0 12px 28px rgba(17,42,68,.11)}#v-comprobantes .comp-detail-box strong{color:#173a5e}
  #v-comprobantes .comp-log{background:#f7f9fb;border-color:#dce5ed;color:#41576b}
  .comp-control-card{margin-bottom:18px;padding:0;overflow:hidden}.comp-control-head{display:flex;justify-content:space-between;gap:16px;align-items:center;padding:17px 19px;border-bottom:1px solid #e3e9ef}.comp-control-head strong{display:block;color:#173a5e}.comp-control-head span{display:block;color:#708293;font-size:11px;margin-top:3px}.comp-filter-meta{font-weight:800;color:#145ca8!important;text-align:right;white-space:nowrap}
  .comp-view-tabs{display:flex;gap:8px;flex-wrap:wrap;padding:14px 19px 8px}.comp-view-tabs button{border:1px solid #d4dee7;border-radius:999px;background:#f7f9fb;color:#526a7f;padding:8px 12px;font-size:11px;font-weight:800;cursor:pointer}.comp-view-tabs button:hover{border-color:#7eacd5;color:#145ca8}.comp-view-tabs button.act{border-color:#145ca8;background:#e8f3ff;color:#145ca8;box-shadow:0 4px 12px rgba(20,92,168,.1)}.comp-view-tabs b{display:inline-flex;align-items:center;justify-content:center;min-width:20px;margin-left:5px;padding:1px 5px;border-radius:999px;background:#fff;font-size:10px}
  .comp-filter-grid{display:grid;grid-template-columns:180px 160px minmax(220px,1fr);gap:12px;padding:10px 19px 18px}.comp-filter-grid label{display:flex;flex-direction:column;gap:5px;color:#607487;font-size:10px;font-weight:850;text-transform:uppercase;letter-spacing:.04em}.comp-filter-grid select,.comp-filter-grid input{width:100%;height:39px;border:1px solid #cbd5de;border-radius:9px;background:#fbfcfd;color:#173a5e;padding:0 11px;font-family:inherit;font-size:12px;font-weight:600;outline:none}.comp-filter-grid select:focus,.comp-filter-grid input:focus{border-color:#1871c9;box-shadow:0 0 0 3px rgba(24,113,201,.12)}
  .comp-empty-filter{padding:28px!important;text-align:center!important;color:#708293!important}.comp-priority-note{display:block;margin-top:4px;font-size:10px;color:#708293}.comp-priority-note.action{color:#167553;font-weight:800}.comp-priority-note.wait{color:#7557a8;font-weight:800}.comp-priority-note.review{color:#9a6211;font-weight:800}
  @media(max-width:1100px){.comp-stats-v2{grid-template-columns:repeat(2,minmax(0,1fr))}.comp-config-grid{grid-template-columns:1fr}.comp-config-grid .comp-actions{grid-column:auto}}
  @media(max-width:760px){.comp-control-head{align-items:flex-start;flex-direction:column}.comp-filter-meta{text-align:left}.comp-filter-grid{grid-template-columns:1fr}.comp-view-tabs{flex-wrap:nowrap;overflow:auto}}

  /* Archivo Quipux: bandeja automática de oficio, estado y formulario. */
  .quipux-nav-badge{display:inline-flex;min-width:19px;height:19px;padding:0 6px;align-items:center;justify-content:center;border-radius:999px;background:#a63e48;color:#fff;font-size:10px;font-weight:850;margin-left:5px}.quipux-nav-badge.hidden{display:none}
  #v-archivoquipux{min-height:calc(100vh - 76px);margin:-28px -30px;padding:28px 30px;background:#eef3f7;color:#24364a}
  #v-archivoquipux h3.sec{color:#0d2743}#v-archivoquipux .sub{color:#607487}
  .quipux-head{display:flex;justify-content:space-between;gap:18px;align-items:flex-start;flex-wrap:wrap}.quipux-head-actions{display:flex;gap:10px;align-items:center;flex-wrap:wrap}.quipux-last{color:#6e8091;font-size:12px;max-width:300px}
  .quipux-notice{display:flex;gap:8px 16px;align-items:center;flex-wrap:wrap;margin:18px 0;padding:14px 16px;border:1px solid #bfe6d5;border-radius:12px;background:#e8f7f0;color:#3e6c5d}.quipux-notice strong{color:#135c45}.quipux-notice span{font-size:12px}.quipux-notice.warn{border-color:#eed7a2;background:#fff5dd;color:#765313}.quipux-notice.warn strong{color:#765313}.quipux-notice.bad{border-color:#f0c5ca;background:#ffedf0;color:#8f3540}.quipux-notice.bad strong{color:#8f3540}
  .quipux-stats{grid-template-columns:repeat(5,minmax(0,1fr));margin-bottom:18px}
  #v-archivoquipux .card,#v-archivoquipux .stat{background:#fff;border-color:#dce5ed;box-shadow:0 8px 24px rgba(17,42,68,.07);backdrop-filter:none}#v-archivoquipux .stat .n{color:#173a5e}#v-archivoquipux .stat.ok .n{color:#167553}#v-archivoquipux .stat.gold .n{color:#9a6211}#v-archivoquipux .stat.bad .n{color:#a63e48}#v-archivoquipux .stat .l{color:#708293}
  .quipux-config details>summary{display:flex;justify-content:space-between;gap:12px;cursor:pointer;color:#173a5e;list-style:none}.quipux-config details>summary::-webkit-details-marker{display:none}.quipux-config details>summary span{color:#145ca8;font-size:12px}.quipux-config-grid{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin-top:18px;padding-top:16px;border-top:1px solid #e3e9ef}.quipux-config-grid .fld input{width:100%;font-family:Consolas,monospace;font-size:11px;background:#fbfcfd;border-color:#cbd5de;color:#173a5e}.quipux-config-grid .quipux-actions{grid-column:1/-1}.quipux-check{display:flex;gap:9px;align-items:center;color:#4b6073;font-weight:700}.quipux-check input{width:16px;height:16px;accent-color:#145ca8}.quipux-actions{display:flex;gap:10px;flex-wrap:wrap}
  #v-archivoquipux .btn.gh{background:#fff;border-color:#cbd5de;color:#173a5e;box-shadow:0 4px 12px rgba(15,43,70,.07)}#v-archivoquipux .btn:not(.gh){background:#145ca8;color:#fff;box-shadow:0 7px 18px rgba(20,92,168,.2)}
  #v-archivoquipux .diaghead{background:#f7f9fb;border-color:#dce5ed;color:#173a5e}#v-archivoquipux .diaghead small{color:#708293}#v-archivoquipux .scrollx{background:#fff;border-color:#dce5ed}#v-archivoquipux .mtable th,#v-archivoquipux .mtable td{border-bottom-color:#e8edf2;color:#4b6073}#v-archivoquipux .mtable thead th{background:#f7f9fb;color:#607487}#v-archivoquipux .mtable tbody tr:hover td{background:#f7fbff}
  .quipux-table{min-width:1370px}.quipux-table td{vertical-align:top}.quipux-file{display:block;max-width:220px;color:#173a5e;font-weight:800;overflow-wrap:anywhere}.quipux-type{display:inline-flex;padding:4px 8px;border-radius:7px;background:#e8f3fb;color:#145c89;font-size:10px;font-weight:900}.quipux-dest{display:block;max-width:270px;color:#24364a;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.quipux-target{display:block;color:#145ca8;font-size:11px;margin-top:4px}.quipux-match{display:block;max-width:180px;color:#173a5e;font-size:11px;font-weight:750}.quipux-score{color:#7d8d9c;font-size:10.5px}.quipux-status{display:inline-flex;border-radius:999px;padding:5px 9px;font-weight:850;font-size:10.5px;white-space:nowrap;background:#edf2f6;color:#4b6073}.quipux-status.archived{background:#e8f7f0;color:#167553}.quipux-status.review{background:#fff5dd;color:#9a6211}.quipux-status.error,.quipux-status.conflict{background:#ffedf0;color:#a63e48}.quipux-status.ignored,.quipux-status.historical{background:#edf2f6;color:#607487}.quipux-package{display:block;font-size:10.5px;color:#607487;margin-top:5px}.quipux-detail summary{cursor:pointer;color:#145ca8;font-weight:750}.quipux-detail-box{width:min(430px,70vw);margin-top:8px;padding:11px;border:1px solid #dce5ed;border-radius:9px;background:#fff;color:#4b6073;box-shadow:0 12px 28px rgba(17,42,68,.11);line-height:1.45;overflow-wrap:anywhere}.quipux-detail-box p{margin:5px 0}.quipux-detail-box strong{color:#173a5e}
  @media(max-width:1100px){.quipux-stats{grid-template-columns:repeat(2,minmax(0,1fr))}.quipux-config-grid{grid-template-columns:1fr}.quipux-config-grid .quipux-actions{grid-column:auto}}

  /* Oficios de Contratos de Agencia Fiscal: monitor separado y seguro. */
  .contract-nav-badge{display:inline-flex;min-width:19px;height:19px;padding:0 6px;align-items:center;justify-content:center;border-radius:999px;background:#a63e48;color:#fff;font-size:10px;font-weight:850;margin-left:5px}.contract-nav-badge.hidden{display:none}
  #v-contratosagencia{min-height:calc(100vh - 76px);margin:-28px -30px;padding:28px 30px;background:#eef3f7;color:#24364a}
  #v-contratosagencia h3.sec{color:#0d2743}#v-contratosagencia .sub{color:#607487}
  .contract-head{display:flex;justify-content:space-between;gap:18px;align-items:flex-start;flex-wrap:wrap}.contract-head-actions{display:flex;gap:10px;align-items:center;flex-wrap:wrap}.contract-last{color:#6e8091;font-size:12px;max-width:320px}
  .contract-notice{display:flex;gap:8px 16px;align-items:center;flex-wrap:wrap;margin:18px 0;padding:14px 16px;border:1px solid #bfe6d5;border-radius:12px;background:#e8f7f0;color:#3e6c5d}.contract-notice strong{color:#135c45}.contract-notice span{font-size:12px}.contract-notice.warn{border-color:#eed7a2;background:#fff5dd;color:#765313}.contract-notice.warn strong{color:#765313}.contract-notice.bad{border-color:#f0c5ca;background:#ffedf0;color:#8f3540}.contract-notice.bad strong{color:#8f3540}
  .contract-stats{grid-template-columns:repeat(5,minmax(0,1fr));margin-bottom:18px}
  #v-contratosagencia .card,#v-contratosagencia .stat{background:#fff;border-color:#dce5ed;box-shadow:0 8px 24px rgba(17,42,68,.07);backdrop-filter:none}#v-contratosagencia .stat .n{color:#173a5e}#v-contratosagencia .stat.ok .n{color:#167553}#v-contratosagencia .stat.gold .n{color:#9a6211}#v-contratosagencia .stat.bad .n{color:#a63e48}#v-contratosagencia .stat .l{color:#708293}
  .contract-config details>summary{display:flex;justify-content:space-between;gap:12px;cursor:pointer;color:#173a5e;list-style:none}.contract-config details>summary::-webkit-details-marker{display:none}.contract-config details>summary span{color:#145ca8;font-size:12px}.contract-config-grid{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin-top:18px;padding-top:16px;border-top:1px solid #e3e9ef}.contract-config-grid .fld input{width:100%;font-family:Consolas,monospace;font-size:11px;background:#fbfcfd;border-color:#cbd5de;color:#173a5e}.contract-config-grid .contract-actions{grid-column:1/-1}.contract-check{display:flex;gap:9px;align-items:center;color:#4b6073;font-weight:700}.contract-check input{width:16px;height:16px;accent-color:#145ca8}.contract-actions{display:flex;gap:10px;flex-wrap:wrap}
  #v-contratosagencia .btn.gh{background:#fff;border-color:#cbd5de;color:#173a5e;box-shadow:0 4px 12px rgba(15,43,70,.07)}#v-contratosagencia .btn:not(.gh){background:#145ca8;color:#fff;box-shadow:0 7px 18px rgba(20,92,168,.2)}
  #v-contratosagencia .diaghead{background:#f7f9fb;border-color:#dce5ed;color:#173a5e}#v-contratosagencia .diaghead small{color:#708293}#v-contratosagencia .scrollx{background:#fff;border-color:#dce5ed}#v-contratosagencia .mtable th,#v-contratosagencia .mtable td{border-bottom-color:#e8edf2;color:#4b6073}#v-contratosagencia .mtable thead th{background:#f7f9fb;color:#607487}#v-contratosagencia .mtable tbody tr:hover td{background:#f7fbff}
  .contract-table{min-width:1380px}.contract-table td{vertical-align:top}.contract-file{display:block;max-width:210px;color:#173a5e;font-weight:800;overflow-wrap:anywhere}.contract-office{display:block;color:#145ca8;font-weight:800}.contract-subject{display:block;max-width:270px;font-size:11px;line-height:1.4}.contract-ref{display:block;font-weight:750;color:#173a5e}.contract-dest{display:block;max-width:280px;color:#24364a;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.contract-target{display:block;max-width:285px;color:#145ca8;font-size:11px;overflow-wrap:anywhere}.contract-match{display:block;max-width:190px;color:#173a5e;font-size:11px;font-weight:750}.contract-score{color:#7d8d9c;font-size:10.5px}.contract-patterns{display:flex;gap:7px;flex-wrap:wrap;margin-top:12px}.contract-pattern{display:inline-flex;gap:6px;padding:6px 9px;border-radius:999px;background:#eef4f8;color:#526a7f;font-size:10.5px}.contract-pattern b{color:#145ca8}
  @media(max-width:1100px){.contract-stats{grid-template-columns:repeat(2,minmax(0,1fr))}.contract-config-grid{grid-template-columns:1fr}.contract-config-grid .contract-actions{grid-column:auto}}

  /* Mensajes firmados v2: flujo controlado importado del Gestor 5088. */
  .firm-stats-v2{grid-template-columns:repeat(5,minmax(0,1fr))}
  .firm-flow-tabs{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px;margin:4px 0 14px}
  .firm-flow-tabs button{display:flex;align-items:center;justify-content:space-between;gap:12px;min-height:64px;padding:12px 15px;border:1px solid var(--line);border-radius:12px;background:#101a36;color:#c7d2ea;text-align:left;cursor:pointer;transition:.18s ease}
  .firm-flow-tabs button:hover{border-color:#d3ab56;transform:translateY(-1px)}.firm-flow-tabs button.act{border-color:var(--gold);background:#2a220f;box-shadow:0 0 0 2px #d3ab5625}
  .firm-flow-tabs button span{font-weight:850}.firm-flow-tabs button small{display:inline-flex;padding:4px 8px;border-radius:999px;background:#1d2946;color:var(--muted);font-size:10px;font-weight:800;white-space:nowrap}.firm-flow-tabs button.act small{background:#5b461a;color:#fde68a}
  .firm-flow-heading{display:flex;align-items:center;justify-content:space-between;gap:14px;margin-bottom:12px}.firm-flow-heading>div{display:grid;gap:3px}.firm-flow-heading strong{font-size:17px;color:#fff}.firm-flow-heading span:not(.firm-monitor-badge){font-size:12px;color:var(--muted)}
  .firm-monitor-badge{display:inline-flex;padding:7px 10px;border-radius:999px;border:1px solid #315c4d;background:#0e3025;color:#86efac;font-size:11px;font-weight:850;white-space:nowrap}
  .firm-route-card{display:grid;grid-template-columns:110px minmax(0,1fr) auto;gap:9px 13px;align-items:center}
  .firm-route-card .label{color:var(--muted);font-size:12px;font-weight:800}.firm-route-card code{color:#cdd8f0;font-size:11px;overflow-wrap:anywhere}
  .firm-route-card .firm-date{color:var(--gold2);font-weight:800}.firm-head-actions{display:flex;gap:10px;flex-wrap:wrap;justify-content:flex-end;margin-bottom:14px}
  .firm-signature-panel{border:1px solid #d69a2d;border-radius:12px;background:#2a1c08;padding:15px;margin-bottom:14px}
  .firm-signature-head{display:flex;justify-content:space-between;align-items:center;gap:12px}.firm-signature-head strong{color:#fde68a}.firm-signature-head span{color:#fbbf24;font-weight:800}
  .firm-signature-list{display:grid;grid-template-columns:repeat(auto-fit,minmax(250px,1fr));gap:9px;margin-top:12px}.firm-signature-item{border:1px solid #6b4a16;border-left:4px solid #f59e0b;border-radius:9px;background:#101a36;padding:10px 12px}.firm-signature-item.second_signature{border-left-color:#ef4444}.firm-signature-item strong,.firm-signature-item small{display:block}.firm-signature-item small{color:var(--muted);margin-top:3px}
  .firm-table-v2{min-width:1160px}.firm-table-v2 td code{display:block;max-width:390px;white-space:normal;overflow-wrap:anywhere;color:#cdd8f0;font-size:11px}.firm-table-v2 input[type=checkbox]{width:16px;height:16px}
  .firm-tag.ready{background:#12361f;color:#86efac}.firm-tag.already_exists{background:#0f2f3d;color:#7dd3fc}.firm-tag.conflict{background:#3b1620;color:#fca5a5}
  .firm-file-main{display:block;color:#fff;font-weight:750}.firm-file-sub{display:block;color:var(--muted);font-size:11px;margin-top:3px;overflow-wrap:anywhere}.firm-last-run{color:var(--muted);font-size:12px;margin-top:8px}
  .firm-nav-badge{display:inline-grid;place-items:center;min-width:20px;height:20px;margin-left:6px;padding:0 6px;border-radius:999px;background:#dc3545;color:#fff;font-size:10px;font-weight:900}.firm-nav-badge.pulse{animation:firmPulse 2s infinite}@keyframes firmPulse{50%{box-shadow:0 0 0 6px #dc35452e}}
  .firm-nav-badge.hidden{display:none}.firm-flow-tabs button:focus{outline:none}.firm-flow-tabs button:focus-visible{box-shadow:0 0 0 3px rgba(20,92,168,.18)}
  @media(max-width:1100px){.firm-flow-tabs{grid-template-columns:1fr}.firm-flow-heading{align-items:flex-start;flex-direction:column}.firm-stats-v2{grid-template-columns:repeat(2,minmax(0,1fr))}.firm-route-card{grid-template-columns:1fr}.firm-head-actions{justify-content:flex-start}}
  #v-agenda{width:100%;max-width:calc(100vw - 326px);min-width:0;overflow:hidden}.agenda-head-row{width:100%;max-width:100%;display:flex;flex-wrap:wrap;justify-content:space-between;align-items:flex-start;gap:18px;margin-bottom:14px}.agenda-head-row>div:first-child{flex:1 1 430px;min-width:0}.agenda-actions{display:flex;flex:1 1 420px;gap:9px;flex-wrap:wrap;justify-content:flex-end}.agenda-last-run{display:block;color:#718395;font-size:12px;margin-top:5px}
  .agenda-notice{display:flex;align-items:center;justify-content:space-between;gap:16px;padding:13px 15px;margin-bottom:14px;border:1px solid #bfe6d5;border-radius:12px;background:#e8f7f0;color:#167553}.agenda-notice.busy{border-color:#c8d8e7;background:#eef6fc;color:#145ca8}.agenda-notice.error{border-color:#efc2c7;background:#fff0f2;color:#a63e48}.agenda-notice strong,.agenda-notice span{display:block}.agenda-notice span{margin-top:2px;font-size:12px}.agenda-live{display:inline-flex;align-items:center;gap:7px;font-size:11px;font-weight:850;white-space:nowrap}.agenda-live-dot{width:8px;height:8px;border-radius:50%;background:currentColor;box-shadow:0 0 0 5px currentColor;opacity:.55}
  .agenda-stats{grid-template-columns:repeat(auto-fit,minmax(145px,1fr))}.agenda-today-grid{display:grid;grid-template-columns:minmax(0,1.3fr) minmax(280px,.7fr);gap:14px;margin-bottom:14px}.agenda-today-card{display:flex;justify-content:space-between;align-items:center;gap:20px;border-left:5px solid #145ca8}.agenda-today-card .kicker,.agenda-total-card .kicker{color:#607487;font-size:11px;font-weight:850;letter-spacing:.7px;text-transform:uppercase}.agenda-today-card h4{margin:5px 0;color:#173a5e;font-size:20px}.agenda-today-card p{margin:0;color:#607487;font-size:12px}.agenda-today-date{min-width:126px;padding:14px;border-radius:11px;background:#eef6fc;color:#145ca8;text-align:center;font-weight:900}.agenda-today-date strong,.agenda-today-date span{display:block}.agenda-today-date strong{font-size:28px;line-height:1}.agenda-today-date span{margin-top:5px;font-size:11px;text-transform:uppercase}
  .agenda-total-card h4{margin:5px 0 10px;color:#173a5e;font-size:16px}.agenda-total-list{display:flex;gap:8px;flex-wrap:wrap}.agenda-total-pill{display:inline-flex;gap:7px;align-items:center;padding:8px 10px;border-radius:9px;background:#eef4f8;color:#173a5e;font-size:12px}.agenda-total-pill strong{color:#145ca8}.agenda-total-empty{color:#718395;font-size:12px}
  .agenda-panel-head{display:flex;flex-wrap:wrap;align-items:end;justify-content:space-between;gap:16px;margin-bottom:14px}.agenda-panel-head>div:first-child{flex:1 1 260px}.agenda-panel-head h4{margin:0;color:#173a5e;font-size:17px}.agenda-panel-head p{margin:4px 0 0;color:#718395;font-size:12px}.agenda-filters{display:flex;flex:1 1 500px;gap:9px;flex-wrap:wrap;justify-content:flex-end}.agenda-filters label{display:grid;flex:1 1 145px;gap:4px;color:#607487;font-size:10px;font-weight:850;text-transform:uppercase}.agenda-filters select,.agenda-filters input{width:100%;min-width:0;padding:8px 10px;font-size:12px}.agenda-filters label:last-child{flex-basis:220px}
  .agenda-table{min-width:1450px}.agenda-table td{vertical-align:top}.agenda-loan{display:block;color:#173a5e;font-weight:850}.agenda-cell-sub{display:block;margin-top:3px;color:#718395;font-size:11px;overflow-wrap:anywhere}.agenda-date-main{display:block;color:#314a60;font-weight:800}.agenda-amount{color:#173a5e;text-align:right;font-variant-numeric:tabular-nums;font-weight:850}.agenda-currency{display:inline-flex;padding:4px 8px;border-radius:7px;background:#eef4f8;color:#526a7f;font-weight:850}.agenda-correspondent{display:inline-flex;padding:5px 9px;border-radius:999px;background:#eaf4ff;color:#145ca8;font-size:10px;font-weight:900}.agenda-correspondent.review{background:#fff5dd;color:#9a6211}.agenda-status{display:inline-flex;padding:5px 9px;border-radius:999px;font-size:10px;font-weight:900;white-space:nowrap}.agenda-status.today{background:#eaf4ff;color:#145ca8}.agenda-status.scheduled{background:#edf2f6;color:#526a7f}.agenda-status.in_process{background:#fff5dd;color:#9a6211}.agenda-status.overdue{background:#ffedf0;color:#a63e48}.agenda-status.paid{background:#e8f7f0;color:#167553}.agenda-status.review{background:#fff5dd;color:#9a6211}.agenda-error{padding:11px 13px;margin-bottom:12px;border:1px solid #efc2c7;border-radius:10px;background:#fff0f2;color:#a63e48;font-size:12px}.agenda-source-line{display:flex;align-items:center;gap:10px;margin-top:11px;color:#607487;font-size:11px}.agenda-source-line code{color:#415b72;overflow-wrap:anywhere}.agenda-nav-badge{display:inline-grid;place-items:center;min-width:20px;height:20px;margin-left:6px;padding:0 6px;border-radius:999px;background:#dc3545;color:#fff;font-size:10px;font-weight:900}.agenda-nav-badge.hidden{display:none}.agenda-nav-badge.pulse{animation:firmPulse 2s infinite}
  @media(max-width:1100px){.agenda-head-row,.agenda-panel-head{align-items:flex-start;flex-direction:column}.agenda-actions,.agenda-filters{justify-content:flex-start}.agenda-stats{grid-template-columns:repeat(2,minmax(0,1fr))}.agenda-today-grid{grid-template-columns:1fr}.agenda-today-card{align-items:flex-start;flex-direction:column}}

  /* Activaciones de cuentas y aprendizaje local de respuestas. */
  .activation-head,.mail-head{display:flex;align-items:flex-start;justify-content:space-between;gap:18px;flex-wrap:wrap}.activation-head>div,.mail-head>div:first-child{flex:1 1 420px}.activation-source{display:grid;gap:7px;padding:10px 13px;border:1px solid #cbd8e2;border-radius:10px;background:#fff;color:#607487;font-size:10.5px;max-width:390px}.activation-source strong{color:#173a5e}.activation-source code{overflow-wrap:anywhere}.activation-swiftref{display:inline-flex;justify-content:center;align-items:center;margin-top:3px;text-decoration:none}.activation-frame-card{padding:0;overflow:hidden}.activation-frame{display:block;width:100%;height:760px;border:0;background:#f4f7f9}.activation-error{padding:14px;border:1px solid #efc2c7;background:#fff0f2;color:#a63e48}
  .mail-flow-tabs{margin-top:2px}.mail-panel{display:none}.mail-panel.act{display:block}.mail-grid{display:grid;grid-template-columns:minmax(300px,.9fr) minmax(360px,1.1fr);gap:16px}.mail-field{display:grid;gap:6px;margin-bottom:13px}.mail-field label{color:#607487;font-size:10.5px;font-weight:850;text-transform:uppercase}.mail-field input,.mail-field select{width:100%}.mail-textarea{height:150px;min-height:110px;resize:vertical;font-family:"Libre Franklin","Segoe UI",sans-serif;line-height:1.5}.mail-response{height:225px}.mail-actions{display:flex;align-items:center;gap:10px;flex-wrap:wrap}.mail-status{padding:12px 14px;margin-bottom:14px;border:1px solid #cbd8e2;border-radius:10px;background:#f7f9fb;color:#526a7f;font-size:12px}.mail-status.ok{border-color:#bfe6d5;background:#e8f7f0;color:#167553}.mail-status.warn{border-color:#eed7a2;background:#fff5dd;color:#765313}.mail-confidence{display:inline-flex;padding:5px 9px;border-radius:999px;background:#eaf4ff;color:#145ca8;font-size:10.5px;font-weight:900}.mail-match-list{display:grid;gap:9px}.mail-match{padding:11px 12px;border:1px solid #d8e1e9;border-radius:10px;background:#f8fafc}.mail-match strong,.mail-match small{display:block}.mail-match strong{color:#173a5e;font-size:12px}.mail-match small{margin-top:4px;color:#718395}.mail-table{min-width:920px}.mail-table td{vertical-align:top}.mail-table .mail-preview{display:block;max-width:310px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.mail-privacy{display:flex;align-items:center;gap:8px;color:#167553;font-size:11px;font-weight:800}
  @media(max-width:980px){.mail-grid{grid-template-columns:1fr}.activation-frame{height:850px}}

  .muted{color:var(--muted)}

  /* Tema claro global del Sistema de Conciliación. */
  :root{--bg:#eef3f7;--bg2:#f7f9fb;--panel:#fff;--panel2:#f7f9fb;--line:#d8e1e9;--txt:#21364a;--muted:#6e8091;--gold:#c49a45;--gold2:#8a611c;--cyan:#1871c9;--ok:#167553;--okbg:#e8f7f0;--bad:#a63e48;--badbg:#ffedf0;--blue:#145ca8}
  body{color:var(--txt);background:#eef3f7}
  .side{background:#fff;border-right-color:#d8e1e9;box-shadow:5px 0 22px rgba(17,42,68,.05)}
  .brand{border-bottom-color:#e1e7ed}.brand h2,.brand small,.brand .ln{display:none}.brand .ln{background:linear-gradient(90deg,transparent,#c49a45,transparent)}
  .nav .it{color:#526a7f}.nav .it:hover{background:#f3f7fa}.nav .it.act{background:#eaf4ff;border-color:#cfe0ee;color:#173a5e;box-shadow:inset 3px 0 0 #145ca8}.nav .it .num{background:#edf2f6;color:#47637c;border-color:#d8e1e9}.nav .it.act .num{background:#145ca8;color:#fff}.nav .it .ss{color:#8292a1}.nav-group-head{background:#f7f9fb;border-color:#d8e1e9;color:#173a5e}.nav-group-head .group-icon{background:#eaf4ff;color:#145ca8}.nav-group-head small{color:#8292a1}.nav-group-items{border-left-color:#d8e1e9}
  .conc-flow-tabs button{background:#fff;border-color:#d8e1e9;color:#526a7f}.conc-flow-tabs button:hover{border-color:#4b83b5}.conc-flow-tabs button.act{background:#eef6fc;border-color:#145ca8;color:#173a5e;box-shadow:0 0 0 2px rgba(20,92,168,.08)}.conc-tab-icon{background:#edf2f6;color:#47637c}.conc-flow-tabs button.act .conc-tab-icon{background:#145ca8;color:#fff}.conc-flow-tabs button small{color:#8292a1}.conc-portfolio-card{background:#fff;border-color:#d8e1e9}.conc-portfolio-label strong{color:#173a5e}.conc-portfolio-list .e{background:#f7f9fb;border-color:#d8e1e9;color:#526a7f}
  .perbox{border-color:#b9cce0;color:#145ca8;background:#eef6fd}.estado{border-top-color:#e1e7ed}.estado .h{color:#47637c}.estado .e{color:#526a7f}
  .main{background:#eef3f7}.top{background:#fff;border-bottom-color:#d8e1e9;box-shadow:0 4px 16px rgba(17,42,68,.04)}.top h1{color:#173a5e}.top h1 b{color:#8a611c}.uchip{color:#607487}
  .badge{border-color:#d8e1e9;color:#607487;background:#f7f9fb}.badge.set{border-color:#d7bc83;color:#765313;background:#fff8e8}.badge.ok{border-color:#bfe6d5;color:#167553;background:#e8f7f0}
  .content{background:#eef3f7}h3.sec{color:#173a5e}.sub{color:#607487}
  .card{background:#fff;border-color:#d8e1e9;box-shadow:0 10px 28px rgba(17,42,68,.07);backdrop-filter:none}
  .stat{background:#fff;border-color:#d8e1e9;box-shadow:0 6px 18px rgba(17,42,68,.05)}.stat .n{color:#173a5e}.stat .l{color:#718395}
  .drop{border-color:#9eb4c7;background:#fbfcfd}.drop:hover{border-color:#145ca8;background:#f2f8fd}.drop .ti{color:#173a5e}.drop .de{color:#718395}.chipf{background:#eef4f8;border-color:#cdd9e3;color:#314a60}
  .btn.gh,.btn.alt{background:#fff;color:#173a5e;border-color:#bfcdda;box-shadow:0 4px 12px rgba(17,42,68,.06)}.btn.alt{color:#8a611c;border-color:#d7bc83}.btn:not(.gh):not(.alt){box-shadow:0 7px 18px rgba(138,97,28,.16)}
  .info{border-left-color:#c49a45;background:#fff8e8;color:#5d5139}.fld label{color:#607487}
  input[type=text],input[type=number],select,textarea{background:#fbfcfd;border-color:#c6d2dc;color:#173a5e}input[type=text]:focus,input[type=number]:focus,select:focus,textarea:focus{border-color:#1871c9;outline:3px solid rgba(24,113,201,.11)}input[type=checkbox]{accent-color:#145ca8}
  .scrollx{background:#fff;border-color:#d8e1e9}.mtable th,.mtable td{border-bottom-color:#e6ecf1;color:#4b6073}.mtable thead th,.mtable thead th.acr,.mtable thead th.est{background:#f7f9fb;color:#607487}.mtable thead tr.sub th{background:#f1f5f8;color:#6e8091}.mtable td.acr{background:#fff;color:#173a5e}.mtable td.num{color:#314a60}.mtable tbody tr:hover td,.mtable tbody tr:hover td.acr{background:#f5f9fc}.mtable tr.tot td,.mtable tr.tot td.acr{background:#eef4f8;color:#765313}.mtable td.sep{border-left-color:#e2e8ee}
  .diaghead{background:#f7f9fb;border-color:#d8e1e9;color:#173a5e}.diaghead small{color:#718395}
  .dcard,.dobs,.dbox,.loan{background:#f8fafc;border-color:#d8e1e9}.dcard .dname,.loan .lh{color:#173a5e}.dobs,.dbox .bd{color:#4b6073}.dtag.ot{background:#e8eef4;color:#526a7f}.dtag.dc{background:#eaf4ff;color:#145ca8}.dtag.pd{background:#fff5dd;color:#9a6211}
  .note-ok{background:#e8f7f0;border-color:#bfe6d5;color:#167553}.note-warn{background:#fff5dd;border-color:#eed7a2;color:#765313}
  .source-pill{background:#eef4f8;border-color:#d8e1e9;color:#526a7f}.caf-check,.firm-check,.firm-paths,.comp-check{color:#526a7f}.caf-table td code,.firm-table td code,.firm-table-v2 td code,.firm-route-card code{color:#415b72}
  .caf-log,.firm-log,.comp-log{background:#f7f9fb;border-color:#d8e1e9;color:#41576b}
  .firm-slicer button{background:#fff;border-color:#cbd6df;color:#526a7f}.firm-slicer button.act{background:#145ca8;border-color:#145ca8;color:#fff}
  .firm-alert{background:#f7f9fb;border-color:#d8e1e9;color:#526a7f}.firm-alert.on{border-color:#eed7a2;background:#fff5dd;color:#765313}.firm-alert.ready{border-color:#bfe6d5;background:#e8f7f0;color:#167553}
  .firm-signature-panel{border-color:#eed7a2;background:#fff8e8}.firm-signature-head strong,.firm-signature-head span{color:#765313}.firm-signature-item{border-color:#e4c98f;background:#fff}.firm-signature-item strong{color:#173a5e}.firm-file-main{color:#173a5e}.firm-file-sub,.firm-last-run{color:#718395}
  .firm-tag{background:#edf2f6;color:#526a7f}.firm-tag.move,.firm-tag.ready{background:#e8f7f0;color:#167553}.firm-tag.clean{background:#fff5dd;color:#9a6211}.firm-tag.done,.firm-tag.already_exists{background:#eaf4ff;color:#145ca8}.firm-tag.warn,.firm-tag.conflict{background:#ffedf0;color:#a63e48}
  .firm-flow-tabs button{background:#fff;border-color:#d8e1e9;color:#314a60}.firm-flow-tabs button:hover{border-color:#4b83b5}.firm-flow-tabs button.act{background:#eef6fc;border-color:#145ca8;box-shadow:0 0 0 2px rgba(20,92,168,.08)}.firm-flow-tabs button small{background:#edf2f6;color:#607487}.firm-flow-tabs button.act small{background:#dcecf8;color:#145ca8}.firm-flow-heading strong{color:#173a5e}.firm-monitor-badge{border-color:#bfe6d5;background:#e8f7f0;color:#167553}
  #v-comprobantes{background:#eef3f7}
  #v-matrizprestamos{min-height:calc(100vh - 76px);margin:-28px -30px;padding:28px 30px;background:#eef3f7;color:#24364a}
  .matrix-head{display:flex;gap:18px;align-items:flex-start;justify-content:space-between;margin-bottom:18px}.matrix-head .sub{margin-bottom:4px}.matrix-path{max-width:690px;color:#607487;font-size:11px;overflow-wrap:anywhere}.matrix-grid{display:grid;grid-template-columns:minmax(300px,.78fr) minmax(520px,1.22fr);gap:18px;align-items:start}.matrix-card{background:#fff;border:1px solid #d8e1e9;border-radius:16px;padding:20px;box-shadow:0 9px 28px rgba(17,42,68,.07)}
  .matrix-drop{min-height:220px;border:2px dashed #b9cad8;border-radius:14px;background:#f7fafc;display:grid;place-items:center;text-align:center;padding:22px;cursor:pointer;transition:.2s}.matrix-drop:hover,.matrix-drop.drag{border-color:#145ca8;background:#eef7ff}.matrix-drop strong{display:block;color:#173a5e;font-size:16px;margin:8px}.matrix-drop span{display:block;color:#718395;font-size:12px;line-height:1.5}.matrix-drop input{display:none}.matrix-preview{display:none;width:100%;max-height:310px;object-fit:contain;margin-top:14px;border:1px solid #d8e1e9;border-radius:10px;background:#f4f7f9}.matrix-preview.on{display:block}
  .matrix-notice{display:flex;align-items:center;justify-content:space-between;gap:12px;margin:14px 0;padding:12px 14px;border:1px solid #cbdce8;border-radius:11px;background:#f4f9fd;color:#45627a;font-size:12px}.matrix-notice.ok{border-color:#bfe6d5;background:#e8f7f0;color:#167553}.matrix-notice.warn{border-color:#eed7a2;background:#fff5dd;color:#765313}.matrix-notice.bad{border-color:#efc2c7;background:#fff0f2;color:#a63e48}.matrix-score{font-size:16px;font-weight:900;white-space:nowrap}
  .matrix-form{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}.matrix-form .wide{grid-column:span 2}.matrix-form .full{grid-column:1/-1}.matrix-form label{display:block;margin:0 0 6px;color:#607487;font-size:10px;font-weight:850;text-transform:uppercase;letter-spacing:.35px}.matrix-form input,.matrix-form select,.matrix-form textarea{width:100%;padding:10px 11px;border:1px solid #cbd6df;border-radius:9px;background:#fbfcfd;color:#173a5e;font-size:12px}.matrix-form textarea{min-height:66px;resize:vertical}.matrix-form input[readonly]{background:#eef3f7;color:#607487}.matrix-required:after{content:" *";color:#a63e48}.matrix-actions{display:flex;gap:10px;justify-content:flex-end;flex-wrap:wrap;margin-top:15px}.matrix-actions .btn:disabled{opacity:.5;cursor:not-allowed}.matrix-warnings{margin:10px 0 0;padding-left:18px;color:#765313;font-size:11px;line-height:1.5}.matrix-history{margin-top:18px}.matrix-history table{width:100%;border-collapse:collapse;font-size:11px}.matrix-history th,.matrix-history td{padding:10px;border-bottom:1px solid #e6ecf1;text-align:left;color:#4b6073}.matrix-history th{background:#f7f9fb;color:#607487;font-size:9px;text-transform:uppercase}.matrix-empty{text-align:center;padding:24px;color:#789}.matrix-live{display:inline-flex;align-items:center;gap:7px;color:#167553;font-size:11px;font-weight:750}.matrix-live:before{content:"";width:8px;height:8px;border-radius:50%;background:#19a36d;box-shadow:0 0 0 4px rgba(25,163,109,.1)}
  @media(max-width:1100px){.matrix-grid{grid-template-columns:1fr}.matrix-form{grid-template-columns:repeat(2,minmax(0,1fr))}}@media(max-width:700px){.matrix-head{display:block}.matrix-form{grid-template-columns:1fr}.matrix-form .wide{grid-column:auto}}
  .control-head{display:flex;align-items:flex-start;justify-content:space-between;gap:18px;flex-wrap:wrap;margin-bottom:18px}.control-head>div{flex:1 1 480px}.control-updated{display:block;margin-top:5px;color:#718395;font-size:11px}
  .control-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(190px,1fr));gap:13px;margin-bottom:18px}.control-card{min-height:128px;padding:17px;border:1px solid #d8e1e9;border-radius:14px;background:#fff;box-shadow:0 6px 18px rgba(17,42,68,.05);text-align:left;cursor:pointer;transition:.16s}.control-card:hover{transform:translateY(-2px);border-color:#8fb5d6;box-shadow:0 10px 24px rgba(17,42,68,.09)}.control-card .control-value{display:block;color:#173a5e;font:800 29px Archivo,"Segoe UI",sans-serif}.control-card .control-label{display:block;margin-top:8px;color:#526a7f;font-size:12px;font-weight:750}.control-card .control-action{display:block;margin-top:9px;color:#145ca8;font-size:10.5px;font-weight:850}.control-card.warning{border-left:5px solid #d49b2f}.control-card.urgent{border-left:5px solid #145ca8}.control-card.danger{border-left:5px solid #b43e4b}.control-card.ok{border-left:5px solid #2b8a63}
  .control-layout{display:grid;grid-template-columns:minmax(0,1.35fr) minmax(280px,.65fr);gap:16px}.control-issues{display:grid;gap:9px}.control-issue{display:grid;grid-template-columns:auto minmax(0,1fr) auto;gap:10px;align-items:center;padding:12px;border:1px solid #dde5ec;border-radius:10px;background:#f8fafc}.control-issue-dot{width:9px;height:9px;border-radius:50%;background:#2b8a63}.control-issue.warning .control-issue-dot{background:#d49b2f}.control-issue.urgent .control-issue-dot{background:#145ca8}.control-issue.danger .control-issue-dot{background:#b43e4b}.control-issue strong,.control-issue small{display:block}.control-issue strong{color:#173a5e;font-size:12px}.control-issue small{margin-top:3px;color:#718395}.control-issue button{border:0;background:transparent;color:#145ca8;font-size:11px;font-weight:850;cursor:pointer}.control-empty{padding:24px;text-align:center;color:#167553;border:1px solid #bfe6d5;border-radius:11px;background:#e8f7f0}.control-health{display:grid;gap:10px}.control-health-row{display:flex;justify-content:space-between;gap:12px;padding:10px 0;border-bottom:1px solid #e6ecf1;color:#526a7f;font-size:11px}.control-health-row strong{color:#173a5e}.control-health-row time{color:#718395;text-align:right}.control-health-badge{display:inline-flex;padding:6px 10px;border-radius:999px;background:#e8f7f0;color:#167553;font-size:10px;font-weight:900}.control-health-badge.bad{background:#ffedf0;color:#a63e48}@media(max-width:980px){.control-layout{grid-template-columns:1fr}}
</style></head>
<body data-user-role="{{ user_role }}">
  <svg class="icon-sprite" aria-hidden="true" xmlns="http://www.w3.org/2000/svg">
    <symbol id="ico-conciliacion" viewBox="0 0 24 24"><path d="M4 7h16M7 4 4 7l3 3M20 17H4m13-3 3 3-3 3"/></symbol>
    <symbol id="ico-recepcion" viewBox="0 0 24 24"><path d="M12 3v12m-4-4 4 4 4-4M4 17v3h16v-3"/></symbol>
    <symbol id="ico-carteras" viewBox="0 0 24 24"><path d="M4 6h11m-3-3 3 3-3 3M20 18H9m3-3-3 3 3 3"/></symbol>
    <symbol id="ico-observaciones" viewBox="0 0 24 24"><path d="M8 4h8l2 2v14H6V6l2-2Z"/><path d="M9 4v3h6V4m-6 8h6m-6 4h4"/></symbol>
    <symbol id="ico-oficio" viewBox="0 0 24 24"><path d="M6 3h9l3 3v15H6Z"/><path d="M15 3v4h4M9 12h6m-6 4h4"/></symbol>
    <symbol id="ico-caf" viewBox="0 0 24 24"><ellipse cx="9" cy="6" rx="5" ry="2.5"/><path d="M4 6v4c0 1.4 2.2 2.5 5 2.5S14 11.4 14 10V6M4 10v4c0 1.4 2.2 2.5 5 2.5 1 0 1.9-.1 2.6-.4"/><path d="M14 13h6v7h-6z"/></symbol>
    <symbol id="ico-agenda" viewBox="0 0 24 24"><rect x="3" y="5" width="18" height="16" rx="2"/><path d="M7 3v4m10-4v4M3 10h18m-14 4h3m4 0h3m-10 3h3"/></symbol>
    <symbol id="ico-firma" viewBox="0 0 24 24"><path d="m4 17 3.5-.8L18 5.7 15.3 3 4.8 13.5 4 17Z"/><path d="m13.5 4.8 2.7 2.7M4 21h16"/></symbol>
    <symbol id="ico-comprobante" viewBox="0 0 24 24"><path d="M5 3h14v18l-2-1.5L15 21l-2-1.5L11 21l-2-1.5L5 21Z"/><path d="M8 7h8m-8 4h8m-8 4h5"/></symbol>
    <symbol id="ico-cuentas" viewBox="0 0 24 24"><rect x="3" y="5" width="18" height="14" rx="2"/><path d="M3 9h18M7 14h3m6-2v5m-2.5-2.5h5"/></symbol>
    <symbol id="ico-correo" viewBox="0 0 24 24"><rect x="3" y="5" width="18" height="14" rx="2"/><path d="m4 7 8 6 8-6M8 17h4m4-2 1 1 2-3"/></symbol>
    <symbol id="ico-archivo-quipux" viewBox="0 0 24 24"><path d="M4 4h6l2 2h8v14H4Z"/><path d="M8 11h8m-8 4h5M16 3v6m-3-3h6"/></symbol>
    <symbol id="ico-contrato-agencia" viewBox="0 0 24 24"><path d="M6 3h9l3 3v15H6Z"/><path d="M15 3v4h4M9 11h6m-6 4h3"/><circle cx="16.5" cy="16.5" r="2.5"/><path d="m15 19-1 2 2.5-.8L19 21l-1-2"/></symbol>
    <symbol id="ico-matriz" viewBox="0 0 24 24"><rect x="3" y="4" width="18" height="16" rx="2"/><path d="M3 9h18M9 9v11m6-11v11M6 2v4m12-4v4"/></symbol>
  </svg>
  <aside class="side">
    <div class="brand"><span class="logo-slot"><img class="bce-logo" src="/logo" alt="Banco Central del Ecuador" onerror="this.style.display='none'" onload="var sib=this.nextElementSibling; if(sib) sib.style.display='none';"><svg class="bce-sun" viewBox="0 0 140 140" xmlns="http://www.w3.org/2000/svg" aria-label="Banco Central del Ecuador"><g fill="#E6B422"><polygon points="130.0,70.0 91.8,67.1 91.8,72.9"/><polygon points="116.2,89.1 91.1,76.3 89.4,80.4"/><polygon points="112.4,112.4 87.4,83.4 83.4,87.4"/><polygon points="89.1,116.2 80.4,89.4 76.3,91.1"/><polygon points="70.0,130.0 72.9,91.8 67.1,91.8"/><polygon points="50.9,116.2 63.7,91.1 59.6,89.4"/><polygon points="27.6,112.4 56.6,87.4 52.6,83.4"/><polygon points="23.8,89.1 50.6,80.4 48.9,76.3"/><polygon points="10.0,70.0 48.2,72.9 48.2,67.1"/><polygon points="23.8,50.9 48.9,63.7 50.6,59.6"/><polygon points="27.6,27.6 52.6,56.6 56.6,52.6"/><polygon points="50.9,23.8 59.6,50.6 63.7,48.9"/><polygon points="70.0,10.0 67.1,48.2 72.9,48.2"/><polygon points="89.1,23.8 76.3,48.9 80.4,50.6"/><polygon points="112.4,27.6 83.4,52.6 87.4,56.6"/><polygon points="116.2,50.9 89.4,59.6 91.1,63.7"/><circle cx="70" cy="70" r="22"/></g><g fill="#1a2440"><circle cx="62" cy="66" r="3.4"/><circle cx="78" cy="66" r="3.4"/><path d="M58 76 Q70 86 82 76" stroke="#1a2440" stroke-width="3.2" fill="none" stroke-linecap="round"/><path d="M64 58 Q70 54 76 58" stroke="#1a2440" stroke-width="2.6" fill="none" stroke-linecap="round"/></g></svg></span><div class="ln"></div><h2 style="font-size:12px;letter-spacing:1px;color:var(--gold2)">CONCILIACION DE DEUDA EXTERNA</h2><small>Servicios Financieros Internacionales</small></div>
    <nav class="nav" id="nav">
      <div class="it act" data-v="control"><div class="num nav-symbol"><svg><use href="#ico-agenda"></use></svg></div><div><div class="tt">Centro de Control</div><div class="ss">Pendientes y alertas operativas</div></div></div>
      <div class="it" data-v="cargar" data-conc-root="true"><div class="num nav-symbol"><svg><use href="#ico-conciliacion"></use></svg></div><div><div class="tt">Conciliación Deuda Externa</div><div class="ss">BCE &middot; MEF &middot; CAF</div></div></div>
      <div class="it" data-v="agenda"><div class="num nav-symbol"><svg><use href="#ico-agenda"></use></svg></div><div><div class="tt">Agenda Diaria <span id="agendaNavBadge" class="agenda-nav-badge hidden">0</span></div><div class="ss">Pagos previstos del día</div></div></div>
      <div class="it" data-v="archivoquipux"><div class="num nav-symbol"><svg><use href="#ico-archivo-quipux"></use></svg></div><div><div class="tt">Archivo Quipux <span id="quipuxNavBadge" class="quipux-nav-badge hidden">0</span></div><div class="ss">Oficio · Estado · Formulario</div></div></div>
      <div class="it" data-v="contratosagencia"><div class="num nav-symbol"><svg><use href="#ico-contrato-agencia"></use></svg></div><div><div class="tt">Contratos Agencia Fiscal <span id="contractNavBadge" class="contract-nav-badge hidden">0</span></div><div class="ss">Oficios firmados · Archivo automático</div></div></div>
      <div class="it" data-v="matrizprestamos"><div class="num nav-symbol"><svg><use href="#ico-matriz"></use></svg></div><div><div class="tt">Matriz de Préstamos</div><div class="ss">Captura SGI · Excel controlado</div></div></div>
      <div class="it" data-v="activaciones"><div class="num nav-symbol"><svg><use href="#ico-cuentas"></use></svg></div><div><div class="tt">Activaciones Cuentas</div><div class="ss">Validador Swift &middot; ISO 20022</div></div></div>
      <div class="it" data-v="respuestas"><div class="num nav-symbol"><svg><use href="#ico-correo"></use></svg></div><div><div class="tt">Análisis de Correos</div><div class="ss">Respuestas asistidas y aprendizaje</div></div></div>
      <div class="it" data-v="firmados"><div class="num nav-symbol"><svg><use href="#ico-firma"></use></svg></div><div><div class="tt">Firma EC <span id="firmNavBadge" class="firm-nav-badge hidden">0</span></div><div class="ss">Deuda Externa · Automáticos · Bancos</div></div></div>
      <div class="it" data-v="comprobantes"><div class="num nav-symbol"><svg><use href="#ico-comprobante"></use></svg></div><div><div class="tt">Comprobantes Contables</div><div class="ss">BRYAN &middot; STEVEN &middot; ACK</div></div></div>
    </nav>
    <div class="perbox" id="perBox"></div>
  </aside>
  <div class="main">
    <div class="top"><div class="topbrand"><span class="logo-slot"><img class="bce-logo" src="/logo" alt="BCE" onerror="this.style.display='none'" onload="var sib=this.nextElementSibling; if(sib) sib.style.display='none';"><svg class="bce-sun" viewBox="0 0 140 140" xmlns="http://www.w3.org/2000/svg" aria-label="Banco Central del Ecuador"><g fill="#E6B422"><polygon points="130.0,70.0 91.8,67.1 91.8,72.9"/><polygon points="116.2,89.1 91.1,76.3 89.4,80.4"/><polygon points="112.4,112.4 87.4,83.4 83.4,87.4"/><polygon points="89.1,116.2 80.4,89.4 76.3,91.1"/><polygon points="70.0,130.0 72.9,91.8 67.1,91.8"/><polygon points="50.9,116.2 63.7,91.1 59.6,89.4"/><polygon points="27.6,112.4 56.6,87.4 52.6,83.4"/><polygon points="23.8,89.1 50.6,80.4 48.9,76.3"/><polygon points="10.0,70.0 48.2,72.9 48.2,67.1"/><polygon points="23.8,50.9 48.9,63.7 50.6,59.6"/><polygon points="27.6,27.6 52.6,56.6 56.6,52.6"/><polygon points="50.9,23.8 59.6,50.6 63.7,48.9"/><polygon points="70.0,10.0 67.1,48.2 72.9,48.2"/><polygon points="89.1,23.8 76.3,48.9 80.4,50.6"/><polygon points="112.4,27.6 83.4,52.6 87.4,56.6"/><polygon points="116.2,50.9 89.4,59.6 91.1,63.7"/><circle cx="70" cy="70" r="22"/></g><g fill="#1a2440"><circle cx="62" cy="66" r="3.4"/><circle cx="78" cy="66" r="3.4"/><path d="M58 76 Q70 86 82 76" stroke="#1a2440" stroke-width="3.2" fill="none" stroke-linecap="round"/><path d="M64 58 Q70 54 76 58" stroke="#1a2440" stroke-width="2.6" fill="none" stroke-linecap="round"/></g></svg></span><h1><b>GIDEP BCE</b></h1></div>
      <div class="right"><span class="badge" id="periodoBadge">Periodo no detectado</span>
        <span class="uchip"><span class="uav">{{ user_initial }}</span><span class="user-meta"><strong>{{ user_name }}</strong><small>{{ user_role }}</small></span><a class="logout-link" href="/logout">Cerrar sesión</a></span></div></div>
    <div class="content">

      <section id="v-control">
        <div class="control-head">
          <div><h3 class="sec"><span class="bar"></span>Centro de Control Diario</h3><p class="sub">Una sola bandeja para revisar pagos, documentos, firmas, comprobantes y ACK que requieren atención.</p><span class="control-updated" id="controlUpdated">Consultando los módulos...</span></div>
          <button class="btn" id="controlRefresh" type="button">Actualizar control</button>
        </div>
        <div class="control-grid" id="controlCards"><div class="control-card"><span class="control-value">—</span><span class="control-label">Cargando estado operativo</span></div></div>
        <div class="control-layout">
          <div class="card"><div class="diaghead">Bandeja de atención<small>Seleccione un elemento para abrir directamente el módulo responsable.</small></div><div class="control-issues" id="controlIssues"><div class="control-empty">Revisando pendientes...</div></div></div>
          <div class="card"><div class="diaghead">Salud de los monitores<small>Última información disponible por módulo.</small></div><div id="controlHealth" class="control-health"></div></div>
        </div>
      </section>

      <div class="conc-module-head hidden" id="concModuleHead">
        <div class="conc-module-title">
          <div><h3 class="sec"><span class="bar"></span>Conciliación Deuda Externa</h3><p>Flujo integral de recepción, conciliación, observaciones, emisión del oficio y condonados CAF.</p></div>
          <span class="badge">Proceso mensual</span>
        </div>
        <div class="conc-flow-tabs" id="concFlowTabs" role="tablist" aria-label="Procesos de Conciliación de Deuda Externa">
          <button class="act" type="button" data-conc-view="cargar" role="tab" aria-selected="true"><span class="conc-tab-icon"><svg><use href="#ico-recepcion"></use></svg></span><span><strong>Recepción de Reportes</strong><small>MEF · BCE</small></span></button>
          <button type="button" data-conc-view="resultados" role="tab" aria-selected="false"><span class="conc-tab-icon"><svg><use href="#ico-carteras"></use></svg></span><span><strong>Conciliación de Carteras</strong><small>Cruce y resultados</small></span></button>
          <button type="button" data-conc-view="ajustes" role="tab" aria-selected="false"><span class="conc-tab-icon"><svg><use href="#ico-observaciones"></use></svg></span><span><strong>Gestión de Observaciones</strong><small>Ajustes y pagos directos</small></span></button>
          <button type="button" data-conc-view="quipux" role="tab" aria-selected="false"><span class="conc-tab-icon"><svg><use href="#ico-oficio"></use></svg></span><span><strong>Emisión del Oficio</strong><small>Quipux de respuesta</small></span></button>
          <button type="button" data-conc-view="condonados" role="tab" aria-selected="false"><span class="conc-tab-icon"><svg><use href="#ico-caf"></use></svg></span><span><strong>Condonados CAF</strong><small>Preparación BCE</small></span></button>
        </div>
        <div class="card conc-portfolio-card">
          <div class="conc-portfolio-label"><strong>Estado de carteras</strong><span id="concPortfolioPeriod">Periodo aún no detectado</span></div>
          <div id="estadoSide" class="conc-portfolio-list">Sin datos</div>
        </div>
      </div>

      <section id="v-cargar">
        <h3 class="sec"><span class="bar"></span>Recepcion de Reportes</h3>
        <p class="sub">Cargue los reportes de conciliacion del MEF y del BCE en una sola zona. El periodo se reconoce automaticamente y los calculos se realizan sin depender del orden de columnas o filas.</p>
        <div class="card">
          <div class="drop" id="dz" onclick="document.getElementById('files').click()" ondrop="onDrop(event)" ondragover="onOver(event)" ondragleave="onLeave(event)">
            <div class="ic">&#128228;</div>
            <div class="ti">Arrastre aqui los dos reportes &mdash; MEF y BCE</div>
            <div class="de">o haga clic para seleccionarlos &middot; no importa el orden, el sistema reconoce cada uno &middot; .xls .xlsx</div>
            <div class="fl" id="fnList"></div>
            <input type="file" id="files" accept=".xls,.xlsx" multiple onchange="recibir(this.files)">
          </div>
        </div>
        <div class="card footbar"><span class="muted" id="cargarMsg">Cargue ambos reportes para continuar.</span>
          <div style="display:flex;gap:10px;flex-wrap:wrap;justify-content:flex-end">
            <button class="btn gh" id="btnAuto" onclick="conciliarAutomatico()" disabled>Conciliar automático y generar reporte</button>
            <button class="btn" id="btnRun" onclick="ejecutar()" disabled>Ejecutar Conciliacion &rarr;</button>
          </div></div>
        <div id="msg"></div>
      </section>

      <section id="v-resultados" class="hidden">
        <h3 class="sec"><span class="bar"></span>Conciliacion de Carteras</h3>
        <p class="sub" id="resSub">&mdash;</p>
        <div id="resAlert"><div class="info">No hay una conciliación activa. Ingrese los reportes nuevos en “Recepción de Reportes” para generar esta vista.</div></div>
        <div class="card"><div class="stats" id="stats"></div>
          <div id="condBox"></div>
          <div class="scrollx"><table class="mtable" id="mtable"></table></div></div>
      </section>

      <section id="v-ajustes" class="hidden">
        <h3 class="sec"><span class="bar"></span>Gestion de Observaciones</h3>
        <p class="sub">Se habilita unicamente cuando una o mas carteras no concilian.</p>
        <div id="ajContenido"></div>
      </section>

      <section id="v-quipux" class="hidden">
        <h3 class="sec"><span class="bar"></span>Emision del Oficio de Respuesta</h3>
        <div id="quipuxContenido"></div>
      </section>

      <section id="v-condonados" class="hidden">
        <h3 class="sec"><span class="bar"></span>Condonados CAF</h3>
        <p class="sub">Carga mensual de Financiamiento Compensatorio para Corporacion Andina de Fomento.</p>
        <div class="caf-panel-grid">
          <div class="card">
            <div class="frow">
              <div class="fld"><label>Anio</label><input id="cafAnio" type="number" min="2000" max="2100"></div>
              <div class="fld"><label>Mes</label><select id="cafMes"></select></div>
              <label class="caf-check"><input id="cafAuto" type="checkbox">Rutas automaticas</label>
              <label class="caf-check"><input id="cafBackup" type="checkbox">Crear respaldo</label>
              <div class="fld caf-wide"><label>Raiz pagos CAF</label><textarea id="cafBase" class="caf-path"></textarea></div>
              <div class="fld caf-wide"><label>Carpeta CAF del mes</label><textarea id="cafDir" class="caf-path"></textarea></div>
              <div class="fld caf-wide"><label>Reporte conciliacion BCE</label><textarea id="cafReporte" class="caf-path"></textarea></div>
              <div class="fld caf-wide"><label>Hoja destino</label><input id="cafHoja" type="text" style="width:100%"></div>
            </div>
            <div class="caf-actions">
              <button class="btn gh" id="cafGuardar" type="button">Guardar</button>
              <button class="btn" id="cafEscanear" type="button">Escanear carpetas</button>
              <button class="btn alt" id="cafAplicar" type="button" disabled>Aplicar al reporte</button>
            </div>
          </div>
          <div>
            <div class="stats caf-stats">
              <div class="stat"><div class="n" id="cafTotalCarpetas">0</div><div class="l">Estados cuenta</div></div>
              <div class="stat gold"><div class="n" id="cafTotalValor">0.00</div><div class="l">Total condonado</div></div>
              <div class="stat"><div class="n" id="cafEstado">Listo</div><div class="l">Estado del modulo</div></div>
            </div>
            <div class="note-warn">Revise los valores antes de aplicar. Las filas en manual pueden editarse directamente en la tabla.</div>
          </div>
        </div>
        <div class="card">
          <div class="scrollx"><table class="mtable caf-table">
            <thead><tr><th>Referencia</th><th>SIGADE</th><th>Carpeta</th><th>Condonado</th><th>Fuente</th><th>Estado cuenta</th></tr></thead>
            <tbody id="cafRows"><tr><td colspan="6" class="muted">Sin datos cargados.</td></tr></tbody>
          </table></div>
        </div>
        <div class="card">
          <div class="diaghead">Registro<small>Actividad del modulo de condonados CAF.</small></div>
          <pre class="caf-log" id="cafLog"></pre>
        </div>
      </section>

      <section id="v-agenda" class="hidden">
        <div class="agenda-head-row">
          <div>
            <h3 class="sec"><span class="bar"></span>Agenda Diaria de Pagos</h3>
            <p class="sub">Resumen preventivo de los préstamos que deben pagarse según la fecha valor o fecha de débito indicada en los oficios.</p>
            <span class="agenda-last-run" id="agendaLastRun">Esperando la primera lectura automática.</span>
          </div>
          <div class="agenda-actions">
            <button class="btn gh" id="agendaOpenSource" type="button">Abrir Acreedores</button>
            <button class="btn gh" id="agendaEnableAlerts" type="button">Activar alertas</button>
            <button class="btn" id="agendaRefresh" type="button">Actualizar ahora</button>
          </div>
        </div>

        <div class="agenda-notice busy" id="agendaNotice">
          <div><strong id="agendaMonitorTitle">Analizando los oficios</strong><span id="agendaMonitorDetail">La revisión automática se ejecuta cada cinco minutos.</span></div>
          <div class="agenda-live"><span class="agenda-live-dot"></span><span id="agendaNextRun">Monitor iniciando</span></div>
        </div>

        <div class="stats agenda-stats">
          <div class="stat gold"><div class="n" id="agendaToday">0</div><div class="l">Pagos para hoy</div></div>
          <div class="stat"><div class="n" id="agendaUpcoming">0</div><div class="l">Próximos 7 días</div></div>
          <div class="stat bad"><div class="n" id="agendaOverdue">0</div><div class="l">Vencidos</div></div>
          <div class="stat ok"><div class="n" id="agendaPaid">0</div><div class="l">Procesados</div></div>
          <div class="stat gold"><div class="n" id="agendaReview">0</div><div class="l">Por revisar</div></div>
        </div>

        <div class="agenda-today-grid">
          <div class="card agenda-today-card">
            <div><span class="kicker">Resumen del día</span><h4 id="agendaTodayHeadline">Revisando pagos programados...</h4><p id="agendaTodayDetail">Se prioriza fecha de débito y luego fecha valor.</p></div>
            <div class="agenda-today-date"><strong id="agendaTodayDay">--</strong><span id="agendaTodayMonth">Fecha actual</span></div>
          </div>
          <div class="card agenda-total-card">
            <span class="kicker">Total pendiente por moneda</span><h4>Pagos de hoy</h4><div class="agenda-total-list" id="agendaTodayTotals"><span class="agenda-total-empty">Sin valores identificados.</span></div>
          </div>
        </div>

        <div class="card">
          <div class="agenda-panel-head">
            <div><h4>Préstamos identificados</h4><p id="agendaResultDescription">Esperando los resultados del monitor.</p></div>
            <div class="agenda-filters">
              <label>Estado<select id="agendaStatusFilter"><option value="pending">Pendientes</option><option value="today">Para hoy</option><option value="upcoming">Próximos</option><option value="overdue">Vencidos</option><option value="paid">Procesados</option><option value="review">Por revisar</option><option value="all">Todos</option></select></label>
              <label>Acreedor<select id="agendaCreditorFilter"><option value="all">Todos</option></select></label>
              <label>Buscar<input id="agendaSearchFilter" type="text" placeholder="Préstamo, oficio o carpeta"></label>
            </div>
          </div>
          <div class="agenda-error hidden" id="agendaError"></div>
          <div class="scrollx"><table class="mtable agenda-table">
            <thead><tr><th>Préstamo</th><th>Fecha valor / débito</th><th>Vencimiento</th><th>Monto</th><th>Moneda</th><th>Corresponsal</th><th>Acreedor / oficio</th><th>Estado</th></tr></thead>
            <tbody id="agendaRows"><tr><td colspan="8" class="muted">Analizando oficios, estados de cuenta y formularios automáticamente...</td></tr></tbody>
          </table></div>
          <div class="agenda-source-line"><strong>Fuente única:</strong><code id="agendaSource">Z:\GISI\SSFI\GESTIÓN PAGOS INTERNACIONALES\2026\DEUDA EXTERNA PÚBLICA\Acreedores Internacionales</code></div>
        </div>

        <div class="card">
          <div class="diaghead">Reglas del resumen<small>Fecha de débito tiene prioridad sobre fecha valor. La fecha de vencimiento es informativa. Los casos sin préstamo o con datos incompletos permanecen en revisión y nunca se asignan por suposición.</small></div>
        </div>
      </section>

      <section id="v-archivoquipux" class="hidden">
        <div class="quipux-head">
          <div>
            <h3 class="sec"><span class="bar"></span>Archivo automático de documentos Quipux</h3>
            <p class="sub">Vigila la bandeja Deuda Quipux, reconoce el oficio, estado de cuenta y formulario, y los copia con el nombre institucional en la carpeta exacta del préstamo.</p>
          </div>
          <div class="quipux-head-actions">
            <span class="quipux-last" id="quipuxLastRun">Monitor iniciando...</span>
            <button class="btn gh" id="quipuxRefresh" type="button">Actualizar vista</button>
            <button class="btn" id="quipuxScan" type="button">Revisar ahora</button>
          </div>
        </div>

        <div class="quipux-notice" id="quipuxNotice">
          <strong>Monitor automático activo</strong>
          <span>Revisa archivos nuevos cada minuto. Conserva los originales y nunca sobrescribe un destino.</span>
        </div>

        <div class="stats quipux-stats">
          <div class="stat"><div class="n" id="quipuxDetected">0</div><div class="l">Pagos de deuda detectados</div></div>
          <div class="stat ok"><div class="n" id="quipuxArchived">0</div><div class="l">Archivados</div></div>
          <div class="stat bad"><div class="n" id="quipuxReview">0</div><div class="l">Requieren revisión</div></div>
          <div class="stat gold"><div class="n" id="quipuxIncomplete">0</div><div class="l">Expedientes incompletos</div></div>
          <div class="stat"><div class="n" id="quipuxHistorical">0</div><div class="l">Históricos protegidos</div></div>
        </div>

        <div class="card quipux-config">
          <details>
            <summary><strong>Configuración y protecciones</strong><span>Mostrar / ocultar</span></summary>
            <div class="quipux-config-grid">
              <div class="fld"><label>Bandeja exclusiva de descarga</label><input id="quipuxSource" type="text" spellcheck="false"></div>
              <div class="fld"><label>Archivo de Acreedores Internacionales</label><input id="quipuxDestination" type="text" spellcheck="false"></div>
              <label class="quipux-check"><input id="quipuxOcr" type="checkbox" checked>Usar OCR cuando el PDF esté escaneado</label>
              <label class="quipux-check"><input id="quipuxAutomatic" type="checkbox" checked>Procesar automáticamente cada minuto</label>
              <div class="quipux-actions"><button class="btn gh" id="quipuxSaveConfig" type="button">Guardar configuración</button></div>
            </div>
          </details>
        </div>

        <div class="card">
          <div class="diaghead">Documentos controlados<small id="quipuxMeta">Los documentos existentes al activar el módulo se protegen como históricos; solamente se procesan descargas nuevas.</small></div>
          <div class="scrollx"><table class="mtable quipux-table">
            <thead><tr><th>Descarga</th><th>Archivo original</th><th>Tipo</th><th>Préstamo / SIGADE</th><th>Fecha valor</th><th>Carpeta destino</th><th>Nombre final</th><th>Método de match</th><th>Estado</th><th>Detalle</th></tr></thead>
            <tbody id="quipuxRows"><tr><td colspan="10" class="muted">Cargando la bandeja de documentos Quipux...</td></tr></tbody>
          </table></div>
        </div>

        <div class="card">
          <div class="diaghead">Reglas de seguridad<small>Match automático solo con préstamo, SIGADE u oficio exacto. Un soporte genérico hereda el destino únicamente si pertenece a un lote con un solo oficio exacto y un solo documento de cada tipo.</small></div>
        </div>
      </section>

      <section id="v-contratosagencia" class="hidden">
        <div class="contract-head">
          <div>
            <h3 class="sec"><span class="bar"></span>Oficios de Contratos de Agencia Fiscal</h3>
            <p class="sub">Identifica únicamente oficios finales firmados, aprende el nombre usado en los expedientes del acreedor y los archiva en la carpeta contractual exacta.</p>
          </div>
          <div class="contract-head-actions">
            <span class="contract-last" id="contractLastRun">Monitor iniciando...</span>
            <button class="btn gh" id="contractRefresh" type="button">Actualizar vista</button>
            <button class="btn" id="contractScan" type="button">Revisar ahora</button>
          </div>
        </div>

        <div class="contract-notice" id="contractNotice">
          <strong>Monitor automático activo</strong>
          <span>Revisa la bandeja cada minuto. Solo guarda oficios finales firmados y nunca sobrescribe un destino.</span>
        </div>

        <div class="stats contract-stats">
          <div class="stat"><div class="n" id="contractCandidates">0</div><div class="l">Oficios contractuales</div></div>
          <div class="stat ok"><div class="n" id="contractArchived">0</div><div class="l">Archivados</div></div>
          <div class="stat bad"><div class="n" id="contractReview">0</div><div class="l">Requieren revisión</div></div>
          <div class="stat gold"><div class="n" id="contractFolders">0</div><div class="l">Carpetas aprendidas</div></div>
          <div class="stat"><div class="n" id="contractPatterns">0</div><div class="l">Nombres de referencia</div></div>
        </div>

        <div class="card contract-config">
          <details>
            <summary><strong>Configuración y protecciones</strong><span>Mostrar / ocultar</span></summary>
            <div class="contract-config-grid">
              <div class="fld"><label>Bandeja de descargas Quipux</label><input id="contractSource" type="text" spellcheck="false"></div>
              <div class="fld"><label>Raíz de Acreedores Internacionales</label><input id="contractDestination" type="text" spellcheck="false"></div>
              <label class="contract-check"><input id="contractOcr" type="checkbox" checked>Usar OCR cuando el oficio esté escaneado</label>
              <label class="contract-check"><input id="contractAutomatic" type="checkbox" checked>Procesar automáticamente cada minuto</label>
              <div class="contract-actions"><button class="btn gh" id="contractSaveConfig" type="button">Guardar configuración</button></div>
            </div>
          </details>
        </div>

        <div class="card">
          <div class="diaghead">Oficios contractuales controlados<small id="contractMeta">Solo se muestran documentos relacionados con Contratos de Agencia Fiscal; los demás PDF de Descargas se ignoran.</small></div>
          <div class="scrollx"><table class="mtable contract-table">
            <thead><tr><th>Descarga</th><th>Oficio original</th><th>Asunto / etapa</th><th>Contrato / SIGADE</th><th>Acreedor / carpeta</th><th>Nombre final</th><th>Método de match</th><th>Estado</th><th>Detalle</th></tr></thead>
            <tbody id="contractRows"><tr><td colspan="9" class="muted">Analizando la bandeja y aprendiendo los expedientes contractuales...</td></tr></tbody>
          </table></div>
        </div>

        <div class="card">
          <div class="diaghead">Patrones de nombres aprendidos<small>Se obtienen de los PDF que ya existen en las carpetas “Contratos Agencia Fiscal” de cada acreedor.</small></div>
          <div class="contract-patterns" id="contractLearnedPatterns"><span class="muted">Cargando patrones...</span></div>
        </div>

        <div class="card">
          <div class="diaghead">Reglas de seguridad<small>Se excluyen memorandos, documentos TEMP, archivos sin firma electrónica, contratos adjuntos y oficios de pago. Una coincidencia ambigua permanece en revisión y nunca se copia.</small></div>
        </div>
      </section>

      <section id="v-matrizprestamos" class="hidden">
        <div class="matrix-head">
          <div>
            <h3 class="sec"><span class="bar"></span>Matriz de Préstamos Realizados</h3>
            <p class="sub">Cargue una captura completa del SGI. El sistema extrae los datos localmente y solo actualiza el Excel después de su revisión y confirmación.</p>
            <div class="matrix-live">OCR local · respaldo automático · control de duplicados</div>
          </div>
          <div><div class="matrix-path" id="matrixWorkbookPath">Consultando la matriz...</div><button class="btn gh" id="matrixConfigButton" type="button">Cambiar ruta</button></div>
        </div>

        <div class="matrix-grid">
          <div class="matrix-card">
            <div class="diaghead">1. Captura del SGI<small>Use la pantalla completa “Mantenimiento Giros Al Exterior - Deuda Externa”.</small></div>
            <div class="matrix-drop" id="matrixDrop">
              <div><div style="font-size:34px">▣</div><strong>Arrastre o seleccione el screenshot</strong><span>PNG, JPG, WEBP o BMP · la imagen permanece en este equipo</span></div>
              <input id="matrixCapture" type="file" accept="image/png,image/jpeg,image/webp,image/bmp">
            </div>
            <img class="matrix-preview" id="matrixPreview" alt="Vista previa de la captura SGI">
            <div class="matrix-notice" id="matrixNotice"><span>Seleccione una captura para comenzar.</span><b class="matrix-score" id="matrixScore">—</b></div>
            <ul class="matrix-warnings" id="matrixWarnings"></ul>
          </div>

          <div class="matrix-card">
            <div class="diaghead">2. Revisar información<small>Los campos con asterisco son obligatorios. Puede corregir cualquier dato antes de guardar.</small></div>
            <div class="matrix-form">
              <div><label class="matrix-required">Fecha ingreso SGI</label><input id="matrixEntryDate" type="date"></div>
              <div><label class="matrix-required">Fecha valor</label><input id="matrixValueDate" type="date"></div>
              <div><label class="matrix-required">Moneda</label><select id="matrixCurrency"><option>USD</option><option>EUR</option><option>JPY</option><option>CNY</option><option>CHF</option><option>GBP</option></select></div>
              <div><label class="matrix-required">Préstamo</label><input id="matrixLoan" type="text" placeholder="BIRF-9595-EC"></div>
              <div class="wide"><label class="matrix-required">Referencia SWIFT</label><input id="matrixSwift" type="text" placeholder="TF-01-7712600697"></div>
              <div><label>Operación local</label><input id="matrixLocal" type="text" placeholder="Opcional"></div>
              <div><label class="matrix-required">Acreedor</label><input id="matrixLender" type="text" placeholder="BIRF"></div>
              <div><label class="matrix-required">Prestatario</label><input id="matrixBorrower" type="text" placeholder="MEF"></div>
              <div><label class="matrix-required">Monto</label><input id="matrixAmount" type="number" min="0.01" step="0.01"></div>
              <div class="wide"><label>Corresponsal</label><input id="matrixCorrespondent" type="text"></div>
              <div class="wide"><label>Banco beneficiario identificado</label><input id="matrixBeneficiaryBank" type="text" readonly></div>
              <div class="full"><div class="matrix-notice" id="matrixCorrespondentEvidence"><span>El corresponsal se validará con la guía y la matriz histórica.</span><b class="matrix-score">—</b></div></div>
              <div><label>Contabilizar otras monedas</label><input id="matrixAccountingDate" type="date"></div>
              <div><label class="matrix-required">Estado</label><select id="matrixStatus"><option>Ingresado</option><option>Procesado</option><option>Pendiente</option></select></div>
              <div><label>Oficio identificado</label><input id="matrixOffice" type="text" readonly></div>
              <div><label>Pedido identificado</label><input id="matrixOrder" type="text" readonly></div>
              <div class="full"><label>Notas</label><textarea id="matrixNotes"></textarea></div>
            </div>
            <div class="matrix-actions"><button class="btn gh" id="matrixReset" type="button">Nueva captura</button><button class="btn" id="matrixSave" type="button" disabled>Guardar fila en Excel</button></div>
          </div>
        </div>

        <div class="matrix-card matrix-history">
          <div class="diaghead">Últimas filas agregadas desde GIDEP<small>Cada guardado conserva un respaldo anterior de la matriz.</small></div>
          <div class="scrollx"><table><thead><tr><th>Fecha</th><th>Usuario</th><th>Préstamo</th><th>Referencia SWIFT</th><th>Fecha valor</th><th>Monto</th><th>Fila</th></tr></thead><tbody id="matrixHistory"><tr><td colspan="7" class="matrix-empty">Aún no hay filas agregadas desde este módulo.</td></tr></tbody></table></div>
        </div>
      </section>

      <section id="v-activaciones" class="hidden">
        <div class="activation-head">
          <div>
            <h3 class="sec"><span class="bar"></span>Activaciones Cuentas</h3>
            <p class="sub">Validación institucional de cuentas para mensajes Swift e ISO 20022. Incluye IBAN, ABA Routing, CLABE, Transit de Canadá y BSB de Australia.</p>
          </div>
          <div class="activation-source"><strong>Validador integrado</strong><code>C:\Users\bromo\Desktop\08 Otros\Validador Swift FINAL.HTML</code><span id="activationStatus">Comprobando disponibilidad...</span><a class="btn gh activation-swiftref" href="https://www.swiftref.com/en/bicsearch" target="_blank" rel="noopener noreferrer">Consultar BIC en SWIFTRef ↗</a></div>
        </div>
        <div class="card activation-frame-card">
          <iframe class="activation-frame" id="activationFrame" src="/activaciones-cuentas/validador" title="Validador institucional Swift e ISO 20022" loading="lazy"></iframe>
        </div>
      </section>

      <section id="v-respuestas" class="hidden">
        <div class="mail-head">
          <div>
            <h3 class="sec"><span class="bar"></span>Análisis de respuestas de correos</h3>
            <p class="sub">Asistente local que aprende de las respuestas aprobadas y propone borradores para correos similares.</p>
          </div>
          <div class="mail-privacy"><span class="dot ok"></span>Aprendizaje local · los correos no salen del equipo</div>
        </div>
        <div class="firm-flow-tabs mail-flow-tabs" id="mailFlowTabs" role="tablist" aria-label="Flujo de análisis de correos">
          <button class="act" type="button" data-mail-flow="analyze"><span>Analizar correo</span><small>Preparar respuesta</small></button>
          <button type="button" data-mail-flow="learn"><span>Enseñar respuesta</span><small>Guardar ejemplo aprobado</small></button>
          <button type="button" data-mail-flow="library"><span>Base aprendida</span><small id="mailExampleCount">0 ejemplos</small></button>
        </div>

        <div class="mail-panel act" data-mail-panel="analyze">
          <div class="mail-grid">
            <div class="card">
              <div class="diaghead">Correo recibido<small>Copie el asunto y el mensaje que necesita responder.</small></div>
              <div class="mail-field"><label>Categoría</label><select id="mailAnalyzeCategory"><option>General</option><option>Conciliación</option><option>Pagos</option><option>Comprobantes</option><option>Firmas</option><option>Activación de cuentas</option><option>Observaciones</option></select></div>
              <div class="mail-field"><label>Asunto</label><input id="mailAnalyzeSubject" type="text" placeholder="Ej.: Observación sobre pago BID"></div>
              <div class="mail-field"><label>Contenido del correo</label><textarea id="mailAnalyzeIncoming" class="mail-textarea" placeholder="Pegue aquí el correo recibido..."></textarea></div>
              <div class="mail-actions"><button class="btn" id="mailSuggest" type="button">Analizar y sugerir respuesta</button><button class="btn gh" id="mailReset" type="button">Nuevo correo</button><span class="mail-confidence" id="mailConfidence">Sin análisis</span></div>
            </div>
            <div class="card">
              <div class="diaghead">Borrador sugerido<small>Revíselo y edítelo antes de utilizarlo.</small></div>
              <div class="mail-status" id="mailAnalyzeStatus">Ingrese un correo para buscar respuestas aprobadas similares.</div>
              <div class="mail-field"><label>Respuesta propuesta</label><textarea id="mailSuggestedResponse" class="mail-textarea mail-response" placeholder="La sugerencia aparecerá aquí..."></textarea></div>
              <div class="mail-actions"><button class="btn gh" id="mailCopySuggestion" type="button">Copiar respuesta</button><button class="btn alt" id="mailLearnSuggestion" type="button">Guardar como respuesta aprobada</button></div>
              <div class="mail-match-list" id="mailMatches"></div>
            </div>
          </div>
        </div>

        <div class="mail-panel" data-mail-panel="learn">
          <div class="mail-grid">
            <div class="card">
              <div class="diaghead">Ejemplo recibido<small>Este texto se usará para reconocer casos parecidos.</small></div>
              <div class="mail-field"><label>Categoría</label><select id="mailLearnCategory"><option>General</option><option>Conciliación</option><option>Pagos</option><option>Comprobantes</option><option>Firmas</option><option>Activación de cuentas</option><option>Observaciones</option></select></div>
              <div class="mail-field"><label>Asunto</label><input id="mailLearnSubject" type="text" placeholder="Asunto del correo"></div>
              <div class="mail-field"><label>Correo recibido</label><textarea id="mailLearnIncoming" class="mail-textarea" placeholder="Correo que originó la respuesta..."></textarea></div>
            </div>
            <div class="card">
              <div class="diaghead">Respuesta aprobada<small>Guarde únicamente versiones que ya considere correctas.</small></div>
              <div class="mail-field"><label>Respuesta final</label><textarea id="mailLearnResponse" class="mail-textarea mail-response" placeholder="Respuesta aprobada..."></textarea></div>
              <div class="mail-field"><label>Observaciones</label><input id="mailLearnNotes" type="text" placeholder="Criterio, excepción o dato que debe conservarse"></div>
              <div class="mail-status" id="mailLearnStatus">La base de aprendizaje está lista para recibir ejemplos.</div>
              <button class="btn" id="mailSaveExample" type="button">Guardar y enseñar al asistente</button>
            </div>
          </div>
        </div>

        <div class="mail-panel" data-mail-panel="library">
          <div class="card">
            <div class="diaghead">Base de respuestas aprendidas<small>Ejemplos guardados localmente y número de veces que fueron utilizados como referencia.</small></div>
            <div class="scrollx"><table class="mtable mail-table"><thead><tr><th>Fecha</th><th>Categoría</th><th>Asunto</th><th>Correo recibido</th><th>Respuesta aprobada</th><th>Usos</th></tr></thead><tbody id="mailLibraryRows"><tr><td colspan="6" class="muted">Aún no hay ejemplos guardados.</td></tr></tbody></table></div>
          </div>
        </div>
      </section>

      <section id="v-firmados" class="hidden" data-firm-version="3">
        <h3 class="sec"><span class="bar"></span>Firma EC</h3>
        <p class="sub">Control independiente de los mensajes de Deuda Externa, Automáticos y Bancos. El sistema vigila automáticamente los tres orígenes casi en tiempo real y avisa cuando un documento requiere firma.</p>
        <div class="firm-flow-tabs" id="firmFlowTabs" role="tablist" aria-label="Tipos de firma EC">
          <button class="act" type="button" data-flow="deuda_externa"><span>Firma Deuda Externa</span><small id="firmFlowCount-deuda_externa">0 pendientes</small></button>
          <button type="button" data-flow="automaticos"><span>Firmas Automáticos</span><small id="firmFlowCount-automaticos">0 pendientes</small></button>
          <button type="button" data-flow="bancos"><span>Firma Bancos</span><small id="firmFlowCount-bancos">0 pendientes</small></button>
        </div>
        <div class="firm-flow-heading">
          <div><strong id="firmFlowTitle">Firma Deuda Externa</strong><span id="firmFlowDescription">Mensajes de pagos de deuda externa ingresados por Bryan.</span></div>
          <span class="firm-monitor-badge">Monitor en tiempo real · cada 3 segundos</span>
        </div>
        <div class="firm-alert" id="firmAlert">Revisando los tres orígenes de Firma EC...</div>

        <div class="stats firm-stats firm-stats-v2">
          <div class="stat gold"><div class="n" id="firmPending">0</div><div class="l">Por firmar</div></div>
          <div class="stat"><div class="n" id="firmDetected">0</div><div class="l">Doble firma</div></div>
          <div class="stat ok"><div class="n" id="firmReady">0</div><div class="l">Listos para copiar</div></div>
          <div class="stat"><div class="n" id="firmExisting">0</div><div class="l">Ya en destino</div></div>
          <div class="stat bad"><div class="n" id="firmConflicts">0</div><div class="l">Conflictos</div></div>
        </div>

        <div class="card firm-route-card">
          <span class="label">Origen del flujo</span><code id="firmSource">-</code><button class="btn gh" id="firmOpenSource" type="button">Abrir origen</button>
          <span class="label">Destino diario</span><code id="firmTarget">-</code><button class="btn gh" id="firmOpenTarget" type="button">Abrir destino</button>
          <span class="label">Fecha destino</span><span class="firm-date" id="firmTargetDate">-</span><span></span>
          <span class="label">Carpetas revisadas</span><span class="firm-date firm-scanned" id="firmScannedFolders">-</span><span></span>
        </div>

        <div class="firm-signature-panel hidden" id="firmSignaturePanel">
          <div class="firm-signature-head"><strong>Mensajes que requieren su firma</strong><span id="firmSignatureCount">0 pendientes</span></div>
          <div class="firm-signature-list" id="firmSignatureList"></div>
        </div>

        <div class="card">
          <div class="firm-head-actions">
            <button class="btn gh" id="firmEnableAlerts" type="button">Activar alertas</button>
            <button class="btn gh" id="firmRefresh" type="button">Actualizar vista</button>
            <button class="btn gh" id="firmSelectAll" type="button">Seleccionar listos</button>
            <button class="btn" id="firmMove" type="button" disabled>Copiar seleccionados</button>
          </div>
          <div class="scrollx"><table class="mtable firm-table-v2">
            <thead><tr><th></th><th>Carpeta origen</th><th>Referencia SWIFT</th><th>Archivo</th><th>Fecha archivo</th><th>Tamaño</th><th>Estado</th></tr></thead>
            <tbody id="firmRows"><tr><td colspan="7" class="muted">Cargando mensajes firmados...</td></tr></tbody>
          </table></div>
          <div class="firm-last-run" id="firmLastRun">Pendiente de lectura.</div>
        </div>
        <div class="card">
          <div class="diaghead">Protecciones<small>El monitoreo es automático. La copia no sobrescribe y conserva en el origen las versiones de una y dos firmas.</small></div>
          <pre class="firm-log" id="firmLog"></pre>
        </div>
      </section>

      <section id="v-comprobantes" class="hidden">
        <div class="comp-head-row">
          <div>
            <h3 class="sec"><span class="bar"></span>Comprobantes Contables y ACK</h3>
            <p class="sub">Revisa automaticamente los pagos de BRYAN y STEVEN, encuentra la carpeta del prestamo y completa el ACK cuando este disponible.</p>
          </div>
          <div class="comp-actions">
            <span class="comp-last-run" id="compLastRun">Monitor iniciando...</span>
            <button class="btn gh" id="compRefrescar" type="button">Actualizar vista</button>
            <button class="btn" id="compAnalizar" type="button">Analizar ahora</button>
          </div>
        </div>

        <div class="comp-notice" id="compAutoNotice">
          <strong>Modo automatico activo</strong>
          <span>Cada cinco minutos revisa BRYAN, STEVEN y los ACK. Solo archiva coincidencias exactas de 99-100%.</span>
        </div>

        <div class="stats comp-stats-v2">
          <div class="stat"><div class="n" id="compDetectados">0</div><div class="l">Detectados</div></div>
          <div class="stat ok"><div class="n" id="compListos">0</div><div class="l">Listos</div></div>
          <div class="stat"><div class="n" id="compArchivados">0</div><div class="l">Archivados</div></div>
          <div class="stat gold"><div class="n" id="compAckPendientes">0</div><div class="l">ACK pendientes</div></div>
          <div class="stat bad"><div class="n" id="compRevision">0</div><div class="l">Requieren revision</div></div>
        </div>

        <div class="card comp-config-card">
          <details>
            <summary><strong>Configuracion de rutas</strong><span>Mostrar / ocultar</span></summary>
            <div class="comp-config-grid">
              <div class="fld"><label>Carpeta de BRYAN</label><input id="compBryan" type="text" spellcheck="false"></div>
              <div class="fld"><label>Carpeta de STEVEN</label><input id="compSteven" type="text" spellcheck="false"></div>
              <div class="fld"><label>Base de ACK</label><input id="compAckBase" type="text" spellcheck="false"></div>
              <div class="fld"><label>Archivo de Acreedores Internacionales</label><input id="compDestino" type="text" spellcheck="false"></div>
              <label class="comp-check"><input id="compRecursivo" type="checkbox">Incluir niveles historicos adicionales</label>
              <label class="comp-check"><input id="compOcr" type="checkbox" checked>Usar OCR para PDF escaneados</label>
              <label class="comp-check"><input id="compAutomatico" type="checkbox" checked>Archivo automatico cada cinco minutos</label>
              <div class="comp-actions"><button class="btn gh" id="compGuardarConfig" type="button">Guardar configuracion</button></div>
            </div>
          </details>
        </div>

        <div class="card comp-control-card">
          <div class="comp-control-head">
            <div><strong>Control diario de archivo</strong><span>La vista inicia con los pendientes de la fecha actual; el historial permanece disponible en “Todos”.</span></div>
            <span class="comp-filter-meta" id="compFilterMeta">0 comprobantes visibles</span>
          </div>
          <div class="comp-view-tabs" id="compViewTabs" role="tablist" aria-label="Estado de comprobantes">
            <button class="act" type="button" data-comp-view="pending">Pendientes <b id="compCountPending">0</b></button>
            <button type="button" data-comp-view="ready">Por archivar <b id="compCountReady">0</b></button>
            <button type="button" data-comp-view="ack">ACK pendientes <b id="compCountAck">0</b></button>
            <button type="button" data-comp-view="review">Revisión <b id="compCountReview">0</b></button>
            <button type="button" data-comp-view="archived">Archivados <b id="compCountArchived">0</b></button>
            <button type="button" data-comp-view="all">Todos <b id="compCountAll">0</b></button>
          </div>
          <div class="comp-filter-grid">
            <label>Fecha valor<select id="compDateMode"><option value="today">Hoy</option><option value="last7">Últimos 7 días</option><option value="all">Todas las fechas</option></select></label>
            <label>Responsable<select id="compOwnerFilter"><option value="all">BRYAN y STEVEN</option><option value="BRYAN">BRYAN</option><option value="STEVEN">STEVEN</option></select></label>
            <label>Buscar comprobante, referencia o préstamo<input id="compSearch" type="search" placeholder="Ej. 771-1057, TF-01 o BIRF-9595"></label>
          </div>
        </div>

        <div class="card">
          <div class="diaghead">
            Comprobantes según los filtros
            <small id="compMeta">Esperando el primer analisis.</small>
          </div>
          <div class="comp-table-actions">
            <label class="comp-check"><input id="compSeleccionarTodos" type="checkbox">Seleccionar coincidencias seguras</label>
            <span id="compSeleccionCuenta" class="muted">0 seleccionados</span>
            <button class="btn" id="compArchivar" type="button" disabled>Archivar seleccionados</button>
          </div>
          <div class="scrollx"><table class="mtable comp-table-v2">
            <thead><tr><th></th><th>Responsable</th><th>Comprobante</th><th>Referencia</th><th>Fecha valor</th><th>Prestamo</th><th>Carpeta destino</th><th>Metodo de match</th><th>Estado</th><th>Detalle</th></tr></thead>
            <tbody id="compRows"><tr><td colspan="10" class="muted">Cargando comprobantes y ACK...</td></tr></tbody>
          </table></div>
        </div>

        <div class="card">
          <div class="diaghead">Registro<small>Actividad visible de este modulo.</small></div>
          <pre class="comp-log" id="compLog"></pre>
        </div>
      </section>
    </div>
  </div>
<script>
const CSRF_TOKEN={{ csrf_token()|tojson }};
const CURRENT_USER_ROLE={{ user_role|tojson }};
const GIDEP_NATIVE_FETCH=window.fetch.bind(window);
window.fetch=(input,init={})=>{
  const options={...init};
  const method=String(options.method||(input instanceof Request?input.method:"GET")).toUpperCase();
  const url=new URL(input instanceof Request?input.url:String(input),window.location.href);
  if(url.origin===window.location.origin&&["POST","PUT","PATCH","DELETE"].includes(method)){
    const headers=new Headers(options.headers||(input instanceof Request?input.headers:undefined));
    headers.set("X-CSRF-Token",CSRF_TOKEN);
    options.headers=headers;
  }
  return GIDEP_NATIVE_FETCH(input,options);
};

const CONTROL={timer:null};
const controlEsc=value=>String(value??"").replace(/[&<>"']/g,char=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#039;"}[char]));
const controlDate=value=>{if(!value)return "Sin lectura todavía";const parsed=new Date(value);return Number.isNaN(parsed.getTime())?String(value):new Intl.DateTimeFormat("es-EC",{dateStyle:"short",timeStyle:"short"}).format(parsed);};
function controlOpen(view){abrirVista(view);}
function controlRender(payload){
  const cards=document.getElementById("controlCards");
  cards.innerHTML=(payload.cards||[]).map(item=>`<button type="button" class="control-card ${controlEsc(item.severity||"ok")}" data-control-view="${controlEsc(item.view)}"><span class="control-value">${controlEsc(item.value||0)}</span><span class="control-label">${controlEsc(item.label)}</span><span class="control-action">Abrir módulo →</span></button>`).join("");
  cards.querySelectorAll("[data-control-view]").forEach(button=>button.addEventListener("click",()=>controlOpen(button.dataset.controlView)));
  const issues=document.getElementById("controlIssues"),items=payload.issues||[];
  issues.innerHTML=items.length?items.map(item=>`<div class="control-issue ${controlEsc(item.severity||"warning")}"><span class="control-issue-dot"></span><div><strong>${controlEsc(item.module)}</strong><small>${controlEsc(item.message)}</small></div><button type="button" data-control-view="${controlEsc(item.view)}">Revisar</button></div>`).join(""):'<div class="control-empty">No hay pendientes ni errores reportados por los monitores.</div>';
  issues.querySelectorAll("[data-control-view]").forEach(button=>button.addEventListener("click",()=>controlOpen(button.dataset.controlView)));
  const labels={agenda:"Agenda",quipux:"Archivo Quipux",contratos:"Contratos",firmas:"Firma EC",comprobantes:"Comprobantes"};
  document.getElementById("controlHealth").innerHTML=Object.entries(payload.last_updates||{}).map(([key,value])=>`<div class="control-health-row"><strong>${controlEsc(labels[key]||key)}</strong><time>${controlEsc(controlDate(value))}</time></div>`).join("")+`<span class="control-health-badge ${payload.ok?"":"bad"}">${controlEsc(payload.healthy_modules||0)}/${controlEsc(payload.total_modules||5)} monitores sin error</span>`;
  document.getElementById("controlUpdated").textContent=`Control actualizado: ${controlDate(payload.generated_at)}`;
}
async function controlLoad(){const button=document.getElementById("controlRefresh");if(button)button.disabled=true;try{const response=await fetch("/api/control-operativo/status",{cache:"no-store"});const payload=await response.json();if(!response.ok)throw new Error(payload.error||`HTTP ${response.status}`);controlRender(payload);}catch(error){document.getElementById("controlIssues").innerHTML=`<div class="control-issue danger"><span class="control-issue-dot"></span><div><strong>No se pudo actualizar el control</strong><small>${controlEsc(error.message)}</small></div></div>`;}finally{if(button)button.disabled=false;}}
document.getElementById("controlRefresh").addEventListener("click",controlLoad);
CONTROL.timer=window.setInterval(controlLoad,30000);
if(CURRENT_USER_ROLE!=="Administrador"){
  ["cafGuardar","quipuxSaveConfig","contractSaveConfig","matrixConfigButton","compGuardarConfig"].forEach(id=>{
    const element=document.getElementById(id);if(element){element.disabled=true;element.title="La configuración requiere rol Administrador";}
  });
}

const CONCEPTOS=["Desembolsos","Amortizaciones","Intereses","Comisiones","Intereses Condonados","Interés por Mora"];
// La matriz principal NO incluye Condonados: es un apartado previo, aparte.
const MATRIZ=["Desembolsos","Amortizaciones","Intereses","Comisiones","Interés por Mora"];
const CONDONADOS="Intereses Condonados";
const RUBRO_TXT={"Desembolsos":"desembolsos","Amortizaciones":"amortizaciones","Intereses":"intereses","Comisiones":"comisiones","Intereses Condonados":"condonados","Interés por Mora":"intereses por mora"};
const CARTERA_TXT={"AMAZON DAC":"AMAZON"};
const ORDEN_QUIPUX=["AIIB","AMAZON DAC","BANCOS","BID","BIRF","BONOS","CAF","FIDA","FLAR","FMI","GOBIERNOS","GPS"];
let FILES=[], DATA=[], TOTALES=[], GTOT=null, DIAG=[], PREST={}, INFO=null, RES=null, PERIODO="", ACTIVE_CONCILIATION=false;
const fmt=n=>(n||0).toLocaleString("es-EC",{minimumFractionDigits:2,maximumFractionDigits:2});
function msg(t,err){const m=document.getElementById("msg");m.textContent=t;m.style.color=err?"#f87171":"#34d399";}
const valido=f=>f&&/\.(xls|xlsx)$/i.test(f.name);
function recibir(files){[...files].forEach(f=>{if(valido(f)&&FILES.length<2&&!FILES.some(x=>x.name===f.name))FILES.push(f);});pintarFiles();}
function pintarFiles(){
  const c=document.getElementById("fnList");
  c.innerHTML=FILES.map((f,i)=>`<span class="chipf">&#128196; ${f.name} <a onclick="quitar(${i});event.stopPropagation()" style="cursor:pointer;color:#f87171">&times;</a></span>`).join("");
  document.getElementById("dz").classList.toggle("set",FILES.length>=2);
  document.getElementById("btnRun").disabled=FILES.length<2;
  document.getElementById("btnAuto").disabled=FILES.length<2;
  document.getElementById("cargarMsg").textContent=FILES.length>=2?"Listo para conciliar.":"Cargue ambos reportes para continuar.";
}
function quitar(i){FILES.splice(i,1);pintarFiles();}
function onOver(e){e.preventDefault();e.currentTarget.classList.add("over");}
function onLeave(e){e.currentTarget.classList.remove("over");}
function onDrop(e){e.preventDefault();e.currentTarget.classList.remove("over");recibir(e.dataTransfer.files);}

async function ejecutar(){
  if(FILES.length<2)return;
  const fd=new FormData();FILES.forEach(f=>fd.append("archivos",f));
  const btn=document.getElementById("btnRun");btn.disabled=true;msg("Procesando...");
  try{const j=await (await fetch("/api/conciliar",{method:"POST",body:fd})).json();
    if(!j.ok){msg(j.error||"Error",true);btn.disabled=false;return;}
    ACTIVE_CONCILIATION=true;aplicarResultado(j);msg("Conciliacion ejecutada para "+(INFO?INFO.texto:PERIODO)+".");irA("resultados");
  }catch(e){msg(e.message,true);}btn.disabled=false;
}

async function conciliarAutomatico(){
  if(FILES.length<2){msg("Cargue los dos reportes nuevos para conciliar.",true);return;}
  const fd=new FormData();FILES.forEach(f=>fd.append("archivos",f));const opts={method:"POST",body:fd};
  const btn=document.getElementById("btnAuto");btn.disabled=true;msg("Conciliando automaticamente y generando reporte...");
  try{const j=await (await fetch("/api/conciliar_automatico",opts)).json();
    if(!j.ok){msg(j.error||"Error",true);btn.disabled=false;return;}
    ACTIVE_CONCILIATION=true;aplicarResultado(j);msg("Conciliacion automatica completa para "+(INFO?INFO.texto:PERIODO)+".");irA("resultados");
  }catch(e){msg(e.message,true);}btn.disabled=false;
}

async function cargarUltimoConciliado(silencioso=true){
  if(!silencioso)alert("Por seguridad no se cargan conciliaciones anteriores. Ingrese reportes nuevos en Recepción de Reportes.");
  return false;
}

function limpiarConciliacionActiva(){
  DATA=[];TOTALES=[];GTOT=null;DIAG=[];PREST={};INFO=null;RES=null;PERIODO="";ACTIVE_CONCILIATION=false;
  const period=document.getElementById("concPortfolioPeriod");if(period)period.textContent="Sin conciliación activa";
  const state=document.getElementById("estadoSide");if(state)state.textContent="Cargue reportes nuevos para consultar el estado.";
  const resSub=document.getElementById("resSub");if(resSub)resSub.textContent="Sin reportes cargados en esta sesión";
  const resAlert=document.getElementById("resAlert");if(resAlert)resAlert.innerHTML='<div class="info">No hay una conciliación activa. Ingrese los reportes nuevos en “Recepción de Reportes” para generar esta vista.</div>';
  const stats=document.getElementById("stats");if(stats)stats.innerHTML="";
  const cond=document.getElementById("condBox");if(cond)cond.innerHTML="";
  const matrix=document.getElementById("mtable");if(matrix)matrix.innerHTML="";
  const adjustments=document.getElementById("ajContenido");if(adjustments)adjustments.innerHTML='<div class="info">Las observaciones se habilitarán después de ejecutar una conciliación nueva.</div>';
  const quipux=document.getElementById("quipuxContenido");if(quipux)quipux.innerHTML='<div class="info">La emisión del oficio se habilitará después de ejecutar una conciliación nueva y completa.</div>';
}

function aplicarResultado(j){
  RES=j;DATA=j.registros;TOTALES=j.totales||[];GTOT=j.gran_total;DIAG=j.diagnostico||[];PREST=j.prestamos||{};INFO=j.info_periodo||INFO;PERIODO=j.periodo;
  const concPeriod=document.getElementById("concPortfolioPeriod");if(concPeriod)concPeriod.textContent=`Conciliación activa: ${INFO?INFO.texto:PERIODO}`;
  pintarSidebar();pintarMatriz();pintarAjustes();pintarQuipux();
  const quipuxTab=document.querySelector('#concFlowTabs [data-conc-view="quipux"]');if(quipuxTab)quipuxTab.classList.remove("dis");
}
function pintarSidebar(){const es=document.getElementById("estadoSide");
  if(es)es.innerHTML=TOTALES.length?TOTALES.map(t=>`<div class="e"><span class="dot ${t.estado==='CONCILIADO'?'ok':'bad'}"></span>${CARTERA_TXT[t.acreedor]||t.acreedor}</div>`).join(""):"Sin datos";}
function pintarMatriz(){
  const piv={};DATA.forEach(r=>{(piv[r.acreedor]=piv[r.acreedor]||{})[r.concepto]={mef:r.mef,bce:r.bce,dif:r.diferencia};});
  const estado={};TOTALES.forEach(t=>estado[t.acreedor]=t.estado);
  // ---- Apartado previo: Intereses Condonados (aparte de la matriz) ----
  const condRows=DATA.filter(r=>r.concepto===CONDONADOS);
  const cb=document.getElementById("condBox");
  if(condRows.length){
    let ch=`<div class="diaghead" style="margin-top:4px">Apartado previo &middot; Intereses Condonados<small>Se revisa antes de la conciliaci&oacute;n de los rubros; es un apartado independiente.</small></div>
      <div class="scrollx" style="max-height:220px;margin-bottom:16px"><table class="mtable"><thead><tr>
      <th class="acr">Acreedor</th><th>MEF</th><th>BCE</th><th>&Delta;</th><th>Estado</th></tr></thead><tbody>`;
    condRows.forEach(r=>{const z=Math.abs(r.diferencia)<0.005;
      ch+=`<tr class="${z?'ok':'bad'}"><td class="acr">${CARTERA_TXT[r.acreedor]||r.acreedor}</td>
        <td class="num">${fmt(r.mef)}</td><td class="num">${fmt(r.bce)}</td>
        <td class="d ${z?'ok':''}">${z?'&#10003;':'+'+fmt(Math.abs(r.diferencia))}</td>
        <td><span class="est-pill ${z?'ok':'bad'}">${z?'CONCILIADO':'DIFERENCIA'}</span></td></tr>`;});
    ch+=`</tbody></table></div>`;
    cb.innerHTML=ch;
  } else cb.innerHTML="";

  let h=`<thead><tr><th class="acr" rowspan="2">Acreedor</th><th class="est" rowspan="2">Estado</th>`;
  MATRIZ.forEach(c=>h+=`<th colspan="3" class="sep">${c}</th>`);
  h+=`</tr><tr class="sub">`;MATRIZ.forEach(()=>h+=`<th class="sep">MEF</th><th>BCE</th><th>&Delta;</th>`);
  h+=`</tr></thead><tbody>`;
  TOTALES.map(t=>t.acreedor).forEach(ac=>{const ok=estado[ac]==="CONCILIADO";
    h+=`<tr class="${ok?'ok':'bad'}"><td class="acr">${CARTERA_TXT[ac]||ac}</td><td><span class="est-pill ${ok?'ok':'bad'}">${ok?'CONCILIADO':'DIFERENCIA'}</span></td>`;
    MATRIZ.forEach(c=>{const d=piv[ac]&&piv[ac][c];
      if(!d){h+=`<td class="num sep muted">&mdash;</td><td class="num muted">&mdash;</td><td class="d ok">&#10003;</td>`;}
      else{const z=Math.abs(d.dif)<0.005;h+=`<td class="num sep">${fmt(d.mef)}</td><td class="num">${fmt(d.bce)}</td><td class="d ${z?'ok':''}">${z?'&#10003;':'+'+fmt(Math.abs(d.dif))}</td>`;}});
    h+=`</tr>`;});
  const tot={};MATRIZ.forEach(c=>tot[c]={mef:0,bce:0});
  DATA.forEach(r=>{if(tot[r.concepto]){tot[r.concepto].mef+=r.mef;tot[r.concepto].bce+=r.bce;}});
  h+=`<tr class="tot"><td class="acr">TOTAL GENERAL</td><td></td>`;
  MATRIZ.forEach(c=>{const m=tot[c].mef,b=tot[c].bce,d=Math.round((m-b)*100)/100,z=Math.abs(d)<0.005;
    h+=`<td class="num sep">${fmt(m)}</td><td class="num">${fmt(b)}</td><td class="d ${z?'ok':''}">${z?'&#10003;':'+'+fmt(Math.abs(d))}</td>`;});
  h+=`</tr></tbody>`;document.getElementById("mtable").innerHTML=h;
  const cartOk=TOTALES.filter(t=>t.estado==='CONCILIADO').length;
  const cartDif=TOTALES.length-cartOk;
  const nombresDif=TOTALES.filter(t=>t.estado!=='CONCILIADO').map(t=>CARTERA_TXT[t.acreedor]||t.acreedor).join(", ");
  document.getElementById("stats").innerHTML=
    `<div class="stat ${cartDif===0?'ok':''}"><div class="n">${cartOk}/${TOTALES.length}</div><div class="l">Carteras conciliadas</div></div>
     <div class="stat bad"><div class="n">${cartDif}</div><div class="l">Carteras con diferencia</div></div>
     <div class="stat gold"><div class="n">${fmt(GTOT.mef)}</div><div class="l">Total MEF (USD)</div></div>
     <div class="stat gold"><div class="n">${fmt(GTOT.bce)}</div><div class="l">Total BCE (USD)</div></div>`;
  document.getElementById("resSub").textContent=`Periodo ${INFO?INFO.texto:PERIODO} · ${cartOk} de ${TOTALES.length} carteras conciliadas`;
  const autoInfo=(RES.auto_ajustes&&RES.auto_ajustes.length)?`<div class="note-ok" style="margin-top:10px">&#9881; Ajustes automaticos aplicados: ${RES.auto_ajustes.length}. El detalle queda en la columna Nota / Ajuste del reporte.</div>`:"";
  const anexInfo=(RES.anexos_aplicados&&RES.anexos_aplicados.length)?`<div class="note-ok" style="margin-top:10px">&#128206; Anexos incorporados al BCE: ${RES.anexos_aplicados.length}.</div>`:"";
  const repBtn=RES.reporte_url?`<div style="margin-top:10px"><button class="btn alt" onclick="window.location=RES.reporte_url">&#11015; Descargar reporte final de conciliacion</button></div>`:"";
  document.getElementById("resAlert").innerHTML = (cartDif===0
    ? `<div class="note-ok">&#10004; Conciliacion completa: las ${TOTALES.length} carteras cuadran al centavo. Puede emitir el oficio de respuesta.</div>`
    : `<div class="note-warn">&#9888; ${cartDif} cartera(s) sin conciliar: ${nombresDif}. Gestione las observaciones antes de emitir el oficio.</div>`) + autoInfo + anexInfo + repBtn;
}

async function resolverConAnexos(){
  if(!PERIODO){alert("Ejecute primero la conciliacion.");return;}
  const btn=document.getElementById("btnAnexos");if(btn)btn.disabled=true;
  msg("Buscando anexos del mes y actualizando el reporte BCE...");
  try{
    const j=await (await fetch("/api/resolver_observaciones_anexos",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify({periodo:PERIODO})})).json();
    if(!j.ok){msg(j.error||"Error",true);if(btn)btn.disabled=false;return;}
    aplicarResultado({...j,info_periodo:INFO});
    const n=(j.anexos_aplicados||[]).length;
    const avis=(j.anexos_avisos||[]).length;
    msg(n?`Anexos aplicados: ${n}.`:(avis?`No se aplicaron anexos; revise avisos.`:"No habia observaciones de pago directo pendientes."));
    if(avis){alert((j.anexos_avisos||[]).join("\n"));}
  }catch(e){msg(e.message,true);}
  if(btn)btn.disabled=false;
}

function pintarAjustes(){
  const cont=document.getElementById("ajContenido");
  if(RES&&RES.conciliado_total){
    const dl=RES.bce_ajustado?`<div style="margin-top:16px"><button class="btn alt" onclick="window.location='/api/descargar_bce_ajustado?periodo=${encodeURIComponent(PERIODO)}'">&#11015; Descargar Reporte Conciliacion BCE (modificado y conciliado)</button></div>`:"";
    const rep=RES.reporte_url?`<div style="margin-top:12px"><button class="btn alt" onclick="window.location=RES.reporte_url">&#11015; Descargar reporte final de conciliacion</button></div>`:"";
    const aj=(RES.auto_ajustes&&RES.auto_ajustes.length)?`<br>Se aplicaron ${RES.auto_ajustes.length} ajustes automaticos auditables.`:"";
    const ax=(RES.anexos_aplicados&&RES.anexos_aplicados.length)?`<br>Anexos incorporados al BCE: ${RES.anexos_aplicados.length}.`:"";
    cont.innerHTML=`<div class="note-ok">&#10004; Conciliacion completa.${aj}${ax}</div>`+rep+dl;
    return;}
  const cls=t=>t==="Pago Directo"?"pd":t==="Diferencial Cambiario"?"dc":"ot";
  const difCarteras=TOTALES.filter(t=>t.estado==='DIFERENCIA').map(t=>t.acreedor);
  let h=`<div class="diaghead">Carteras no conciliadas<small>${difCarteras.length} cartera(s) con diferencia &middot; se identifica el credito exacto que no cuadra. Cada ajuste agregado desaparece de la lista.</small></div><div style="margin:0 0 12px;display:flex;gap:10px;flex-wrap:wrap"><button class="btn alt" id="btnAnexos" onclick="resolverConAnexos()">Corregir con anexos del mes</button><span class="muted" style="align-self:center;font-size:12px">Busca en Conciliaciones, Pagos y Desembolsos del periodo.</span></div>`;
  if(!difCarteras.length){ h+=`<div class="note-ok">No hay carteras con diferencia.</div>`; }
  else h+=difCarteras.map(ac=>{
    const d=DIAG.find(x=>x.acreedor===ac);
    const tipo=d?d.tipo:"Revisar"; const c=cls(tipo);
    const obs=d&&d.observacion?`<div class="dobs">&#128221; <b>Observacion MEF:</b> ${d.observacion}</div>`:"";
    // Prestamos que no cuadran (de todos los conceptos de esta cartera)
    let loansHtml="";
    Object.keys(PREST).filter(k=>k.split("|")[0]===ac).forEach(k=>{
      const concepto=k.split("|")[1];
      PREST[k].forEach((p,idx)=>{
        const id=(ac+concepto+p.credito).replace(/[^A-Za-z0-9]/g,"");
        const refc=p.referencia||(ac+"-"+p.credito);
        const esPD=(tipo==="Pago Directo"&&concepto==="Desembolsos");
        const notaDef=(concepto==="Desembolsos")
          ? `Considerar desembolso realizado a traves de la modalidad "pago directo", credito ${refc}`
          : `Ajuste ${concepto} - credito ${refc}`;
        const accion=esPD
          ? `<input type="file" id="lf${id}" accept=".xls,.xlsx,.pdf" style="display:none" onchange="document.getElementById('lfn${id}').textContent=this.files.length?this.files[0].name:''">
             <button class="btn gh" onclick="document.getElementById('lf${id}').click()">&#128206; Subir respaldo</button>
             <span class="muted" id="lfn${id}" style="font-size:11.5px"></span>
             <button class="btn" onclick="aplicarCredito('${ac}','${concepto}','${id}')">Agregar al BCE</button>`
          : `<button class="btn" onclick="aplicarCredito('${ac}','${concepto}','${id}')">Agregar al BCE</button>`;
        loansHtml+=`<div class="loan"><div class="lh"><b>${refc}</b> &middot; ${concepto}
          <span class="lm">MEF ${fmt(p.mef)} &nbsp;|&nbsp; BCE ${fmt(p.bce)} &nbsp;|&nbsp; falta <b style="color:var(--bad)">${fmt(p.dif)}</b></span></div>
          <div class="frow" style="margin-top:8px">
            <div class="fld"><label>Referencia</label><input type="text" id="lr${id}" value="${refc}" style="width:160px"></div>
            <div class="fld"><label>Valor a agregar (USD)</label><input type="number" id="lv${id}" value="${p.dif}" step="0.01" style="width:150px"></div>
            <div class="fld" style="flex:1;min-width:220px"><label>Nota / Observacion</label><input type="text" id="ln${id}" value="${notaDef.replace(/"/g,'&quot;')}" style="width:100%"></div>
            ${accion}
          </div></div>`;
      });
    });
    if(!loansHtml) loansHtml=`<div class="dobs muted">No pude desglosar el credito (revise el detalle del MEF/BCE).</div>`;
    return `<div class="dcard ${c}"><div class="dh"><span class="dname">${CARTERA_TXT[ac]||ac}</span><span class="dtag ${c}">${tipo}</span></div>${obs}${loansHtml}</div>`;
  }).join("");
  h+=`<div class="note-warn" style="margin-top:6px">La descarga del Reporte BCE modificado se habilita cuando TODAS las carteras esten conciliadas.</div>`;
  cont.innerHTML=h;
}
async function aplicarCredito(acreedor,concepto,id){
  const valor=parseFloat(document.getElementById("lv"+id).value)||0;
  let nota=document.getElementById("ln"+id).value.trim();
  const referencia=document.getElementById("lr"+id).value.trim();
  if(!valor){alert("Ingrese el valor a agregar.");return;}
  // Si hay respaldo adjunto (pago directo), subirlo y anexarlo a la nota
  const fin=document.getElementById("lf"+id);
  if(fin&&fin.files&&fin.files.length){
    const fd=new FormData();fd.append("archivos",fin.files[0]);
    try{await fetch("/api/pago_directo",{method:"POST",body:fd});}catch(e){}
    nota+="  (respaldo: "+fin.files[0].name+")";
  }
  const j=await (await fetch("/api/aplicar_pago_directo",{method:"POST",headers:{"Content-Type":"application/json"},
    body:JSON.stringify({periodo:PERIODO,acreedor,concepto,valor,nota,referencia})})).json();
  if(!j.ok){alert(j.error||"Error");return;}
  const mod=j.archivo_modificado;
  aplicarResultado({...j,info_periodo:INFO});
  alert((mod?"✓ Agregado al reporte del BCE ("+referencia+")":"⚠ Registrado, pero el archivo no se modifico: "+(j.archivo_motivo||"")));
}
function onDropPd(e){e.preventDefault();e.currentTarget.classList.remove("over");if(e.dataTransfer.files.length)subirPd(e.dataTransfer.files);}
async function subirPd(files){
  if(!PERIODO){alert("Ejecute primero la conciliacion.");return;}
  const fd=new FormData();[...files].forEach(f=>fd.append("archivos",f));
  const prev=document.getElementById("pdPrev");prev.innerHTML="<p class='muted'>Leyendo respaldo...</p>";
  const j=await (await fetch("/api/pago_directo",{method:"POST",body:fd})).json();
  if(!j.ok){prev.innerHTML="<p style='color:#f87171'>Error</p>";return;}
  prev.innerHTML=j.previos.map((p,i)=>{if(p.error)return `<div class="dcard"><b>${p.archivo}</b>: <span style="color:#f87171">${p.error}</span></div>`;
    const det=(p.detalle||[]).map(d=>`<div class="muted" style="font-size:11.5px">&middot; ${d.beneficiario}: ${fmt(d.monto)}</div>`).join("");
    return `<div class="dcard pd" id="pd${i}"><div class="dh"><span class="dname">${p.acreedor||'?'}</span><span class="dtag pd">Pago Directo</span><span class="dmonto" style="color:var(--gold2)">${p.prestamo||''}</span></div>${det}
      <div class="frow" style="margin-top:10px">
        <div class="fld"><label>Cartera</label><input type="text" id="pdac${i}" value="${p.acreedor||''}" style="width:120px"></div>
        <div class="fld"><label>Referencia</label><input type="text" id="pdref${i}" value="${p.referencia||''}" style="width:110px"></div>
        <div class="fld"><label>Valor USD</label><input type="number" id="pdval${i}" value="${p.valor||0}" step="0.01" style="width:150px"></div>
        <div class="fld" style="flex:1;min-width:230px"><label>Nota</label><input type="text" id="pdnota${i}" value="${(p.nota||'').replace(/"/g,'&quot;')}" style="width:100%"></div>
        <button class="btn" onclick="aplicarPd(${i})">Agregar al BCE</button></div></div>`;}).join("");
}
async function aplicarPd(i){
  const acreedor=document.getElementById("pdac"+i).value.trim(),valor=parseFloat(document.getElementById("pdval"+i).value)||0;
  const nota=document.getElementById("pdnota"+i).value.trim(),referencia=document.getElementById("pdref"+i).value.trim();
  if(!acreedor||!valor){alert("Falta cartera o valor");return;}
  const j=await (await fetch("/api/aplicar_pago_directo",{method:"POST",headers:{"Content-Type":"application/json"},
    body:JSON.stringify({periodo:PERIODO,acreedor,concepto:"Desembolsos",valor,nota,referencia})})).json();
  if(!j.ok){alert(j.error||"Error");return;}
  aplicarResultado({...j,info_periodo:INFO});mostrarDescargaBce();
  alert("Pago directo agregado a "+acreedor+" en el mismo archivo del BCE.");
}
function mostrarDescargaBce(){const el=document.getElementById("bceDl");if(!el)return;
  el.innerHTML=`<button class="btn alt" onclick="window.location='/api/descargar_bce_ajustado?periodo=${encodeURIComponent(PERIODO)}'">&#11015; Descargar reporte BCE modificado</button>`;}
function pintarQuipux(){
  const cont=document.getElementById("quipuxContenido");
  if(!(RES&&RES.conciliado_total)){const pend=DIAG.map(d=>CARTERA_TXT[d.acreedor]||d.acreedor).join(", ")||(RES?RES.diferencias+" concepto(s)":"");
    cont.innerHTML=`<div class="note-warn"><b>Oficio aún no disponible.</b><br>${RES?"La conciliación nueva no está completa"+(pend?": "+pend:"")+".":"Cargue y ejecute reportes nuevos en esta sesión."}</div>`;return;}
  cont.innerHTML=`<div class="card">
    <div class="frow"><div class="fld"><label>Nro. Oficio BCE</label><input type="text" id="qOfBce" value="BCE-SSFI-${INFO?INFO.anio:''}-XXXX-OF" style="width:200px"></div>
      <div class="fld"><label>Responde Oficio MEF</label><input type="text" id="qOfMef" value="MEF-DNS-${INFO?INFO.anio:''}-XXXX-O" style="width:200px"></div>
      <div class="fld"><label>Fecha Oficio MEF</label><input type="text" id="qFechaMef" value="" style="width:170px"></div></div>
    <div class="frow" style="margin-top:12px"><div class="fld" style="flex:1;min-width:240px"><label>Destinatario</label><input type="text" id="qDest" value="Ana Maria Vallejo Cabezas - Directora Nacional de Seguimiento" style="width:100%"></div>
      <div class="fld" style="flex:1;min-width:240px"><label>Firmante (BCE)</label><input type="text" id="qFirma" value="Mgs. Luis Santiago Vargas Bautista - Subgerente de Servicios Financieros Internacionales" style="width:100%"></div></div>
    <div class="fld" style="margin-top:12px"><label>Copia (una persona por linea: Nombre - Cargo)</label>
      <textarea id="qCopia" style="height:150px">Miguel Rodrigo Hernandez Cobos - Subsecretario de Financiamiento Publico y Analisis de Riesgos
Katherine de los Angeles Castellanos Cevallos - Analista 2 de Negociacion y Financiamiento Publico
Edith Vinueza Herrera - Analista 2 de Seguimiento y Evaluacion del Financiamiento Publico
Yessenia Centeno Masache - Especialista de Financiamiento Publico y Analisis de Riesgos
Jose Patricio Egas Ponce - Analista 1 de Seguimiento y Evaluacion del Financiamiento Publico
Rodolfo Ehmig Santillan - Analista 2 de la Caja Fiscal
Diana Carolina Valdivieso Velasco - Especialista de Sistemas de Pago 2</textarea></div>
    <div style="margin-top:14px;display:flex;gap:10px;flex-wrap:wrap"><button class="btn" onclick="generarQuipux()">Generar oficio</button>
      <button class="btn gh" onclick="copiarQuipux()">Copiar</button><button class="btn gh" onclick="descargarQuipux()">Descargar .txt</button></div>
    <textarea id="qText" style="margin-top:14px" placeholder="Pulse 'Generar oficio'."></textarea></div>`;
}
function unirRubros(a){if(!a.length)return"";if(a.length===1)return a[0];const u=a[a.length-1];const c=/^h?i/i.test(u)?"e":"y";return a.slice(0,-1).join(", ")+" "+c+" "+u;}
function lineaCartera(cart){const rows=DATA.filter(r=>r.acreedor===cart);const nom=CARTERA_TXT[cart]||cart;
  if(!rows.length)return `${nom}: No existen movimientos.`;
  const con=rows.filter(r=>r.estado==="CONCILIADO").map(r=>RUBRO_TXT[r.concepto]||r.concepto.toLowerCase());
  return `${nom}: No existen observaciones en los rubros de ${unirRubros(con)}.`;}
function qv(id){const el=document.getElementById(id);return el?el.value.trim():"";}
function generarQuipux(){if(!(RES&&RES.conciliado_total)){alert("La conciliacion no esta completa.");return;}
  const present=new Set(DATA.map(r=>r.acreedor));const orden=ORDEN_QUIPUX.concat([...present].filter(a=>!ORDEN_QUIPUX.includes(a)));
  const lineas=orden.filter(c=>present.has(c)||["AIIB","GPS"].includes(c)).map(lineaCartera);
  const hoy=new Date().toLocaleDateString("es-EC",{day:"numeric",month:"long",year:"numeric"});const fMef=qv("qFechaMef");
  const copiaRaw=(document.getElementById("qCopia")?document.getElementById("qCopia").value:"").split("\n").map(x=>x.trim()).filter(Boolean);
  const copia=copiaRaw.length?"\n\nCopia:\n"+copiaRaw.map(l=>{const p=l.split(" - ");
    return p.length>1?`Senor(a)\n${p[0].trim()}\n${p.slice(1).join(" - ").trim()}\nMINISTERIO DE ECONOMIA Y FINANZAS\n`:l;}).join("\n"):"";
  const txt=`Oficio Nro. ${qv("qOfBce")}
Quito, D.M., ${hoy}

Asunto: Conciliacion mensual de valores del Servicio de la Deuda a ${INFO?INFO.texto:PERIODO}

Senora Magister
${qv("qDest")}
MINISTERIO DE ECONOMIA Y FINANZAS
En su Despacho

De mi consideracion:

Me refiero al Oficio Nro. ${qv("qOfMef")}${fMef?` de ${fMef}`:""}, mediante el cual su Despacho solicito la ratificacion o rectificacion de la informacion sobre los movimientos de la deuda externa publica correspondiente al mes de ${(INFO?INFO.mes:"").toLowerCase()} de ${INFO?INFO.anio:""}.

Al respecto, me permito indicar lo siguiente:

${lineas.join("\n")}

Por lo expuesto, me permito remitir para su revision la matriz de conciliacion del Banco Central del Ecuador en formato Excel. Este archivo contiene la informacion completa y necesaria para el proceso de conciliacion.

Con sentimientos de distinguida consideracion.

Atentamente,

Documento firmado electronicamente
${qv("qFirma")}${copia}`;
  document.getElementById("qText").value=txt;}
function copiarQuipux(){const t=document.getElementById("qText");if(!t.value){alert("Genere el oficio primero.");return;}t.select();document.execCommand("copy");alert("Texto copiado.");}
function descargarQuipux(){const t=document.getElementById("qText").value;if(!t){alert("Genere el oficio primero.");return;}
  const b=new Blob([t],{type:"text/plain;charset=utf-8"}),u=URL.createObjectURL(b),a=document.createElement("a");a.href=u;a.download="Oficio_Respuesta_"+PERIODO+".txt";a.click();URL.revokeObjectURL(u);}

const CAF_MONTHS={Enero:"Ene",Febrero:"Feb",Marzo:"Mar",Abril:"Abr",Mayo:"May",Junio:"Jun",Julio:"Jul",Agosto:"Ago",Septiembre:"Sep",Octubre:"Oct",Noviembre:"Nov",Diciembre:"Dic"};
let CAF_ITEMS=[];
function cafEl(id){return document.getElementById(id);}
function cafFmt(n){return Number(n||0).toLocaleString("en-US",{minimumFractionDigits:2,maximumFractionDigits:2});}
function cafParse(value){const text=String(value??"").trim();if(!text)return 0;const neg=text.startsWith("-")||(text.startsWith("(")&&text.endsWith(")"));let c=text.replace(/[()$\s]/g,"");if(c.includes(",")&&c.includes(".")){c=c.lastIndexOf(",")>c.lastIndexOf(".")?c.replace(/\./g,"").replace(",","."):c.replace(/,/g,"");}else if(c.includes(",")){c=c.replace(/\./g,"").replace(",",".");}const num=Number(c||0);return Number((neg?-num:num).toFixed(2));}
function cafLog(t){const el=cafEl("cafLog");if(!el)return;const h=new Date().toLocaleTimeString();el.textContent+=`[${h}] ${t}\n`;el.scrollTop=el.scrollHeight;}
function cafSetStatus(t){const el=cafEl("cafEstado");if(el)el.textContent=t;}
function cafBasePath(y){return `Z:\\GISI\\SSFI\\GESTI\u00d3N PAGOS INTERNACIONALES\\${y}\\DEUDA EXTERNA P\u00daBLICA\\Acreedores Internacionales\\CAF\\Pagos`;}
function cafReportPath(y){return `Z:\\GISI\\SSFI\\GESTI\u00d3N PAGOS INTERNACIONALES\\${y}\\DEUDA EXTERNA P\u00daBLICA\\Conciliaciones\\Reporte\\Reporte Conciliaci\u00f3n BCE - ${y}.xlsx`;}
function cafSyncPaths(){if(!cafEl("cafAuto")||!cafEl("cafAuto").checked)return;const y=Number(cafEl("cafAnio").value||new Date().getFullYear());const m=cafEl("cafMes").value||"Enero";cafEl("cafBase").value=cafBasePath(y);cafEl("cafDir").value=`${cafEl("cafBase").value}\\${m}`;cafEl("cafReporte").value=cafReportPath(y);cafEl("cafHoja").value=`Giros al Exterior - ${(CAF_MONTHS[m]||m.slice(0,3))}${y}`;}
function cafCfg(){cafSyncPaths();return{anio:Number(cafEl("cafAnio").value),mes:cafEl("cafMes").value,auto_paths:cafEl("cafAuto").checked,pagos_caf_base:cafEl("cafBase").value.trim(),pagos_caf_dir:cafEl("cafDir").value.trim(),reporte_path:cafEl("cafReporte").value.trim(),sheet_name:cafEl("cafHoja").value.trim(),crear_backup:cafEl("cafBackup").checked};}
async function cafApi(path,payload){const opts=payload?{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)}:{};const r=await fetch(path,opts);const j=await r.json();if(!r.ok||!j.ok)throw new Error(j.error||"Operacion no completada.");return j;}
async function cafLoadConfig(){if(!cafEl("cafMes"))return;const data=await cafApi("/api/caf_condonados/config");const meses=(data.config.meses||[]).map(x=>x.nombre);cafEl("cafMes").innerHTML=(meses.length?meses:Object.keys(CAF_MONTHS)).map(m=>`<option value="${m}">${m}</option>`).join("");const map={anio:"cafAnio",mes:"cafMes",auto_paths:"cafAuto",crear_backup:"cafBackup",pagos_caf_base:"cafBase",pagos_caf_dir:"cafDir",reporte_path:"cafReporte",sheet_name:"cafHoja"};Object.entries(map).forEach(([k,id])=>{const el=cafEl(id);if(!el)return;if(el.type==="checkbox")el.checked=!!data.config[k];else el.value=data.config[k]??"";});cafSyncPaths();cafLog("Configuracion cargada.");}
async function cafSaveConfig(){cafSetStatus("Guardando");const data=await cafApi("/api/caf_condonados/config",cafCfg());const map={anio:"cafAnio",mes:"cafMes",auto_paths:"cafAuto",crear_backup:"cafBackup",pagos_caf_base:"cafBase",pagos_caf_dir:"cafDir",reporte_path:"cafReporte",sheet_name:"cafHoja"};Object.entries(map).forEach(([k,id])=>{const el=cafEl(id);if(!el)return;if(el.type==="checkbox")el.checked=!!data.config[k];else el.value=data.config[k]??"";});cafSyncPaths();cafSetStatus("Guardado");cafLog("Configuracion guardada.");}
function cafTotals(){const total=CAF_ITEMS.reduce((s,it)=>s+cafParse(it.valor),0);cafEl("cafTotalCarpetas").textContent=CAF_ITEMS.length;cafEl("cafTotalValor").textContent=cafFmt(total);cafEl("cafAplicar").disabled=!CAF_ITEMS.length;}
function cafRender(){const body=cafEl("cafRows");if(!CAF_ITEMS.length){body.innerHTML='<tr><td colspan="6" class="muted">Sin datos cargados.</td></tr>';cafTotals();return;}body.innerHTML="";CAF_ITEMS.forEach((it,i)=>{const tr=document.createElement("tr");tr.innerHTML=`<td><strong>${it.clave||"Sin referencia"}</strong></td><td>${it.sigade||""}</td><td><code>${it.carpeta||""}</code></td><td><input class="caf-amount" type="text" value="${cafFmt(it.valor)}" data-i="${i}"></td><td><span class="source-pill">${it.fuente||"manual"}</span></td><td class="${it.existe_estado_cuenta?"ok-text":"warn-text"}">${it.existe_estado_cuenta?"Encontrado":"No encontrado"}</td>`;body.appendChild(tr);});document.querySelectorAll(".caf-amount").forEach(input=>{input.addEventListener("input",()=>{CAF_ITEMS[Number(input.dataset.i)].valor=cafParse(input.value);cafTotals();});input.addEventListener("blur",()=>{input.value=cafFmt(cafParse(input.value));});});cafTotals();}
async function cafScan(){try{cafSetStatus("Escaneando");cafLog("Escaneando carpetas CAF.");const data=await cafApi("/api/caf_condonados/scan",cafCfg());CAF_ITEMS=data.items||[];cafRender();cafSetStatus(CAF_ITEMS.length?"Valores listos":"Sin datos");cafLog(`Escaneo finalizado: ${data.total} carpeta(s).`);}catch(e){cafSetStatus("Error");cafLog("Error: "+e.message);alert(e.message);}}
async function cafProcess(){if(!CAF_ITEMS.length){alert("Escanee las carpetas primero.");return;}if(!confirm("Desea actualizar el reporte de conciliacion BCE ahora?"))return;try{cafSetStatus("Actualizando");cafLog("Aplicando condonados al reporte BCE.");const data=await cafApi("/api/caf_condonados/process",{...cafCfg(),items:CAF_ITEMS});cafSetStatus("Completado");cafLog(`Reporte actualizado: ${data.reporte_path}`);cafLog(`Creditos aplicados: ${data.matched.length}; total: ${cafFmt(data.total)}`);if(data.backup_path)cafLog(`Respaldo: ${data.backup_path}`);if(data.missing&&data.missing.length)cafLog(`No encontrados en la hoja: ${data.missing.join(", ")}`);alert("Reporte actualizado correctamente.");}catch(e){cafSetStatus("Error");cafLog("Error: "+e.message);alert(e.message);}}
function cafInit(){if(!cafEl("cafMes"))return;cafEl("cafGuardar").addEventListener("click",()=>cafSaveConfig().catch(e=>alert(e.message)));cafEl("cafEscanear").addEventListener("click",cafScan);cafEl("cafAplicar").addEventListener("click",cafProcess);["cafAnio","cafMes","cafAuto"].forEach(id=>cafEl(id).addEventListener("change",cafSyncPaths));cafLoadConfig().catch(e=>{cafSetStatus("Error");cafLog("Error cargando configuracion: "+e.message);});}
cafInit();

const AGENDA_UI={items:[],summary:{},monitor:{},errors:[],scan:null,timer:null,audio:null};
const AGENDA_ALERTS_KEY="gidep-agenda-alertas";
const AGENDA_SEEN_KEY="gidep-agenda-pagos-vistos";
function agendaEl(id){return document.getElementById(id);}
function agendaEsc(value){return String(value??"").replace(/[&<>"']/g,ch=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[ch]));}
function agendaFmtDate(value){if(!value)return "No indicada";const parts=String(value).split("-").map(Number);if(parts.length!==3||!parts.every(Number.isFinite))return value;return new Intl.DateTimeFormat("es-EC",{day:"2-digit",month:"2-digit",year:"numeric"}).format(new Date(parts[0],parts[1]-1,parts[2]));}
function agendaFmtDateTime(value){if(!value)return "Pendiente";const parsed=new Date(value);return Number.isNaN(parsed.getTime())?value:new Intl.DateTimeFormat("es-EC",{dateStyle:"short",timeStyle:"short"}).format(parsed);}
function agendaFmtAmount(value){if(value===null||value===undefined||value==="")return "No indicado";if(typeof value==="number")return new Intl.NumberFormat("en-US",{minimumFractionDigits:2,maximumFractionDigits:2}).format(value);const raw=String(value).trim();let normalized=raw;if(raw.includes(".")&&raw.includes(","))normalized=raw.lastIndexOf(".")>raw.lastIndexOf(",")?raw.replaceAll(",",""):raw.replaceAll(".","").replace(",",".");else if(/^\d{1,3}(?:\.\d{3})+,\d{2}$/.test(raw))normalized=raw.replaceAll(".","").replace(",",".");else if(/^\d+,\d{2}$/.test(raw))normalized=raw.replace(",",".");const number=Number(normalized);return Number.isFinite(number)?new Intl.NumberFormat("en-US",{minimumFractionDigits:2,maximumFractionDigits:2}).format(number):raw;}
function agendaMatchesStatus(item,filter){if(filter==="all")return true;if(filter==="pending")return item.status!=="paid";if(filter==="upcoming")return item.days_until>0&&item.status!=="paid";if(filter==="today")return item.days_until===0;return item.status===filter;}
function agendaFiltered(){const status=agendaEl("agendaStatusFilter").value;const creditor=agendaEl("agendaCreditorFilter").value;const query=agendaEl("agendaSearchFilter").value.trim().toLocaleUpperCase("es");return AGENDA_UI.items.filter(item=>{if(!agendaMatchesStatus(item,status))return false;if(creditor!=="all"&&item.responsible!==creditor)return false;if(!query)return true;return [item.loan_reference,item.office_reference,item.voucher_number,item.responsible,item.issuer,item.loan_folder,item.correspondent,item.beneficiary_bank].filter(Boolean).join(" ").toLocaleUpperCase("es").includes(query);});}
function agendaPopulateCreditors(){const select=agendaEl("agendaCreditorFilter");const selected=select.value;const creditors=[...new Set(AGENDA_UI.items.map(item=>item.responsible).filter(Boolean))].sort((a,b)=>a.localeCompare(b,"es"));select.innerHTML='<option value="all">Todos</option>'+creditors.map(value=>`<option value="${agendaEsc(value)}">${agendaEsc(value)}</option>`).join("");select.value=creditors.includes(selected)?selected:"all";}
function agendaRenderRows(){const body=agendaEl("agendaRows");const items=agendaFiltered();agendaEl("agendaResultDescription").textContent=`${items.length} de ${AGENDA_UI.items.length} pago(s) identificados en la ventana de 14 días.`;if(!items.length){body.innerHTML='<tr><td colspan="8" class="muted">No hay pagos que coincidan con los filtros seleccionados.</td></tr>';return;}const labels={today:"Pagar hoy",scheduled:"Programado",in_process:"En proceso",overdue:"Vencido",paid:"Procesado",review:"Revisar"};body.innerHTML=items.map(item=>`<tr><td><span class="agenda-loan">${agendaEsc(item.loan_reference||"Sin referencia")}</span><span class="agenda-cell-sub">${agendaEsc(item.loan_folder||"")}</span>${item.voucher_number&&item.voucher_number!=="—"?`<span class="agenda-cell-sub">Comprobante ${agendaEsc(item.voucher_number)}</span>`:""}</td><td><span class="agenda-date-main">${agendaEsc(agendaFmtDate(item.value_date))}</span><span class="agenda-cell-sub">${agendaEsc(item.date_type||"Fecha valor")}</span></td><td><span class="agenda-date-main">${agendaEsc(agendaFmtDate(item.due_date))}</span></td><td class="agenda-amount">${agendaEsc(agendaFmtAmount(item.amount_value))}</td><td><span class="agenda-currency">${agendaEsc(item.currency||"No indicada")}</span></td><td><span class="agenda-correspondent ${item.correspondent==="REVISAR"?"review":""}">${agendaEsc(item.correspondent||"REVISAR")}</span><span class="agenda-cell-sub">${agendaEsc(item.correspondent_method||"Sin método seguro")} · ${agendaEsc(item.correspondent_confidence||0)}%</span>${item.beneficiary_bank?`<span class="agenda-cell-sub">Banco beneficiario: ${agendaEsc(item.beneficiary_bank)}</span>`:""}<span class="agenda-cell-sub">Respaldos bancarios revisados: ${agendaEsc(item.account_statements_reviewed||0)}</span></td><td><span class="agenda-loan">${agendaEsc(item.responsible||item.issuer||"Institución")}</span><span class="agenda-cell-sub">${agendaEsc(item.office_reference||item.document_name||"")}</span><span class="agenda-cell-sub">${agendaEsc(item.source_document||"")}</span></td><td><span class="agenda-status ${agendaEsc(item.status)}">${agendaEsc(labels[item.status]||item.status)}</span><span class="agenda-cell-sub">${agendaEsc(item.note||"")}</span><span class="agenda-cell-sub">Confianza del pago ${agendaEsc(item.confidence||0)}%${item.ocr_used?" · OCR":""}</span></td></tr>`).join("");}
function agendaRenderTotals(totals){const target=agendaEl("agendaTodayTotals");const entries=Object.entries(totals||{});target.innerHTML=entries.length?entries.map(([currency,value])=>`<span class="agenda-total-pill"><strong>${agendaEsc(currency)}</strong>${agendaEsc(agendaFmtAmount(value))}</span>`).join(""):'<span class="agenda-total-empty">Sin valores pendientes identificados.</span>';}
function agendaRenderSummary(){const summary=AGENDA_UI.summary||{};agendaEl("agendaToday").textContent=summary.today||0;agendaEl("agendaUpcoming").textContent=summary.next_7_days||0;agendaEl("agendaOverdue").textContent=summary.overdue||0;agendaEl("agendaPaid").textContent=summary.paid||0;agendaEl("agendaReview").textContent=summary.review||0;const pending=summary.today_pending||0;const badge=agendaEl("agendaNavBadge");badge.textContent=pending;badge.classList.toggle("hidden",pending===0);badge.classList.toggle("pulse",pending>0);agendaEl("agendaTodayHeadline").textContent=pending?`${pending} pago(s) pendientes para hoy`:(summary.today||0)?`${summary.today} pago(s) de hoy ya procesados`:"No hay pagos programados para hoy";agendaEl("agendaTodayDetail").textContent=pending?"Revise préstamo, oficio, monto y moneda antes de ingresar el pago.":"El monitor continuará revisando nuevos oficios cada cinco minutos.";agendaRenderTotals(summary.today_pending_totals||{});if(AGENDA_UI.scan?.today){const parts=AGENDA_UI.scan.today.split("-").map(Number);const current=new Date(parts[0],parts[1]-1,parts[2]);agendaEl("agendaTodayDay").textContent=String(parts[2]).padStart(2,"0");agendaEl("agendaTodayMonth").textContent=new Intl.DateTimeFormat("es-EC",{month:"long",year:"numeric"}).format(current);}}
function agendaRenderMonitor(){const monitor=AGENDA_UI.monitor||{};const notice=agendaEl("agendaNotice");notice.className=`agenda-notice ${monitor.last_error?"error":monitor.busy?"busy":""}`.trim();agendaEl("agendaMonitorTitle").textContent=monitor.last_error?"La última lectura tuvo observaciones":monitor.busy?"Analizando los oficios":"Agenda automática activa";agendaEl("agendaMonitorDetail").textContent=monitor.last_error||monitor.message||"Revisión automática cada cinco minutos.";agendaEl("agendaNextRun").textContent=monitor.next_run?`Próxima: ${agendaFmtDateTime(monitor.next_run)}`:"Cada 5 minutos";agendaEl("agendaLastRun").textContent=`Última actualización: ${agendaFmtDateTime(monitor.last_run)}`;agendaEl("agendaRefresh").disabled=!!monitor.busy;}
function agendaRenderErrors(){const errors=[...(AGENDA_UI.errors||[])];if(AGENDA_UI.monitor?.last_error)errors.unshift(AGENDA_UI.monitor.last_error);const box=agendaEl("agendaError");box.classList.toggle("hidden",errors.length===0);box.textContent=errors.length?`Observaciones de lectura: ${errors.join(" | ")}`:"";}
function agendaSeen(){try{return new Set(JSON.parse(localStorage.getItem(AGENDA_SEEN_KEY)||"[]"));}catch(_){return new Set();}}
function agendaSaveSeen(seen){try{localStorage.setItem(AGENDA_SEEN_KEY,JSON.stringify([...seen].slice(-300)));}catch(_){}}
function agendaBeep(){try{const AudioCtor=window.AudioContext||window.webkitAudioContext;if(!AudioCtor)return;AGENDA_UI.audio=AGENDA_UI.audio||new AudioCtor();const oscillator=AGENDA_UI.audio.createOscillator();const gain=AGENDA_UI.audio.createGain();oscillator.frequency.value=740;gain.gain.setValueAtTime(.0001,AGENDA_UI.audio.currentTime);gain.gain.exponentialRampToValueAtTime(.13,AGENDA_UI.audio.currentTime+.02);gain.gain.exponentialRampToValueAtTime(.0001,AGENDA_UI.audio.currentTime+.42);oscillator.connect(gain);gain.connect(AGENDA_UI.audio.destination);oscillator.start();oscillator.stop(AGENDA_UI.audio.currentTime+.43);}catch(_){}}
function agendaAnnounce(force=false){if(localStorage.getItem(AGENDA_ALERTS_KEY)!=="true")return;const today=AGENDA_UI.items.filter(item=>item.days_until===0&&!item.completed);const seen=agendaSeen();const fresh=force?today:today.filter(item=>!seen.has(item.payment_id));if(!fresh.length)return;fresh.forEach(item=>seen.add(item.payment_id));agendaSaveSeen(seen);agendaBeep();const first=fresh[0];const extra=fresh.length>1?` y ${fresh.length-1} más`:"";const message=`${first.loan_reference||"Préstamo sin referencia"} · ${first.currency||""} ${agendaFmtAmount(first.amount_value)}${extra}`;if("Notification" in window&&Notification.permission==="granted"){try{new Notification("GIDEP BCE · Pago para hoy",{body:message,tag:"gidep-agenda-hoy"});}catch(_){}}}
function agendaApply(payload){AGENDA_UI.monitor=payload.monitor||{};if(payload.scan){AGENDA_UI.scan=payload.scan;AGENDA_UI.items=payload.scan.items||[];AGENDA_UI.summary=payload.scan.summary||{};AGENDA_UI.errors=payload.scan.errors||[];agendaEl("agendaSource").textContent=payload.scan.source||"-";agendaPopulateCreditors();agendaRenderSummary();agendaRenderRows();agendaAnnounce();}agendaRenderMonitor();agendaRenderErrors();}
async function agendaLoad(){try{const response=await fetch("/api/agenda-pagos/status",{cache:"no-store"});const payload=await response.json();if(!response.ok||!payload.ok)throw new Error(payload.error||`HTTP ${response.status}`);agendaApply(payload);}catch(error){AGENDA_UI.monitor={last_error:`No se pudo consultar la agenda: ${error.message}`};agendaRenderMonitor();agendaRenderErrors();}}
async function agendaRefresh(){try{agendaEl("agendaRefresh").disabled=true;const response=await fetch("/api/agenda-pagos/refresh",{method:"POST"});const payload=await response.json();if(!response.ok&&!payload.ok)throw new Error(payload.error||`HTTP ${response.status}`);agendaApply(payload);window.setTimeout(agendaLoad,2500);}catch(error){alert(error.message);agendaEl("agendaRefresh").disabled=false;}}
async function agendaOpenSource(){try{const response=await fetch("/api/agenda-pagos/open-source",{method:"POST"});const payload=await response.json();if(!response.ok||!payload.ok)throw new Error(payload.error||`HTTP ${response.status}`);}catch(error){alert(error.message);}}
async function agendaEnableAlerts(){try{localStorage.setItem(AGENDA_ALERTS_KEY,"true");const AudioCtor=window.AudioContext||window.webkitAudioContext;if(AudioCtor){AGENDA_UI.audio=AGENDA_UI.audio||new AudioCtor();await AGENDA_UI.audio.resume();}if("Notification" in window&&Notification.permission==="default")await Notification.requestPermission();agendaEl("agendaEnableAlerts").textContent="Alertas activas";agendaAnnounce(true);}catch(error){alert(`No se pudieron activar las alertas: ${error.message}`);}}
function agendaInit(){if(!agendaEl("agendaRows"))return;agendaEl("agendaStatusFilter").addEventListener("change",agendaRenderRows);agendaEl("agendaCreditorFilter").addEventListener("change",agendaRenderRows);agendaEl("agendaSearchFilter").addEventListener("input",agendaRenderRows);agendaEl("agendaRefresh").addEventListener("click",agendaRefresh);agendaEl("agendaOpenSource").addEventListener("click",agendaOpenSource);agendaEl("agendaEnableAlerts").addEventListener("click",agendaEnableAlerts);if(localStorage.getItem(AGENDA_ALERTS_KEY)==="true")agendaEl("agendaEnableAlerts").textContent="Alertas activas";agendaLoad();AGENDA_UI.timer=window.setInterval(agendaLoad,30000);}
agendaInit();

const MAIL={examples:[],initialized:false,selectedExample:null,originalSuggestion:""};
const mailEl=id=>document.getElementById(id);
const mailEsc=value=>String(value??"").replace(/[&<>"']/g,char=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#039;"}[char]));
async function mailApi(path,payload){const options=payload?{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)}:{};const response=await fetch(path,options);const data=await response.json();if(!response.ok||!data.ok)throw new Error(data.error||`HTTP ${response.status}`);return data;}
function mailSetFlow(flow){document.querySelectorAll("#mailFlowTabs [data-mail-flow]").forEach(button=>{const active=button.dataset.mailFlow===flow;button.classList.toggle("act",active);button.setAttribute("aria-selected",String(active));});document.querySelectorAll("[data-mail-panel]").forEach(panel=>panel.classList.toggle("act",panel.dataset.mailPanel===flow));if(flow==="library")mailLoadExamples();}
function mailSetStatus(id,text,tone=""){const element=mailEl(id);element.className=`mail-status ${tone}`.trim();element.textContent=text;}
function mailRenderLibrary(){mailEl("mailExampleCount").textContent=`${MAIL.examples.length} ejemplo(s)`;const body=mailEl("mailLibraryRows");if(!MAIL.examples.length){body.innerHTML='<tr><td colspan="6" class="muted">Aún no hay ejemplos guardados. Use “Enseñar respuesta” para registrar el primero.</td></tr>';return;}body.innerHTML=MAIL.examples.map(item=>`<tr><td>${mailEsc(new Date(item.fecha_creacion).toLocaleDateString("es-EC"))}</td><td><span class="source-pill">${mailEsc(item.categoria)}</span></td><td>${mailEsc(item.asunto||"Sin asunto")}</td><td><span class="mail-preview" title="${mailEsc(item.correo_recibido)}">${mailEsc(item.correo_recibido)}</span></td><td><span class="mail-preview" title="${mailEsc(item.respuesta_aprobada)}">${mailEsc(item.respuesta_aprobada)}</span></td><td>${mailEsc(item.usos||0)}</td></tr>`).join("");}
async function mailLoadExamples(){try{const data=await mailApi("/api/respuestas-correos/examples");MAIL.examples=data.examples||[];mailRenderLibrary();}catch(error){mailSetStatus("mailLearnStatus",`No se pudo leer la base local: ${error.message}`,"warn");}}
async function mailSuggest(){const incoming=mailEl("mailAnalyzeIncoming").value.trim();if(!incoming){mailSetStatus("mailAnalyzeStatus","Ingrese el contenido del correo antes de analizar.","warn");return;}mailEl("mailSuggest").disabled=true;mailSetStatus("mailAnalyzeStatus","Buscando respuestas aprobadas similares...");try{const data=await mailApi("/api/respuestas-correos/suggest",{category:mailEl("mailAnalyzeCategory").value,subject:mailEl("mailAnalyzeSubject").value,incoming});mailEl("mailSuggestedResponse").value=data.suggestion||"";MAIL.selectedExample=data.matches?.[0]?.id||null;MAIL.originalSuggestion=data.suggestion||"";mailEl("mailConfidence").textContent=data.matches?.length?`Confianza controlada ${data.confidence}%`:"Sin ejemplos";mailSetStatus("mailAnalyzeStatus",data.message,data.matches?.length?(data.confidence>=70?"ok":"warn"):"warn");mailEl("mailMatches").innerHTML=(data.matches||[]).map(item=>`<article class="mail-match"><strong>${mailEsc(item.categoria)} · ${mailEsc(item.confidence)}% de similitud</strong><small>${mailEsc(item.asunto||"Sin asunto")}</small><small>${mailEsc(item.observaciones||"Respuesta aprobada guardada localmente")}</small></article>`).join("");await mailLoadExamples();}catch(error){mailSetStatus("mailAnalyzeStatus",error.message,"warn");}finally{mailEl("mailSuggest").disabled=false;}}
function mailPrepareLearning(){mailEl("mailLearnCategory").value=mailEl("mailAnalyzeCategory").value;mailEl("mailLearnSubject").value=mailEl("mailAnalyzeSubject").value;mailEl("mailLearnIncoming").value=mailEl("mailAnalyzeIncoming").value;mailEl("mailLearnResponse").value=mailEl("mailSuggestedResponse").value;mailSetFlow("learn");mailSetStatus("mailLearnStatus","Revise el ejemplo y guárdelo únicamente si la respuesta es correcta.");}
async function mailSaveExample(){const payload={category:mailEl("mailLearnCategory").value,subject:mailEl("mailLearnSubject").value,incoming:mailEl("mailLearnIncoming").value,response:mailEl("mailLearnResponse").value,notes:mailEl("mailLearnNotes").value};mailEl("mailSaveExample").disabled=true;try{const data=await mailApi("/api/respuestas-correos/learn",payload);mailSetStatus("mailLearnStatus",data.message,"ok");await mailLoadExamples();}catch(error){mailSetStatus("mailLearnStatus",error.message,"warn");}finally{mailEl("mailSaveExample").disabled=false;}}
async function mailCopySuggestion(){const response=mailEl("mailSuggestedResponse").value.trim();if(!response){mailSetStatus("mailAnalyzeStatus","No hay una respuesta para copiar.","warn");return;}try{await navigator.clipboard.writeText(response);if(MAIL.selectedExample){const outcome=response===MAIL.originalSuggestion?"accepted":"edited";await mailApi("/api/respuestas-correos/feedback",{example_id:MAIL.selectedExample,outcome});}mailSetStatus("mailAnalyzeStatus","Respuesta copiada y resultado registrado. Revísela antes de enviarla.","ok");}catch(error){mailSetStatus("mailAnalyzeStatus","No se pudo copiar automáticamente; seleccione el texto manualmente.","warn");}}
function mailResetAnalysis(){mailSetFlow("analyze");MAIL.selectedExample=null;MAIL.originalSuggestion="";mailEl("mailAnalyzeCategory").selectedIndex=0;mailEl("mailAnalyzeSubject").value="";mailEl("mailAnalyzeIncoming").value="";mailEl("mailSuggestedResponse").value="";mailEl("mailMatches").innerHTML="";mailEl("mailConfidence").textContent="Sin análisis";mailSetStatus("mailAnalyzeStatus","Ingrese un correo para buscar respuestas aprobadas similares.");mailEl("mailAnalyzeSubject").focus();}
async function activationInit(){try{const data=await mailApi("/api/activaciones-cuentas/status");mailEl("activationStatus").textContent=data.available?"Disponible e integrado en el gestor":"No se encontró el archivo original";}catch(error){mailEl("activationStatus").textContent=`No disponible: ${error.message}`;}}
function mailInit(){if(MAIL.initialized||!mailEl("mailFlowTabs"))return;MAIL.initialized=true;document.querySelectorAll("#mailFlowTabs [data-mail-flow]").forEach(button=>button.addEventListener("click",()=>mailSetFlow(button.dataset.mailFlow)));mailEl("mailSuggest").addEventListener("click",mailSuggest);mailEl("mailReset").addEventListener("click",mailResetAnalysis);mailEl("mailLearnSuggestion").addEventListener("click",mailPrepareLearning);mailEl("mailSaveExample").addEventListener("click",mailSaveExample);mailEl("mailCopySuggestion").addEventListener("click",mailCopySuggestion);mailLoadExamples();activationInit();}
mailInit();

const FIRM={flows:[],flowId:"deuda_externa",messages:[],alerts:[],summary:{},target:"",timer:null,audio:null,initialized:false,lastData:null,loading:false};
const FIRM_ALERTS_KEY="automatiza-traslado-alertas-firma";
const FIRM_SEEN_KEY="automatiza-traslado-firmas-vistas";
function firmEl(id){return document.getElementById(id);}
function firmEsc(value){return String(value??"").replace(/[&<>"']/g,ch=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[ch]));}
function firmLog(text){const el=firmEl("firmLog");if(!el)return;const hour=new Date().toLocaleTimeString();el.textContent+=`[${hour}] ${text}\n`;el.scrollTop=el.scrollHeight;}
function firmFmtBytes(value){if(value<1024)return `${value} B`;if(value<1048576)return `${(value/1024).toFixed(1)} KB`;return `${(value/1048576).toFixed(2)} MB`;}
function firmFmtDate(value){const parsed=new Date(value);return Number.isNaN(parsed.getTime())?value:new Intl.DateTimeFormat("es-EC",{dateStyle:"short",timeStyle:"short"}).format(parsed);}
async function firmApi(path,payload){const opts=payload?{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)}:{cache:"no-store"};const response=await fetch(path,opts);const data=await response.json();if(!response.ok||!data.ok)throw new Error(data.error||`HTTP ${response.status}`);return data;}
function firmSetAlert(text,mode=""){const el=firmEl("firmAlert");el.textContent=text;el.className=`firm-alert ${mode}`.trim();}
function firmSelected(){return [...document.querySelectorAll("#firmRows input[type=checkbox]:checked")].map(input=>input.value);}
function firmUpdateActions(){const count=firmSelected().length;const button=firmEl("firmMove");button.disabled=count===0;button.textContent=count?`Copiar seleccionados (${count})`:"Copiar seleccionados";}
function firmRenderSummary(data){const summary=data.summary||{};FIRM.summary=summary;FIRM.target=data.target||"";firmEl("firmPending").textContent=summary.pending_signatures||0;firmEl("firmDetected").textContent=summary.detected||0;firmEl("firmReady").textContent=summary.ready||0;firmEl("firmExisting").textContent=summary.already_exists||0;firmEl("firmConflicts").textContent=summary.conflicts||0;firmEl("firmSource").textContent=data.source||"-";firmEl("firmTarget").textContent=data.target||"-";firmEl("firmTargetDate").textContent=data.target_date_label||"-";firmEl("firmLastRun").textContent=`Actualizado: ${firmFmtDate(data.generated_at)}`;const badge=firmEl("firmNavBadge");const pending=summary.pending_signatures||0;badge.textContent=pending;badge.classList.toggle("hidden",pending===0);badge.classList.toggle("pulse",pending>0);if(pending)firmSetAlert(`${pending} mensaje(s) requieren firma.`,"on");else if((summary.ready||0)>0)firmSetAlert(`${summary.ready} mensaje(s) con dos firmas listos para copiar.`,"ready");else firmSetAlert("Monitor activo. No hay mensajes pendientes.");}
function firmRenderAlerts(alerts){FIRM.alerts=alerts||[];const panel=firmEl("firmSignaturePanel");const list=firmEl("firmSignatureList");panel.classList.toggle("hidden",FIRM.alerts.length===0);firmEl("firmSignatureCount").textContent=`${FIRM.alerts.length} pendiente(s)`;list.innerHTML=FIRM.alerts.map(item=>`<article class="firm-signature-item ${firmEsc(item.signature_status)}"><strong>${firmEsc(item.signature_label)}</strong><small>${firmEsc(item.voucher_number)} · ${firmEsc(item.operation_reference)}</small><small>${firmEsc(item.filename)}</small></article>`).join("");}
function firmRenderMessages(messages){FIRM.messages=messages||[];const body=firmEl("firmRows");if(!FIRM.messages.length){body.innerHTML='<tr><td colspan="7" class="muted">No hay mensajes con dos firmas en las carpetas actuales de Bryan.</td></tr>';firmUpdateActions();return;}const labels={ready:"Listo",already_exists:"Ya existe",conflict:"Conflicto"};body.innerHTML=FIRM.messages.map(item=>`<tr><td><input type="checkbox" value="${firmEsc(item.message_id)}" ${item.selectable?"":"disabled"}></td><td><strong>${firmEsc(item.voucher_number)}</strong></td><td><strong>${firmEsc(item.operation_reference)}</strong></td><td><span class="firm-file-main">${firmEsc(item.filename)}</span><span class="firm-file-sub">${firmEsc(item.source_path)}</span></td><td>${firmEsc(firmFmtDate(item.modified_at))}</td><td>${firmEsc(firmFmtBytes(item.size))}</td><td><span class="firm-tag ${firmEsc(item.status)}">${firmEsc(labels[item.status]||item.status)}</span><span class="firm-file-sub">${firmEsc(item.note)}</span></td></tr>`).join("");body.querySelectorAll("input[type=checkbox]").forEach(input=>input.addEventListener("change",firmUpdateActions));firmUpdateActions();}
function firmSeen(){try{return new Set(JSON.parse(localStorage.getItem(FIRM_SEEN_KEY)||"[]"));}catch(_){return new Set();}}
function firmSaveSeen(seen){try{localStorage.setItem(FIRM_SEEN_KEY,JSON.stringify([...seen].slice(-250)));}catch(_){} }
function firmBeep(){try{const AudioCtor=window.AudioContext||window.webkitAudioContext;if(!AudioCtor)return;FIRM.audio=FIRM.audio||new AudioCtor();const oscillator=FIRM.audio.createOscillator();const gain=FIRM.audio.createGain();oscillator.frequency.value=880;gain.gain.setValueAtTime(.0001,FIRM.audio.currentTime);gain.gain.exponentialRampToValueAtTime(.15,FIRM.audio.currentTime+.02);gain.gain.exponentialRampToValueAtTime(.0001,FIRM.audio.currentTime+.42);oscillator.connect(gain);gain.connect(FIRM.audio.destination);oscillator.start();oscillator.stop(FIRM.audio.currentTime+.43);}catch(_){} }
function firmAnnounce(alerts,force=false){if(localStorage.getItem(FIRM_ALERTS_KEY)!=="true")return;const seen=firmSeen();const fresh=force?alerts:alerts.filter(item=>!seen.has(item.alert_id));if(!fresh.length)return;fresh.forEach(item=>seen.add(item.alert_id));firmSaveSeen(seen);const first=fresh[0];const extra=fresh.length>1?` y ${fresh.length-1} mas`:"";const message=`${first.voucher_number} · ${first.operation_reference}: ${first.signature_label}${extra}.`;firmSetAlert(message,"on");firmLog(`Alerta: ${message}`);firmBeep();if("Notification" in window&&Notification.permission==="granted"){try{new Notification("Mensajes SWIFT por firmar",{body:message,tag:"automatiza-traslado-firmas"});}catch(_){}}}
function firmApply(data){firmRenderSummary(data);firmRenderAlerts(data.signature_alerts||[]);firmRenderMessages(data.messages||[]);firmAnnounce(data.signature_alerts||[]);}
async function firmLoad(){try{const data=await firmApi("/api/mensajes-firmados/status");firmApply(data);}catch(error){firmSetAlert(`Error de lectura: ${error.message}`,"on");firmLog(`Error: ${error.message}`);}}
async function firmEnableAlerts(){try{localStorage.setItem(FIRM_ALERTS_KEY,"true");const AudioCtor=window.AudioContext||window.webkitAudioContext;if(AudioCtor){FIRM.audio=FIRM.audio||new AudioCtor();await FIRM.audio.resume();}if("Notification" in window&&Notification.permission==="default")await Notification.requestPermission();firmEl("firmEnableAlerts").textContent="Alertas activas";firmLog("Alertas visuales, sonoras y del navegador activadas.");firmAnnounce(FIRM.alerts,true);}catch(error){firmLog(`No se pudieron activar las alertas: ${error.message}`);}}
async function firmOpenFolder(folder){try{const data=await firmApi("/api/mensajes-firmados/open-folder",{folder});firmLog(`Carpeta abierta: ${data.path}`);}catch(error){firmLog(`Error abriendo carpeta: ${error.message}`);alert(error.message);}}
async function firmMove(){const ids=firmSelected();if(!ids.length)return;if(!confirm(`Se copiarán ${ids.length} mensaje(s) con dos firmas a:\n${FIRM.target}\n\nSe conservarán en origen las versiones de una y dos firmas. No se sobrescribirá ningún archivo. ¿Desea continuar?`))return;try{firmEl("firmMove").disabled=true;const data=await firmApi("/api/mensajes-firmados/move",{confirmation:"COPIAR",message_ids:ids});const copied=(data.moved||[]).filter(item=>item.result==="copied").length;firmLog(`${copied} mensaje(s) copiados correctamente; los originales se conservaron.`);firmApply(data);}catch(error){firmLog(`Error al copiar: ${error.message}`);alert(error.message);}finally{firmUpdateActions();}}
function firmSelectAll(){const inputs=[...document.querySelectorAll("#firmRows input[type=checkbox]:not(:disabled)")];const select=inputs.some(input=>!input.checked);inputs.forEach(input=>{input.checked=select;});firmUpdateActions();}
function firmInit(){if(!firmEl("firmRows"))return;firmEl("firmRefresh").addEventListener("click",firmLoad);firmEl("firmSelectAll").addEventListener("click",firmSelectAll);firmEl("firmMove").addEventListener("click",firmMove);firmEl("firmEnableAlerts").addEventListener("click",firmEnableAlerts);firmEl("firmOpenSource").addEventListener("click",()=>firmOpenFolder("source"));firmEl("firmOpenTarget").addEventListener("click",()=>firmOpenFolder("destination"));if(localStorage.getItem(FIRM_ALERTS_KEY)==="true")firmEl("firmEnableAlerts").textContent="Alertas activas";firmLoad();FIRM.timer=setInterval(firmLoad,30000);}
// La inicialización se realiza después de declarar la versión de tres flujos.

// Firma EC v3: tres flujos independientes con vigilancia automática cada tres segundos.
function firmCurrent(){return FIRM.flows.find(flow=>flow.id===FIRM.flowId)||FIRM.flows[0]||null;}
function firmAllAlerts(){return FIRM.flows.flatMap(flow=>flow.signature_alerts||[]);}
function firmRenderTabs(){document.querySelectorAll("#firmFlowTabs button[data-flow]").forEach(button=>{const flow=FIRM.flows.find(item=>item.id===button.dataset.flow);button.classList.toggle("act",button.dataset.flow===FIRM.flowId);const count=firmEl(`firmFlowCount-${button.dataset.flow}`);if(!count)return;if(!flow){count.textContent="Sin lectura";return;}if(!flow.ok){count.textContent="Error de acceso";return;}const pending=(flow.summary||{}).pending_signatures||0;const ready=(flow.summary||{}).ready||0;count.textContent=pending?`${pending} por firmar`:ready?`${ready} listos`:"Sin pendientes";});}
function firmRenderSummary(data){const flow=firmCurrent();const summary=(flow&&flow.summary)||{};FIRM.summary=summary;FIRM.target=data.target||"";firmEl("firmPending").textContent=summary.pending_signatures||0;firmEl("firmDetected").textContent=summary.detected||0;firmEl("firmReady").textContent=summary.ready||0;firmEl("firmExisting").textContent=summary.already_exists||0;firmEl("firmConflicts").textContent=summary.conflicts||0;firmEl("firmSource").textContent=flow?flow.source:"-";firmEl("firmTarget").textContent=data.target||"-";firmEl("firmTargetDate").textContent=data.target_date_label||"-";firmEl("firmScannedFolders").textContent=flow&&flow.scanned_folders&&flow.scanned_folders.length?flow.scanned_folders.join(" · "):"Sin carpetas recientes con archivos";firmEl("firmFlowTitle").textContent=flow?flow.label:"Firma EC";firmEl("firmFlowDescription").textContent=flow?flow.description:"Esperando la primera lectura.";firmEl("firmLastRun").textContent=`Actualizado: ${firmFmtDate(data.generated_at)} · vigilancia automática cada ${data.automatic_refresh_seconds||3} segundos`;const aggregate=data.summary||{};const totalPending=aggregate.pending_signatures||0;const badge=firmEl("firmNavBadge");badge.textContent=totalPending;badge.classList.toggle("hidden",totalPending===0);badge.classList.toggle("pulse",totalPending>0);if(data.monitor_warming)firmSetAlert("Iniciando la vigilancia en tiempo real de los tres orígenes...");else if(flow&&!flow.ok)firmSetAlert(`No se pudo leer ${flow.label}: ${flow.error}`,"on");else if((summary.pending_signatures||0)>0)firmSetAlert(`${summary.pending_signatures} mensaje(s) requieren firma en ${flow.label}.`,"on");else if((summary.ready||0)>0)firmSetAlert(`${summary.ready} mensaje(s) con dos firmas listos para copiar en ${flow.label}.`,"ready");else if(totalPending>0)firmSetAlert(`Este flujo está al día. Hay ${totalPending} pendiente(s) en otro flujo de Firma EC.`,"on");else firmSetAlert("Monitor en tiempo real activo. Los tres flujos están al día.");}
function firmRenderAlerts(alerts){FIRM.alerts=alerts||[];const panel=firmEl("firmSignaturePanel");const list=firmEl("firmSignatureList");panel.classList.toggle("hidden",FIRM.alerts.length===0);firmEl("firmSignatureCount").textContent=`${FIRM.alerts.length} pendiente(s)`;list.innerHTML=FIRM.alerts.map(item=>`<article class="firm-signature-item ${firmEsc(item.signature_status)}"><strong>${firmEsc(item.signature_label)}</strong><small>${firmEsc(item.voucher_number)} · ${firmEsc(item.operation_reference)}</small><small>${firmEsc(item.filename)}</small></article>`).join("");}
function firmRenderMessages(messages){FIRM.messages=messages||[];const body=firmEl("firmRows");const flow=firmCurrent();if(!FIRM.messages.length){body.innerHTML=`<tr><td colspan="7" class="muted">No hay mensajes con dos firmas listos en ${firmEsc(flow?flow.label:"este flujo")}.</td></tr>`;firmUpdateActions();return;}const labels={ready:"Listo",already_exists:"Ya existe",conflict:"Conflicto"};body.innerHTML=FIRM.messages.map(item=>`<tr><td><input type="checkbox" value="${firmEsc(item.message_id)}" ${item.selectable?"":"disabled"}></td><td><strong>${firmEsc(item.voucher_number)}</strong></td><td><strong>${firmEsc(item.operation_reference)}</strong></td><td><span class="firm-file-main">${firmEsc(item.filename)}</span><span class="firm-file-sub">${firmEsc(item.source_path)}</span></td><td>${firmEsc(firmFmtDate(item.modified_at))}</td><td>${firmEsc(firmFmtBytes(item.size))}</td><td><span class="firm-tag ${firmEsc(item.status)}">${firmEsc(labels[item.status]||item.status)}</span><span class="firm-file-sub">${firmEsc(item.note)}</span></td></tr>`).join("");body.querySelectorAll("input[type=checkbox]").forEach(input=>input.addEventListener("change",firmUpdateActions));firmUpdateActions();}
function firmAnnounce(alerts,force=false){if(localStorage.getItem(FIRM_ALERTS_KEY)!=="true")return;const seen=firmSeen();const fresh=force?alerts:alerts.filter(item=>!seen.has(item.alert_id));if(!fresh.length)return;fresh.forEach(item=>seen.add(item.alert_id));firmSaveSeen(seen);const first=fresh[0];const extra=fresh.length>1?` y ${fresh.length-1} más`:"";const message=`${first.flow_label}: ${first.voucher_number} · ${first.operation_reference}, ${first.signature_label}${extra}.`;firmSetAlert(message,"on");firmLog(`Alerta: ${message}`);firmBeep();if("Notification" in window&&Notification.permission==="granted"){try{new Notification("Firma EC · mensajes por firmar",{body:message,tag:"firma-ec-pendientes"});}catch(_){}}}
function firmRenderCurrent(){if(!FIRM.lastData)return;const flow=firmCurrent();firmRenderTabs();firmRenderSummary(FIRM.lastData);firmRenderAlerts(flow?flow.signature_alerts:[]);firmRenderMessages(flow?flow.messages:[]);}
function firmSetFlow(flowId){if(!["deuda_externa","automaticos","bancos"].includes(flowId))return;FIRM.flowId=flowId;if(FIRM.lastData)firmRenderCurrent();else firmRenderTabs();}
function firmApply(data){FIRM.lastData=data;FIRM.flows=data.flows||[];if(!FIRM.flows.some(flow=>flow.id===FIRM.flowId)&&FIRM.flows.length)FIRM.flowId=FIRM.flows[0].id;firmRenderCurrent();firmAnnounce(firmAllAlerts());}
async function firmLoad(){if(FIRM.loading)return;FIRM.loading=true;try{const data=await firmApi("/api/mensajes-firmados/status");firmApply(data);}catch(error){firmSetAlert(`Error de lectura: ${error.message}`,"on");firmLog(`Error: ${error.message}`);}finally{FIRM.loading=false;}}
async function firmEnableAlerts(){try{localStorage.setItem(FIRM_ALERTS_KEY,"true");const AudioCtor=window.AudioContext||window.webkitAudioContext;if(AudioCtor){FIRM.audio=FIRM.audio||new AudioCtor();await FIRM.audio.resume();}if("Notification" in window&&Notification.permission==="default")await Notification.requestPermission();firmEl("firmEnableAlerts").textContent="Alertas activas";firmLog("Alertas visuales, sonoras y del navegador activadas para los tres flujos.");firmAnnounce(firmAllAlerts(),true);}catch(error){firmLog(`No se pudieron activar las alertas: ${error.message}`);}}
async function firmOpenFolder(folder){try{const data=await firmApi("/api/mensajes-firmados/open-folder",{flow_id:FIRM.flowId,folder});firmLog(`Carpeta abierta: ${data.path}`);}catch(error){firmLog(`Error abriendo carpeta: ${error.message}`);alert(error.message);}}
async function firmMove(){const ids=firmSelected();const flow=firmCurrent();if(!ids.length||!flow)return;if(!confirm(`Se copiarán ${ids.length} mensaje(s) de ${flow.label} a:\n${FIRM.target}\n\nSe conservarán en la carpeta de origen tanto la versión de una firma como la de dos firmas. No se sobrescribirá ningún archivo. ¿Desea continuar?`))return;try{firmEl("firmMove").disabled=true;const data=await firmApi("/api/mensajes-firmados/move",{confirmation:"COPIAR",flow_id:flow.id,message_ids:ids});const copied=(data.moved||[]).filter(item=>item.result==="copied").length;const observations=(data.moved||[]).filter(item=>item.result!=="copied").length;firmLog(`${copied} mensaje(s) de ${flow.label} copiados correctamente; se conservaron las versiones de una y dos firmas en origen${observations?` · ${observations} observación(es)`:""}.`);firmApply(data);}catch(error){firmLog(`Error al copiar: ${error.message}`);alert(error.message);}finally{firmUpdateActions();}}
function firmInit(){if(!firmEl("firmRows")||FIRM.initialized)return;FIRM.initialized=true;firmEl("firmRefresh").addEventListener("click",firmLoad);firmEl("firmSelectAll").addEventListener("click",firmSelectAll);firmEl("firmMove").addEventListener("click",firmMove);firmEl("firmEnableAlerts").addEventListener("click",firmEnableAlerts);firmEl("firmOpenSource").addEventListener("click",()=>firmOpenFolder("source"));firmEl("firmOpenTarget").addEventListener("click",()=>firmOpenFolder("destination"));document.querySelectorAll("#firmFlowTabs button[data-flow]").forEach(button=>button.addEventListener("click",()=>firmSetFlow(button.dataset.flow)));if(localStorage.getItem(FIRM_ALERTS_KEY)==="true")firmEl("firmEnableAlerts").textContent="Alertas activas";firmLoad();FIRM.timer=setInterval(firmLoad,3000);document.addEventListener("visibilitychange",()=>{if(!document.hidden)firmLoad();});window.addEventListener("focus",firmLoad);}
firmInit();


const COMP2={analysisId:null,payments:[],busy:false,configLoaded:false,view:"pending",dateMode:"today",owner:"all",query:""};
function compEl(id){return document.getElementById(id);}
function compEsc(value){return String(value??"").replace(/[&<>"']/g,char=>({"&":"&amp;","<":"&lt;",">":"&gt;","\"":"&quot;","'":"&#39;"}[char]));}
function compLog(text){const el=compEl("compLog");if(!el)return;el.textContent+=`[${new Date().toLocaleTimeString("es-EC")}] ${text}\n`;el.scrollTop=el.scrollHeight;}
async function compApi(path,payload){const options=payload===undefined?{cache:"no-store"}:{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)};const response=await fetch(path,options);const data=await response.json().catch(()=>({ok:false,error:"Respuesta invalida del servidor."}));if(!response.ok||!data.ok)throw new Error(data.error||`Error HTTP ${response.status}`);return data;}
function compSetBusy(busy,label){COMP2.busy=busy;compEl("compAnalizar").disabled=busy;compEl("compRefrescar").disabled=busy;compEl("compGuardarConfig").disabled=busy;if(label)compEl("compLastRun").textContent=label;compUpdateSelection();}
function compConfigPayload(){return {bryan_source:compEl("compBryan").value.trim(),steven_source:compEl("compSteven").value.trim(),ack_base:compEl("compAckBase").value.trim(),destination:compEl("compDestino").value.trim(),recursive:compEl("compRecursivo").checked,use_ocr:compEl("compOcr").checked,automatic:compEl("compAutomatico").checked};}
function compApplyConfig(config){if(!config)return;compEl("compBryan").value=config.bryan_source||"";compEl("compSteven").value=config.steven_source||"";compEl("compAckBase").value=config.ack_base||"";compEl("compDestino").value=config.destination||"";compEl("compRecursivo").checked=!!config.recursive;compEl("compOcr").checked=config.use_ocr!==false;compEl("compAutomatico").checked=config.automatic!==false;COMP2.configLoaded=true;}
function compStatusPresentation(status){return {"LISTO":["ready","Listo"],"LISTO PARCIAL":["partial","Listo parcial"],"LISTO COMPROBANTE":["ready","Comprobante listo"],"ACK PENDIENTE":["pending","ACK pendiente"],"COINCIDENCIA PROBABLE":["probable","Coincidencia probable"],"ARCHIVADO":["archived","Archivado"],"YA ARCHIVADO":["archived","Ya archivado"],"REVISAR":["review","Revisar"],"ERROR":["error","Error"],"CONFLICTO":["conflict","Conflicto"]}[status]||["review",status||"Pendiente"];}
const COMP_ARCHIVED=new Set(["ARCHIVADO","YA ARCHIVADO"]);
const COMP_REVIEW=new Set(["REVISAR","ERROR","CONFLICTO"]);
function compDateKey(value){const match=String(value||"").match(/^(\d{2})\/(\d{2})\/(\d{4})$/);return match?`${match[3]}-${match[2]}-${match[1]}`:"";}
function compTodayKey(){const now=new Date();return `${now.getFullYear()}-${String(now.getMonth()+1).padStart(2,"0")}-${String(now.getDate()).padStart(2,"0")}`;}
function compMatchesDate(payment){if(COMP2.dateMode==="all")return true;const key=compDateKey(payment.value_date);if(!key)return false;if(COMP2.dateMode==="today")return key===compTodayKey();const selected=new Date(`${key}T12:00:00`);const today=new Date(`${compTodayKey()}T12:00:00`);const days=Math.round((today-selected)/86400000);return days>=0&&days<7;}
function compMatchesView(payment,view=COMP2.view){if(view==="all")return true;if(view==="archived")return COMP_ARCHIVED.has(payment.status);if(view==="ready")return !!payment.ready;if(view==="ack")return payment.status==="ACK PENDIENTE";if(view==="review")return COMP_REVIEW.has(payment.status);return !COMP_ARCHIVED.has(payment.status);}
function compMatchesCommon(payment){if(!compMatchesDate(payment))return false;if(COMP2.owner!=="all"&&payment.responsible!==COMP2.owner)return false;if(!COMP2.query)return true;const text=[payment.voucher_number,payment.operation_reference,payment.loan_reference,payment.office_reference,payment.destination_name,payment.beneficiary].join(" ").toLocaleLowerCase("es");return text.includes(COMP2.query);}
function compFilteredPayments(){return COMP2.payments.filter(payment=>compMatchesCommon(payment)&&compMatchesView(payment));}
function compActionLabel(payment){if(payment.ready)return ["action","Debe archivarse"];if(payment.status==="ACK PENDIENTE")return ["wait","Comprobante guardado; esperando ACK"];if(COMP_REVIEW.has(payment.status))return ["review","Requiere validación antes de archivar"];if(COMP_ARCHIVED.has(payment.status))return ["","Archivo completado"];return ["review","Pendiente de clasificación"];}
function compUpdateFilterCounts(){const common=COMP2.payments.filter(compMatchesCommon);const views={pending:"compCountPending",ready:"compCountReady",ack:"compCountAck",review:"compCountReview",archived:"compCountArchived",all:"compCountAll"};Object.entries(views).forEach(([view,id])=>{compEl(id).textContent=common.filter(payment=>compMatchesView(payment,view)).length;});document.querySelectorAll("#compViewTabs [data-comp-view]").forEach(button=>button.classList.toggle("act",button.dataset.compView===COMP2.view));}
function compSelected(){return [...document.querySelectorAll("#compRows .comp-select:checked")].map(input=>input.value);}
function compUpdateSelection(){const selected=compSelected();const safe=[...document.querySelectorAll('#compRows .comp-select:not([data-probable="true"])')];const selectedSafe=safe.filter(input=>input.checked).length;compEl("compSeleccionCuenta").textContent=`${selected.length} seleccionado${selected.length===1?"":"s"}`;compEl("compArchivar").disabled=COMP2.busy||selected.length===0;compEl("compSeleccionarTodos").checked=safe.length>0&&selectedSafe===safe.length;compEl("compSeleccionarTodos").indeterminate=selectedSafe>0&&selectedSafe<safe.length;}
function compDetail(payment){const evidence=(payment.match_evidence||[]).length?`<p><strong>Evidencia:</strong> ${(payment.match_evidence||[]).map(compEsc).join(" · ")}</p>`:"";const notes=(payment.notes||[]).length?`<p><strong>Notas:</strong> ${(payment.notes||[]).map(compEsc).join(" · ")}</p>`:"";const ack=payment.requires_ack?compEsc(payment.ack_source||"No encontrado"):"No aplica para transferencia interna";return `<details class="comp-detail"><summary>Ver</summary><div class="comp-detail-box"><p><strong>Oficio:</strong> ${compEsc(payment.office_reference||"No identificado")}</p><p><strong>Pedido:</strong> ${compEsc(payment.order_number||"No identificado")}</p><p><strong>Beneficiario:</strong> ${compEsc(payment.beneficiary||"No identificado")}</p><p><strong>Monto / moneda:</strong> ${compEsc((payment.amounts||[]).join(" · ")||"No identificado")} ${compEsc(payment.currency||"")}</p><p><strong>Concepto:</strong> ${compEsc(payment.concept||"No identificado")}</p><p><strong>Comprobante origen:</strong> ${compEsc(payment.accounting_source||"")}</p><p><strong>ACK origen:</strong> ${ack}</p>${evidence}${notes}</div></details>`;}
function compRenderRows(){const body=compEl("compRows");const payments=compFilteredPayments();compUpdateFilterCounts();compEl("compFilterMeta").textContent=`Mostrando ${payments.length} de ${COMP2.payments.length} comprobante${COMP2.payments.length===1?"":"s"}`;if(!payments.length){const completed=COMP2.payments.filter(payment=>compMatchesCommon(payment)&&COMP_ARCHIVED.has(payment.status)).length;const pendingEmpty=completed?`No hay comprobantes pendientes. ${completed} ya ${completed===1?"fue archivado":"fueron archivados"}; puede verlos en la pestaña Archivados.`:"No hay comprobantes pendientes para los filtros seleccionados.";const messages={pending:pendingEmpty,ready:"No hay comprobantes listos para archivar en esta fecha.",ack:"No hay ACK pendientes para los filtros seleccionados.",review:"No hay comprobantes que requieran revisión.",archived:"No hay comprobantes archivados para los filtros seleccionados.",all:"No se encontraron comprobantes para los filtros seleccionados."};body.innerHTML=`<tr><td colspan="10" class="comp-empty-filter">${messages[COMP2.view]}</td></tr>`;compUpdateSelection();return;}body.innerHTML=payments.map(payment=>{const [statusClass,statusLabel]=compStatusPresentation(payment.status);const [actionClass,actionLabel]=compActionLabel(payment);const checkbox=payment.ready?`<input class="comp-select" type="checkbox" value="${compEsc(payment.payment_id)}" data-probable="${payment.probable_match?"true":"false"}" title="${payment.probable_match?"Requiere revision individual":"Seleccionar"}">`:"";const copy=payment.destination?`<button class="comp-copy" type="button" data-path="${compEsc(payment.destination)}">Copiar ruta</button>`:"";return `<tr><td>${checkbox}</td><td><span class="comp-owner ${payment.responsible==="STEVEN"?"steven":""}">${compEsc(payment.responsible)}</span></td><td><span class="comp-code">${compEsc(payment.voucher_number)}</span><small>${compEsc(payment.document_type==="TRANSFERENCIA INTERNA"?"Transferencia interna":"Pago exterior")}</small></td><td><span class="comp-ref">${compEsc(payment.operation_reference)}</span></td><td>${compEsc(payment.value_date||"—")}</td><td>${compEsc(payment.loan_reference||"—")}</td><td title="${compEsc(payment.destination||"")}"><span class="comp-dest">${compEsc(payment.destination_name||"Sin coincidencia")}</span>${copy}</td><td><span class="comp-match">${compEsc(payment.match_method||"—")}</span><span class="comp-score">${compEsc(payment.match_score||0)}%${payment.ocr_used?" · OCR":""}</span></td><td><span class="comp-status ${statusClass}">${compEsc(statusLabel)}</span><span class="comp-priority-note ${actionClass}">${compEsc(actionLabel)}</span></td><td>${compDetail(payment)}</td></tr>`;}).join("");body.querySelectorAll(".comp-select").forEach(input=>input.addEventListener("change",compUpdateSelection));body.querySelectorAll(".comp-copy").forEach(button=>button.addEventListener("click",async()=>{try{await navigator.clipboard.writeText(button.dataset.path);compLog("Ruta destino copiada.");}catch(_){compLog("No se pudo copiar la ruta; puede verla en Detalle.");}}));compUpdateSelection();}
function compApply(data){if(data.analysis_id)COMP2.analysisId=data.analysis_id;if(data.payments)COMP2.payments=data.payments;if(data.config&&!COMP2.configLoaded)compApplyConfig(data.config);const summary=data.summary||{};compEl("compDetectados").textContent=summary.detected||0;compEl("compListos").textContent=summary.ready||0;compEl("compArchivados").textContent=summary.archived||0;compEl("compAckPendientes").textContent=summary.ack_pending||0;compEl("compRevision").textContent=summary.review||0;const notice=compEl("compAutoNotice");notice.classList.remove("warn","bad");if(data.auto_error){notice.classList.add("bad");notice.querySelector("strong").textContent="Monitor reintentando";notice.querySelector("span").textContent=data.auto_error;}else if(data.auto_status==="PAUSADO"){notice.classList.add("warn");notice.querySelector("strong").textContent="Modo automatico pausado";notice.querySelector("span").textContent="Puede analizar y archivar manualmente desde esta pantalla.";}else{notice.querySelector("strong").textContent="Modo automatico activo";notice.querySelector("span").textContent="Cada cinco minutos revisa BRYAN, STEVEN y los ACK. Solo archiva coincidencias exactas de 99-100%.";}if(data.generated_at){const date=new Date(data.generated_at);const origin=data.analysis_origin==="AUTOMÁTICO"?"Automatico":"Manual";compEl("compLastRun").textContent=`${origin} · Ultima revision: ${date.toLocaleString("es-EC")}`;compEl("compMeta").textContent=`${summary.detected||0} registros · monitor ${String(data.auto_status||"activo").toLowerCase()}`;}else{compEl("compLastRun").textContent=`Monitor ${String(data.auto_status||"iniciando").toLowerCase()}`;}compRenderRows();}
async function compLoad(silent=false){if(COMP2.busy)return;try{const data=await compApi("/api/comprobantes_contables/status");compApply(data);if(!silent)compLog("Vista actualizada.");}catch(error){compEl("compAutoNotice").classList.add("bad");compLog(`Error de lectura: ${error.message}`);}}
async function compAnalyze(){try{compSetBusy(true,"Analizando BRYAN, STEVEN y ACK...");const data=await compApi("/api/comprobantes_contables/analyze",compConfigPayload());compApply(data);compLog(`Analisis manual completado: ${(data.summary||{}).detected||0} registros.`);}catch(error){compLog(`Error al analizar: ${error.message}`);alert(error.message);}finally{compSetBusy(false);}}
async function compSaveConfig(){try{compSetBusy(true,"Guardando configuracion...");const data=await compApi("/api/comprobantes_contables/config",compConfigPayload());compApplyConfig(data.config);compLog("Configuracion guardada; el monitor realizara una nueva lectura.");await compLoad(true);}catch(error){compLog(`Error de configuracion: ${error.message}`);alert(error.message);}finally{compSetBusy(false);}}
async function compArchive(){const selected=compSelected();if(!selected.length)return;const confirmation=prompt(`Se archivaran ${selected.length} registro(s). Los originales no se eliminaran y ningun archivo se sobrescribira.\n\nEscriba ARCHIVAR para continuar:`);if(confirmation!=="ARCHIVAR")return;try{compSetBusy(true,"Copiando y verificando documentos...");const data=await compApi("/api/comprobantes_contables/archive",{analysis_id:COMP2.analysisId,references:selected,confirmation});compApply(data);const archived=(data.processed||[]).filter(item=>item.status==="ARCHIVADO").length;compLog(`Proceso terminado: ${archived} archivado(s), ${(data.processed||[]).length-archived} con observaciones.`);}catch(error){compLog(`Error al archivar: ${error.message}`);alert(error.message);}finally{compSetBusy(false);}}
function compInit(){if(!compEl("compRows"))return;compEl("compRefrescar").addEventListener("click",()=>compLoad());compEl("compAnalizar").addEventListener("click",compAnalyze);compEl("compGuardarConfig").addEventListener("click",compSaveConfig);compEl("compArchivar").addEventListener("click",compArchive);document.querySelectorAll("#compViewTabs [data-comp-view]").forEach(button=>button.addEventListener("click",()=>{COMP2.view=button.dataset.compView;compRenderRows();}));compEl("compDateMode").addEventListener("change",event=>{COMP2.dateMode=event.target.value;compRenderRows();});compEl("compOwnerFilter").addEventListener("change",event=>{COMP2.owner=event.target.value;compRenderRows();});compEl("compSearch").addEventListener("input",event=>{COMP2.query=event.target.value.trim().toLocaleLowerCase("es");compRenderRows();});compEl("compSeleccionarTodos").addEventListener("change",()=>{document.querySelectorAll('#compRows .comp-select:not([data-probable="true"])').forEach(input=>{input.checked=compEl("compSeleccionarTodos").checked;});compUpdateSelection();});compLoad(true);setInterval(()=>compLoad(true),30000);}
compInit();

// Archivo Quipux: inventario persistente, lectura PDF/OCR y archivo seguro.
const QUIPUX={records:[],busy:false,configLoaded:false};
const quipuxEl=id=>document.getElementById(id);
const quipuxEsc=value=>String(value??"").replace(/[&<>"']/g,char=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#039;"}[char]));
async function quipuxApi(path,payload){const options=payload===undefined?{}:{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)};const response=await fetch(path,options);const data=await response.json();if(!response.ok||data.ok===false&&data.last_error)throw new Error(data.error||data.last_error||`Error HTTP ${response.status}`);return data;}
function quipuxStatusClass(status){status=String(status||"").toUpperCase();if(status.includes("ARCHIVADO"))return "archived";if(status==="REVISAR")return "review";if(status==="CONFLICTO")return "conflict";if(status==="ERROR")return "error";if(status==="IGNORADO")return "ignored";if(status.includes("HISTÓRICO"))return "historical";return "";}
function quipuxDate(value){if(!value)return "—";const date=new Date(value);return Number.isNaN(date.getTime())?quipuxEsc(value):date.toLocaleString("es-EC");}
function quipuxDetail(record){const debtEvidence=(record.public_debt_evidence||[]).length?`<p><strong>Validación deuda pública:</strong> ${(record.public_debt_evidence||[]).map(quipuxEsc).join(" · ")}</p>`:"";const evidence=(record.evidence||[]).length?`<p><strong>Evidencia del destino:</strong> ${(record.evidence||[]).map(quipuxEsc).join(" · ")}</p>`:"";const notes=(record.notes||[]).length?`<p><strong>Notas:</strong> ${(record.notes||[]).map(quipuxEsc).join(" · ")}</p>`:"";const candidates=(record.candidate_paths||[]).length?`<p><strong>Candidatos:</strong> ${(record.candidate_paths||[]).map(quipuxEsc).join(" | ")}</p>`:"";const missing=(record.missing_documents||[]).length?`<p><strong>Faltan:</strong> ${(record.missing_documents||[]).map(quipuxEsc).join(" · ")}</p>`:"";return `<details class="quipux-detail"><summary>Ver</summary><div class="quipux-detail-box"><p><strong>Oficio:</strong> ${quipuxEsc(record.office_reference||"No identificado")}</p><p><strong>Moneda / montos:</strong> ${quipuxEsc(record.currency||"—")} ${quipuxEsc((record.amounts||[]).join(" · ")||"—")}</p><p><strong>OCR:</strong> ${record.ocr_used?"Sí":"No"}</p><p><strong>Origen:</strong> ${quipuxEsc(record.source_path||"")}</p>${debtEvidence}${evidence}${missing}${candidates}${notes}</div></details>`;}
function quipuxRenderRows(){const body=quipuxEl("quipuxRows");if(!QUIPUX.records.length){body.innerHTML='<tr><td colspan="10" class="muted">No hay documentos registrados. El monitor está listo para la próxima descarga.</td></tr>';return;}body.innerHTML=QUIPUX.records.map(record=>{const reference=[record.loan_reference,record.sigade].filter(Boolean).join(" · ")||"—";const statusClass=quipuxStatusClass(record.status);return `<tr><td>${quipuxDate(record.downloaded_at)}</td><td><span class="quipux-file">${quipuxEsc(record.source_name)}</span></td><td><span class="quipux-type">${quipuxEsc(record.document_type||"—")}</span></td><td>${quipuxEsc(reference)}</td><td>${quipuxEsc(record.value_date||"—")}</td><td title="${quipuxEsc(record.destination||"")}"><span class="quipux-dest">${quipuxEsc(record.destination_name||"Sin coincidencia")}</span></td><td><span class="quipux-target">${quipuxEsc(record.target_name||"—")}</span></td><td><span class="quipux-match">${quipuxEsc(record.match_method||"—")}</span><span class="quipux-score">${quipuxEsc(record.score||0)}%${record.ocr_used?" · OCR":""}</span></td><td><span class="quipux-status ${statusClass}">${quipuxEsc(record.status||"—")}</span><span class="quipux-package">${quipuxEsc(record.package_status||"")}</span></td><td>${quipuxDetail(record)}</td></tr>`;}).join("");}
function quipuxApplyConfig(config){if(!config)return;quipuxEl("quipuxSource").value=config.source||"";quipuxEl("quipuxDestination").value=config.destination||"";quipuxEl("quipuxOcr").checked=!!config.use_ocr;quipuxEl("quipuxAutomatic").checked=!!config.automatic;QUIPUX.configLoaded=true;}
function quipuxPayload(){return {source:quipuxEl("quipuxSource").value.trim(),destination:quipuxEl("quipuxDestination").value.trim(),use_ocr:quipuxEl("quipuxOcr").checked,automatic:quipuxEl("quipuxAutomatic").checked};}
function quipuxApply(data){QUIPUX.records=data.records||[];if(data.config&&!QUIPUX.configLoaded)quipuxApplyConfig(data.config);const summary=data.summary||{};quipuxEl("quipuxDetected").textContent=summary.detected||0;quipuxEl("quipuxArchived").textContent=summary.archived||0;quipuxEl("quipuxReview").textContent=summary.review||0;quipuxEl("quipuxIncomplete").textContent=summary.incomplete||0;quipuxEl("quipuxHistorical").textContent=summary.historical||0;const badge=quipuxEl("quipuxNavBadge");badge.textContent=summary.review||0;badge.classList.toggle("hidden",!(summary.review||0));const notice=quipuxEl("quipuxNotice");notice.classList.remove("warn","bad");if(data.last_error){notice.classList.add("bad");notice.querySelector("strong").textContent="Monitor reintentando";notice.querySelector("span").textContent=data.last_error;}else if(data.monitor_status==="PAUSADO"){notice.classList.add("warn");notice.querySelector("strong").textContent="Monitor automático pausado";notice.querySelector("span").textContent="Puede ejecutar una revisión manual desde esta pantalla.";}else{notice.querySelector("strong").textContent="Monitor automático activo";notice.querySelector("span").textContent="Revisa archivos nuevos cada minuto. Conserva los originales y nunca sobrescribe un destino.";}quipuxEl("quipuxLastRun").textContent=data.last_scan?`Última revisión: ${quipuxDate(data.last_scan)}`:`Protección inicial: ${quipuxDate(data.initialized_at)}`;quipuxEl("quipuxMeta").textContent=`${summary.detected||0} documentos nuevos · ${summary.historical||0} históricos protegidos · monitor ${String(data.monitor_status||"iniciando").toLowerCase()}`;quipuxRenderRows();}
async function quipuxLoad(silent=false){if(QUIPUX.busy)return;try{const data=await quipuxApi("/api/archivo_quipux/status");quipuxApply(data);}catch(error){const notice=quipuxEl("quipuxNotice");notice.classList.add("bad");notice.querySelector("strong").textContent="No se pudo consultar el monitor";notice.querySelector("span").textContent=error.message;if(!silent)alert(error.message);}}
async function quipuxScan(){if(QUIPUX.busy)return;QUIPUX.busy=true;const button=quipuxEl("quipuxScan");button.disabled=true;button.textContent="Revisando PDF...";try{const data=await quipuxApi("/api/archivo_quipux/scan",{});quipuxApply(data);}catch(error){alert(error.message);}finally{QUIPUX.busy=false;button.disabled=false;button.textContent="Revisar ahora";}}
async function quipuxSaveConfig(){if(QUIPUX.busy)return;QUIPUX.busy=true;try{const data=await quipuxApi("/api/archivo_quipux/config",quipuxPayload());quipuxApplyConfig(data.config);await quipuxLoad(true);}catch(error){alert(error.message);}finally{QUIPUX.busy=false;}}
function quipuxInit(){if(!quipuxEl("quipuxRows"))return;quipuxEl("quipuxRefresh").addEventListener("click",()=>quipuxLoad());quipuxEl("quipuxScan").addEventListener("click",quipuxScan);quipuxEl("quipuxSaveConfig").addEventListener("click",quipuxSaveConfig);quipuxLoad(true);setInterval(()=>quipuxLoad(true),30000);}
quipuxInit();

// Oficios de Contratos de Agencia Fiscal: clasificación estricta y archivo seguro.
const CONTRACTS={records:[],busy:false,configLoaded:false};
const contractEl=id=>document.getElementById(id);
const contractEsc=value=>String(value??"").replace(/[&<>"']/g,char=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#039;"}[char]));
async function contractApi(path,payload){const options=payload===undefined?{}:{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)};const response=await fetch(path,options);const data=await response.json();if(!response.ok)throw new Error(data.error||`Error HTTP ${response.status}`);return data;}
function contractDate(value){if(!value)return "—";const date=new Date(value);return Number.isNaN(date.getTime())?contractEsc(value):date.toLocaleString("es-EC");}
function contractStatusClass(status){status=String(status||"").toUpperCase();if(status.includes("ARCHIVADO"))return "archived";if(status==="REVISAR")return "review";if(status==="CONFLICTO")return "conflict";if(status==="ERROR")return "error";if(status==="IGNORADO")return "ignored";return "";}
function contractDetail(record){const evidence=(record.evidence||[]).length?`<p><strong>Evidencia:</strong> ${(record.evidence||[]).map(contractEsc).join(" · ")}</p>`:"";const offices=(record.referenced_offices||[]).length?`<p><strong>Antecedentes:</strong> ${(record.referenced_offices||[]).map(contractEsc).join(" · ")}</p>`:"";const notes=(record.notes||[]).length?`<p><strong>Notas:</strong> ${(record.notes||[]).map(contractEsc).join(" · ")}</p>`:"";const candidates=(record.candidate_paths||[]).length?`<p><strong>Carpetas candidatas:</strong> ${(record.candidate_paths||[]).map(contractEsc).join(" | ")}</p>`:"";return `<details class="quipux-detail"><summary>Ver</summary><div class="quipux-detail-box"><p><strong>Oficio:</strong> ${contractEsc(record.office_reference||"No identificado")}</p><p><strong>Firma electrónica:</strong> ${record.signature_verified?"Verificada":"No verificada"}</p><p><strong>Etapa / patrón:</strong> ${contractEsc(record.descriptor||"No aplica")} · ${contractEsc(record.descriptor_method||"")}</p><p><strong>Origen:</strong> ${contractEsc(record.source_path||"")}</p><p><strong>Destino:</strong> ${contractEsc(record.destination||"Sin asignar")}</p>${evidence}${offices}${candidates}${notes}</div></details>`;}
function contractRenderRows(){const body=contractEl("contractRows");if(!CONTRACTS.records.length){body.innerHTML='<tr><td colspan="9" class="muted">No hay oficios contractuales registrados. El monitor está listo para la próxima descarga.</td></tr>';return;}body.innerHTML=CONTRACTS.records.map(record=>{const references=[...(record.contract_references||[]),...(record.sigade_numbers||[])].join(" · ")||"—";const creditor=[record.creditor,record.group].filter(Boolean).filter((value,index,array)=>array.indexOf(value)===index).join(" · ")||"—";return `<tr><td>${contractDate(record.downloaded_at)}</td><td><span class="contract-file">${contractEsc(record.source_name)}</span><span class="contract-office">${contractEsc(record.office_reference||"No identificado")}</span></td><td><span class="contract-subject">${contractEsc(record.subject||"—")}</span><span class="quipux-package">${contractEsc(record.descriptor||"Documento excluido")}</span></td><td><span class="contract-ref">${contractEsc(references)}</span></td><td title="${contractEsc(record.destination||"")}"><span class="contract-ref">${contractEsc(creditor)}</span><span class="contract-dest">${contractEsc(record.destination_name||"Sin coincidencia")}</span></td><td><span class="contract-target">${contractEsc(record.target_name||"—")}</span></td><td><span class="contract-match">${contractEsc(record.match_method||"—")}</span><span class="contract-score">${contractEsc(record.score||0)}%${record.ocr_used?" · OCR":""}</span></td><td><span class="quipux-status ${contractStatusClass(record.status)}">${contractEsc(record.status||"—")}</span></td><td>${contractDetail(record)}</td></tr>`;}).join("");}
function contractApplyConfig(config){if(!config)return;contractEl("contractSource").value=config.source||"";contractEl("contractDestination").value=config.destination||"";contractEl("contractOcr").checked=!!config.use_ocr;contractEl("contractAutomatic").checked=!!config.automatic;CONTRACTS.configLoaded=true;}
function contractPayload(){return {source:contractEl("contractSource").value.trim(),destination:contractEl("contractDestination").value.trim(),use_ocr:contractEl("contractOcr").checked,automatic:contractEl("contractAutomatic").checked};}
function contractApply(data){CONTRACTS.records=data.records||[];if(data.config&&!CONTRACTS.configLoaded)contractApplyConfig(data.config);const summary=data.summary||{};contractEl("contractCandidates").textContent=summary.candidates||0;contractEl("contractArchived").textContent=summary.archived||0;contractEl("contractReview").textContent=summary.review||0;contractEl("contractFolders").textContent=summary.folders||0;contractEl("contractPatterns").textContent=summary.patterns||0;const badge=contractEl("contractNavBadge");badge.textContent=summary.review||0;badge.classList.toggle("hidden",!(summary.review||0));const notice=contractEl("contractNotice");notice.classList.remove("warn","bad");if(data.last_error){notice.classList.add("bad");notice.querySelector("strong").textContent="Monitor reintentando";notice.querySelector("span").textContent=data.last_error;}else if(data.monitor_status==="PAUSADO"){notice.classList.add("warn");notice.querySelector("strong").textContent="Monitor automático pausado";notice.querySelector("span").textContent="Puede ejecutar una revisión manual desde esta pantalla.";}else{notice.querySelector("strong").textContent="Monitor automático activo";notice.querySelector("span").textContent="Revisa la bandeja cada minuto. Solo guarda oficios finales firmados y nunca sobrescribe un destino.";}contractEl("contractLastRun").textContent=data.last_scan?`Última revisión: ${contractDate(data.last_scan)}`:`Activado: ${contractDate(data.initialized_at)}`;contractEl("contractMeta").textContent=`${summary.candidates||0} oficio(s) contractual(es) · ${summary.archived||0} archivado(s) · ${summary.excluded||0} PDF excluido(s) por las reglas`;const patterns=contractEl("contractLearnedPatterns");patterns.innerHTML=(data.learned_patterns||[]).length?(data.learned_patterns||[]).map(item=>`<span class="contract-pattern">${contractEsc(item.descriptor)} <b>${contractEsc(item.count)}</b></span>`).join(""):'<span class="muted">Aún no hay nombres institucionales disponibles para aprender.</span>';contractRenderRows();}
async function contractLoad(silent=false){if(CONTRACTS.busy)return;try{const data=await contractApi("/api/contratos_agencia_fiscal/status");contractApply(data);}catch(error){const notice=contractEl("contractNotice");notice.classList.add("bad");notice.querySelector("strong").textContent="No se pudo consultar el monitor";notice.querySelector("span").textContent=error.message;if(!silent)alert(error.message);}}
async function contractScan(){if(CONTRACTS.busy)return;CONTRACTS.busy=true;const button=contractEl("contractScan");button.disabled=true;button.textContent="Analizando oficios...";try{const data=await contractApi("/api/contratos_agencia_fiscal/scan",{});contractApply(data);}catch(error){alert(error.message);}finally{CONTRACTS.busy=false;button.disabled=false;button.textContent="Revisar ahora";}}
async function contractSaveConfig(){if(CONTRACTS.busy)return;CONTRACTS.busy=true;try{const data=await contractApi("/api/contratos_agencia_fiscal/config",contractPayload());contractApplyConfig(data.config);await contractLoad(true);}catch(error){alert(error.message);}finally{CONTRACTS.busy=false;}}
function contractInit(){if(!contractEl("contractRows"))return;contractEl("contractRefresh").addEventListener("click",()=>contractLoad());contractEl("contractScan").addEventListener("click",contractScan);contractEl("contractSaveConfig").addEventListener("click",contractSaveConfig);contractLoad(true);setInterval(()=>contractLoad(true),30000);}
contractInit();

// Matriz de Préstamos: captura SGI, OCR local, revisión y escritura transaccional.
const MATRIX={draftId:"",busy:false,config:null};
const matrixEl=id=>document.getElementById(id);
const matrixEsc=value=>String(value??"").replace(/[&<>"']/g,char=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#039;"}[char]));
async function matrixJson(path,payload){const options=payload===undefined?{}:{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)};const response=await fetch(path,options);const data=await response.json();if(!response.ok||data.ok===false)throw new Error(data.error||`Error HTTP ${response.status}`);return data;}
function matrixNotice(message,type="",score="—"){const box=matrixEl("matrixNotice");box.className=`matrix-notice ${type}`.trim();box.querySelector("span").textContent=message;matrixEl("matrixScore").textContent=score;}
function matrixSetBusy(value,message="Analizando la captura localmente..."){MATRIX.busy=value;matrixEl("matrixCapture").disabled=value;matrixEl("matrixSave").disabled=value||!MATRIX.draftId;if(value)matrixNotice(message,"warn","OCR");}
function matrixSet(id,value){const element=matrixEl(id);if(element)element.value=value??"";}
function matrixCorrespondentProof(classification={}){const box=matrixEl("matrixCorrespondentEvidence");if(!box)return;const evidence=(classification.evidence||[]).join(" · ")||"El corresponsal se validará con la guía y la matriz histórica.";box.className=`matrix-notice ${classification.correspondent==="REVISAR"?"warn":"ok"}`;box.querySelector("span").textContent=evidence;box.querySelector("b").textContent=classification.confidence!==undefined?`${classification.confidence}%`:"—";}
function matrixPopulate(draft){MATRIX.draftId=draft.id||"";const e=draft.entry||{};matrixSet("matrixEntryDate",e.entry_date);matrixSet("matrixValueDate",e.value_date);matrixSet("matrixCurrency",e.currency||"USD");matrixSet("matrixLoan",e.reference_loan);matrixSet("matrixSwift",e.refer_swift);matrixSet("matrixLocal",e.local_operation);matrixSet("matrixLender",e.lender);matrixSet("matrixBorrower",e.borrower);matrixSet("matrixAmount",e.amount_usd??e.amount_other??"");matrixSet("matrixCorrespondent",e.correspondent);matrixSet("matrixBeneficiaryBank",e.beneficiary_bank);matrixCorrespondentProof({correspondent:e.correspondent,confidence:e.correspondent_confidence,evidence:e.correspondent_evidence});matrixSet("matrixAccountingDate",e.other_currency_accounting_date);matrixSet("matrixStatus",e.status||"Ingresado");matrixSet("matrixOffice",e.office_reference);matrixSet("matrixOrder",e.order_number);matrixSet("matrixNotes",e.notes);const warnings=draft.warnings||[];matrixEl("matrixWarnings").innerHTML=warnings.map(item=>`<li>${matrixEsc(item)}</li>`).join("");matrixNotice(warnings.length?"Extracción terminada con datos inferidos que debe revisar.":"Extracción terminada. Revise todos los campos antes de guardar.",warnings.length?"warn":"ok",`${draft.confidence||0}%`);matrixEl("matrixSave").disabled=false;if(draft.id){matrixEl("matrixPreview").src=`/api/matriz_prestamos/capture/${encodeURIComponent(draft.id)}?t=${Date.now()}`;matrixEl("matrixPreview").classList.add("on");}}
function matrixPayload(){const currency=matrixEl("matrixCurrency").value;const amount=Number(matrixEl("matrixAmount").value);return {entry_date:matrixEl("matrixEntryDate").value,reference_loan:matrixEl("matrixLoan").value.trim(),refer_swift:matrixEl("matrixSwift").value.trim().toUpperCase(),local_operation:matrixEl("matrixLocal").value.trim().toUpperCase(),lender:matrixEl("matrixLender").value.trim().toUpperCase(),borrower:matrixEl("matrixBorrower").value.trim().toUpperCase(),amount_usd:currency==="USD"?amount:null,amount_other:currency!=="USD"?amount:null,currency,correspondent:matrixEl("matrixCorrespondent").value.trim(),beneficiary_bank:matrixEl("matrixBeneficiaryBank").value.trim(),other_currency_accounting_date:matrixEl("matrixAccountingDate").value,value_date:matrixEl("matrixValueDate").value,status:matrixEl("matrixStatus").value,notes:matrixEl("matrixNotes").value.trim(),office_reference:matrixEl("matrixOffice").value,order_number:matrixEl("matrixOrder").value};}
function matrixReset(){MATRIX.draftId="";matrixEl("matrixCapture").value="";matrixEl("matrixPreview").removeAttribute("src");matrixEl("matrixPreview").classList.remove("on");["matrixEntryDate","matrixValueDate","matrixLoan","matrixSwift","matrixLocal","matrixLender","matrixBorrower","matrixAmount","matrixCorrespondent","matrixBeneficiaryBank","matrixAccountingDate","matrixOffice","matrixOrder","matrixNotes"].forEach(id=>matrixSet(id,""));matrixSet("matrixCurrency","USD");matrixSet("matrixStatus","Ingresado");matrixEl("matrixWarnings").innerHTML="";matrixEl("matrixSave").disabled=true;matrixCorrespondentProof();matrixNotice("Seleccione una captura para comenzar.","","—");}
function matrixRenderHistory(history){const body=matrixEl("matrixHistory");if(!(history||[]).length){body.innerHTML='<tr><td colspan="7" class="matrix-empty">Aún no hay filas agregadas desde este módulo.</td></tr>';return;}body.innerHTML=history.map(item=>{const when=item.saved_at?new Date(item.saved_at).toLocaleString("es-EC"):"—";const amount=Number(item.amount||0).toLocaleString("es-EC",{minimumFractionDigits:2,maximumFractionDigits:2});return `<tr><td>${matrixEsc(when)}</td><td>${matrixEsc(item.username)}</td><td><b>${matrixEsc(item.reference_loan)}</b></td><td>${matrixEsc(item.refer_swift)}</td><td>${matrixEsc(item.value_date)}</td><td>${matrixEsc(item.currency)} ${matrixEsc(amount)}</td><td>${matrixEsc(item.row||"—")}</td></tr>`;}).join("");}
async function matrixLoad(silent=false){try{const data=await matrixJson("/api/matriz_prestamos/status");MATRIX.config=data.config;matrixEl("matrixWorkbookPath").textContent=`Matriz: ${data.config.workbook}${data.workbook_exists?" · disponible":" · no encontrada"}`;matrixRenderHistory(data.history||[]);}catch(error){matrixEl("matrixWorkbookPath").textContent=error.message;if(!silent)alert(error.message);}}
async function matrixAnalyze(file){if(!file||MATRIX.busy)return;matrixSetBusy(true);const form=new FormData();form.append("capture",file);try{const response=await fetch("/api/matriz_prestamos/analyze",{method:"POST",body:form});const data=await response.json();if(!response.ok||data.ok===false)throw new Error(data.error||`Error HTTP ${response.status}`);matrixPopulate(data.draft);}catch(error){matrixNotice(error.message,"bad","Error");MATRIX.draftId="";matrixEl("matrixSave").disabled=true;}finally{matrixSetBusy(false);matrixEl("matrixSave").disabled=!MATRIX.draftId;}}
async function matrixSave(){if(!MATRIX.draftId||MATRIX.busy)return;const entry=matrixPayload();if(!confirm(`Se agregará ${entry.refer_swift||"esta operación"} a la matriz de préstamos. ¿Confirma que revisó los datos?`))return;matrixSetBusy(true,"Creando respaldo y actualizando el Excel...");try{const data=await matrixJson("/api/matriz_prestamos/save",{draft_id:MATRIX.draftId,entry});matrixNotice(`Guardado correctamente en la fila ${data.saved.row}. Se creó un respaldo anterior.`,"ok","Listo");MATRIX.draftId="";matrixEl("matrixSave").disabled=true;await matrixLoad(true);}catch(error){matrixNotice(error.message,"bad","Error");}finally{matrixSetBusy(false);matrixEl("matrixSave").disabled=!MATRIX.draftId;}}
async function matrixChangeConfig(){const current=MATRIX.config&&MATRIX.config.workbook||"";const workbook=prompt("Ruta completa del archivo Prestamo Realizados.xlsx",current);if(workbook===null)return;try{const data=await matrixJson("/api/matriz_prestamos/config",{workbook});MATRIX.config=data.config;await matrixLoad(true);}catch(error){alert(error.message);}}
let matrixClassifyTimer=null;
async function matrixReclassify(){if(!MATRIX.draftId||MATRIX.busy)return;try{const data=await matrixJson("/api/corresponsales/classify",{loan_reference:matrixEl("matrixLoan").value,lender:matrixEl("matrixLender").value,borrower:matrixEl("matrixBorrower").value,currency:matrixEl("matrixCurrency").value,beneficiary_bank:matrixEl("matrixBeneficiaryBank").value});const c=data.classification||{};matrixSet("matrixCorrespondent",c.correspondent);matrixCorrespondentProof(c);}catch(error){matrixCorrespondentProof({correspondent:"REVISAR",confidence:0,evidence:[error.message]});}}
function matrixScheduleReclassify(){clearTimeout(matrixClassifyTimer);matrixClassifyTimer=setTimeout(matrixReclassify,350);}
function matrixInit(){if(!matrixEl("matrixDrop"))return;const drop=matrixEl("matrixDrop"),input=matrixEl("matrixCapture");drop.addEventListener("click",()=>input.click());input.addEventListener("change",()=>matrixAnalyze(input.files&&input.files[0]));drop.addEventListener("dragover",event=>{event.preventDefault();drop.classList.add("drag");});drop.addEventListener("dragleave",()=>drop.classList.remove("drag"));drop.addEventListener("drop",event=>{event.preventDefault();drop.classList.remove("drag");const file=event.dataTransfer.files&&event.dataTransfer.files[0];if(file)matrixAnalyze(file);});matrixEl("matrixSave").addEventListener("click",matrixSave);matrixEl("matrixReset").addEventListener("click",matrixReset);matrixEl("matrixConfigButton").addEventListener("click",matrixChangeConfig);["matrixLoan","matrixLender","matrixBorrower","matrixCurrency"].forEach(id=>matrixEl(id).addEventListener("change",matrixScheduleReclassify));matrixLoad(true);}
matrixInit();
const CONC_VIEWS=["cargar","resultados","ajustes","quipux","condonados"];
function irA(v){const isConc=CONC_VIEWS.includes(v);
  if(["resultados","ajustes","quipux"].includes(v)&&!ACTIVE_CONCILIATION)limpiarConciliacionActiva();
  document.querySelectorAll(".nav .it").forEach(it=>it.classList.toggle("act",it.dataset.concRoot==="true"?isConc:it.dataset.v===v));
  ["control","cargar","resultados","ajustes","quipux","condonados","agenda","archivoquipux","contratosagencia","matrizprestamos","activaciones","respuestas","firmados","comprobantes"].forEach(s=>document.getElementById("v-"+s).classList.toggle("hidden",s!==v));
  document.getElementById("concModuleHead").classList.toggle("hidden",!isConc);
  document.querySelectorAll("#concFlowTabs [data-conc-view]").forEach(button=>{const active=button.dataset.concView===v;button.classList.toggle("act",active);button.setAttribute("aria-selected",String(active));});
}
function abrirVista(v){irA(v);if(v==="control")controlLoad();if(v==="agenda")agendaLoad();if(v==="archivoquipux")quipuxLoad();if(v==="contratosagencia")contractLoad();if(v==="matrizprestamos")matrixLoad();if(v==="respuestas")mailLoadExamples();if(v==="activaciones")activationInit();}
document.getElementById("nav").addEventListener("click",e=>{const it=e.target.closest(".it");if(it)abrirVista(it.dataset.v);});
document.querySelectorAll("#concFlowTabs [data-conc-view]").forEach(button=>button.addEventListener("click",()=>abrirVista(button.dataset.concView)));
const systemName="Gestor Integral de Deuda Externa Pública - Banco Central del Ecuador";
document.title=systemName;const systemTitle=document.querySelector(".topbrand h1 b");if(systemTitle)systemTitle.textContent=systemName;
function actualizarFechaOperativa(){const now=new Date();const full=new Intl.DateTimeFormat("es-EC",{day:"2-digit",month:"long",year:"numeric"}).format(now);const month=new Intl.DateTimeFormat("es-EC",{month:"long",year:"numeric"}).format(now);const pb=document.getElementById("periodoBadge");pb.textContent=full.charAt(0).toUpperCase()+full.slice(1);pb.className="badge set";const px=document.getElementById("perBox");px.style.display="block";px.textContent=month.charAt(0).toUpperCase()+month.slice(1);}
actualizarFechaOperativa();
limpiarConciliacionActiva();
abrirVista("control");
</script></body></html>"""


@app.route("/logo")
def logo():
    """Sirve el logo OFICIAL del BCE si el usuario deja un archivo
    'logo_bce.(png|jpg|svg)' en la carpeta de la app; si no, 404 y la
    interfaz usa el emblema SVG por defecto."""
    import glob
    mimes = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
             "gif": "image/gif", "svg": "image/svg+xml", "webp": "image/webp"}
    # Acepta logo_bce.png, y también nombres con doble extensión (logo_bce.png.png)
    for ruta in sorted(glob.glob(os.path.join(BASE_DIR, "logo_bce*"))):
        ext = ruta.rsplit(".", 1)[-1].lower()
        if ext in mimes:
            return send_file(ruta, mimetype=mimes[ext])
    from flask import abort
    abort(404)


@app.route("/")
def home():
    display_name = session.get("display_name") or session.get("username") or "Usuario"
    return render_template_string(
        PANEL_HTML,
        user_name=display_name,
        user_role=session.get("role") or "Usuario",
        user_initial=display_name[:1].upper(),
    )


def abrir_navegador():
    webbrowser.open_new("http://127.0.0.1:5001/")


if __name__ == "__main__":
    init_db()
    print("=" * 60)
    print("  PANEL CONCILIACIÓN BCE vs MEF")
    print("=" * 60)
    print(f"  BD:  {DB_PATH}")
    print(f"  URL: http://127.0.0.1:5001/")
    print("=" * 60)
    Timer(1.5, abrir_navegador).start()
    app.run(host="127.0.0.1", port=5001, debug=False)



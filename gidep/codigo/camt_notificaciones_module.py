# -*- coding: utf-8 -*-
"""Modulo CAMT.054 para el Sistema de Conciliacion BCE.

Registra rutas API para:
- Escanear CAMT.054 DBIT.
- Cruzar Refer Swift contra la matriz Prestamo Realizados.
- Mostrar pagos diarios y estado de debito.
- Generar eventos JSON para Power Automate, Teams y correo.
"""
from __future__ import annotations

import csv
import json
import os
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from flask import jsonify, request


CAMT_ROOT = Path(r"Z:\GISI\SSFI\SWIFT\Mensajes Giros del")
WORKBOOK_PATH = Path(r"C:\Users\bromo\OneDrive - BANCO CENTRAL DEL ECUADOR\Prestamo Realizados.xlsx")
QUEUE_DIR = Path(r"C:\Users\bromo\OneDrive - BANCO CENTRAL DEL ECUADOR\CAMT PowerAutomate Queue")
NOTIFICATION_EMAIL = "bromo@bce.ec"
SHARED_DATA_DIR = Path(
    r"C:\Users\bromo\Documents\Codex\2026-06-29\necesito-que-leas-mi-disco-local-2\outputs\camt_gestor\data"
)

CAMT_RE = re.compile(
    r"^(?P<prefix>[O0])_"
    r"(?P<amount>[^_]+)_"
    r"camt\.054\.001\.08_"
    r"(?P<operation>[^_]+)_"
    r"(?P<message>[^_]+)_"
    r"(?P<entry>[^_]+)_"
    r"(?P<bic>[^_]+)_"
    r"DBIT_"
    r"\((?P<sequence>[^)]+)\)_"
    r"\((?P<value_date>\d{4}-\d{2}-\d{2})\)"
    r"\.pdf$",
    re.IGNORECASE,
)

DEBIT_FIELDS = [
    "operation_code",
    "amount",
    "value_date",
    "message_id",
    "entry_id",
    "bic",
    "sequence",
    "file_name",
    "file_path",
    "last_write_time",
]

MATCH_HEADERS = [
    "sheet",
    "row",
    "reference_loan",
    "refer_swift",
    "matrix_amount_usd",
    "matrix_value_date",
    "current_status",
    "gestor_status",
    "match_source",
    "camt_amount",
    "camt_value_date",
    "camt_message_id",
    "camt_sequence",
    "camt_pdf",
]

DUE_HEADERS = [
    "sheet",
    "row",
    "reference_loan",
    "refer_swift",
    "operation_local",
    "lender",
    "borrower",
    "amount_usd",
    "amount_other_currencies",
    "currency",
    "corresponsal",
    "value_date",
    "status",
    "gestor_status",
    "match_source",
    "camt_amount",
    "camt_value_date",
    "camt_message_id",
    "camt_sequence",
    "camt_pdf",
    "notes",
]


@dataclass(frozen=True)
class DebitNotice:
    operation_code: str
    amount: str
    value_date: str
    message_id: str
    entry_id: str
    bic: str
    sequence: str
    file_name: str
    file_path: str
    last_write_time: str


def quitar_acentos(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in text if not unicodedata.combining(char))


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", quitar_acentos(value)).strip().upper()


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value or "").strip()).strip("_") or "CAMT"


def parse_decimal(value: str) -> Decimal:
    return Decimal(str(value).replace(".", "").replace(",", "."))


def parse_notice(path: Path) -> DebitNotice | None:
    match = CAMT_RE.match(path.name)
    if not match:
        return None

    groups = match.groupdict()
    return DebitNotice(
        operation_code=groups["operation"],
        amount=str(parse_decimal(groups["amount"])),
        value_date=groups["value_date"],
        message_id=groups["message"],
        entry_id=groups["entry"],
        bic=groups["bic"],
        sequence=groups["sequence"],
        file_name=path.name,
        file_path=str(path),
        last_write_time=datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
    )


def iter_camt_pdfs(root: Path) -> Iterable[Path]:
    for path in root.rglob("*.pdf"):
        name = path.name.lower()
        if path.parent.name.lower() == "camt.054 dbit" and "camt.054" in name and "dbit" in name:
            yield path


def read_json(path: Path, default: dict[str, Any]) -> dict[str, Any]:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def load_seen(path: Path) -> set[str]:
    return set(read_json(path, {"seen_files": []}).get("seen_files", []))


def save_seen(path: Path, seen: set[str]) -> None:
    write_json(path, {"updated_at": datetime.now().isoformat(timespec="seconds"), "seen_files": sorted(seen)})


def load_events(path: Path) -> set[str]:
    return set(read_json(path, {"events": []}).get("events", []))


def save_events(path: Path, events: set[str]) -> None:
    write_json(path, {"updated_at": datetime.now().isoformat(timespec="seconds"), "events": sorted(events)})


def write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def scan_camt(root: Path, state_path: Path, output_csv: Path, *, include_all: bool, commit: bool) -> list[DebitNotice]:
    if not root.exists():
        raise FileNotFoundError(f"No existe la carpeta CAMT: {root}")

    seen = load_seen(state_path)
    current_seen = set(seen)
    notices: list[DebitNotice] = []

    for pdf in iter_camt_pdfs(root):
        notice = parse_notice(pdf)
        if notice is None:
            continue
        if include_all or notice.file_path not in seen:
            notices.append(notice)
        current_seen.add(notice.file_path)

    notices.sort(key=lambda item: (item.value_date, item.operation_code, item.file_name))
    write_csv(output_csv, DEBIT_FIELDS, [asdict(notice) for notice in notices])
    if commit:
        save_seen(state_path, current_seen)
    return notices


def excel_serial_to_date(value: int | float) -> date:
    return date(1899, 12, 30) + timedelta(days=int(value))


def as_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return excel_serial_to_date(value)

    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def cell(sheet: Any, row: int, headers: dict[str, int], *names: str) -> Any:
    for name in names:
        col = headers.get(normalize(name))
        if col:
            return sheet.cell(row, col).value
    return ""


def load_debits(path: Path) -> dict[str, dict[str, str]]:
    return {normalize(row.get("operation_code")): row for row in read_csv(path) if normalize(row.get("operation_code"))}


def available_value_dates(workbook_path: Path) -> list[dict[str, object]]:
    if not workbook_path.exists():
        return []
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    counts: Counter[str] = Counter()
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            headers = {normalize(cell_obj.value): idx for idx, cell_obj in enumerate(sheet[1], start=1)}
            value_col = headers.get("VALUE DATE")
            if not value_col:
                continue
            for row_idx in range(2, sheet.max_row + 1):
                payment_date = as_date(sheet.cell(row_idx, value_col).value)
                if payment_date:
                    counts[payment_date.isoformat()] += 1
    finally:
        workbook.close()
    return [{"date": key, "count": counts[key]} for key in sorted(counts)]


def find_due_payments(workbook_path: Path, output_csv: Path, target_date: date, camt_csv: Path) -> list[dict[str, Any]]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"No existe la matriz Excel: {workbook_path}")

    debits = load_debits(camt_csv)
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            headers = {normalize(cell_obj.value): idx for idx, cell_obj in enumerate(sheet[1], start=1)}
            value_col = headers.get("VALUE DATE")
            if not value_col:
                continue

            for row_idx in range(2, sheet.max_row + 1):
                payment_date = as_date(sheet.cell(row_idx, value_col).value)
                if payment_date != target_date:
                    continue

                refer_swift = cell(sheet, row_idx, headers, "Refer Swift")
                debit = debits.get(normalize(refer_swift))
                rows.append(
                    {
                        "sheet": sheet_name,
                        "row": row_idx,
                        "reference_loan": cell(sheet, row_idx, headers, "Reference Loan"),
                        "refer_swift": refer_swift,
                        "operation_local": cell(sheet, row_idx, headers, "Operacion Local", "Operación Local"),
                        "lender": cell(sheet, row_idx, headers, "Lender"),
                        "borrower": cell(sheet, row_idx, headers, "Borrower"),
                        "amount_usd": cell(sheet, row_idx, headers, "Amount USD"),
                        "amount_other_currencies": cell(sheet, row_idx, headers, "Amount Other Currencies"),
                        "currency": cell(sheet, row_idx, headers, "Currency"),
                        "corresponsal": cell(sheet, row_idx, headers, "Corresponsal"),
                        "value_date": payment_date.isoformat(),
                        "status": cell(sheet, row_idx, headers, "Status"),
                        "gestor_status": "Debitado" if debit else "No debitado",
                        "match_source": "CAMT.054 DBIT" if debit else "",
                        "camt_amount": debit.get("amount", "") if debit else "",
                        "camt_value_date": debit.get("value_date", "") if debit else "",
                        "camt_message_id": debit.get("message_id", "") if debit else "",
                        "camt_sequence": debit.get("sequence", "") if debit else "",
                        "camt_pdf": debit.get("file_path", "") if debit else "",
                        "notes": cell(sheet, row_idx, headers, "Notes"),
                    }
                )
    finally:
        workbook.close()

    write_csv(output_csv, DUE_HEADERS, rows)
    return rows


def match_workbook(workbook_path: Path, debits_csv: Path, output_csv: Path) -> list[dict[str, Any]]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"No existe la matriz Excel: {workbook_path}")
    if not debits_csv.exists():
        raise FileNotFoundError(f"No existe el CSV de debitos: {debits_csv}")

    debits = load_debits(debits_csv)
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    rows: list[dict[str, Any]] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            headers = {normalize(cell_obj.value): idx for idx, cell_obj in enumerate(sheet[1], start=1)}
            refer_col = headers.get("REFER SWIFT")
            if not refer_col:
                continue

            for row_idx in range(2, sheet.max_row + 1):
                swift_code = normalize(sheet.cell(row_idx, refer_col).value)
                if not swift_code or swift_code in {"N/A", "NA"}:
                    continue
                debit = debits.get(swift_code)
                if debit is None:
                    continue

                rows.append(
                    {
                        "sheet": sheet_name,
                        "row": row_idx,
                        "reference_loan": cell(sheet, row_idx, headers, "Reference Loan"),
                        "refer_swift": swift_code,
                        "matrix_amount_usd": cell(sheet, row_idx, headers, "Amount USD"),
                        "matrix_value_date": cell(sheet, row_idx, headers, "Value Date"),
                        "current_status": cell(sheet, row_idx, headers, "Status"),
                        "gestor_status": "Debitado",
                        "match_source": "CAMT.054 DBIT",
                        "camt_amount": debit.get("amount", ""),
                        "camt_value_date": debit.get("value_date", ""),
                        "camt_message_id": debit.get("message_id", ""),
                        "camt_sequence": debit.get("sequence", ""),
                        "camt_pdf": debit.get("file_path", ""),
                    }
                )
    finally:
        workbook.close()

    write_csv(output_csv, MATCH_HEADERS, rows)
    return rows


def camt_event_key(row: dict[str, str]) -> str:
    return "|".join(
        [
            "camt_debited",
            normalize(row.get("refer_swift")),
            normalize(row.get("camt_message_id")),
            normalize(row.get("camt_sequence")),
            str(row.get("camt_value_date", "")).strip(),
        ]
    )


def legacy_camt_event_key(row: dict[str, str]) -> str:
    return "|".join(
        [
            normalize(row.get("refer_swift")),
            normalize(row.get("camt_message_id")),
            normalize(row.get("camt_sequence")),
            str(row.get("camt_value_date", "")).strip(),
        ]
    )


def due_event_key(row: dict[str, str]) -> str:
    return "|".join(
        [
            "daily_due",
            str(row.get("value_date", "")).strip(),
            normalize(row.get("refer_swift")),
            normalize(row.get("reference_loan")),
        ]
    )


def status_event_key(row: dict[str, str], new_status: str) -> str:
    return "|".join(
        [
            "status_update",
            normalize(new_status),
            str(row.get("value_date", "")).strip(),
            normalize(row.get("refer_swift")),
            normalize(row.get("reference_loan")),
        ]
    )


def export_camt_events(matches_csv: Path, queue_dir: Path, state_path: Path, *, force: bool = False) -> list[Path]:
    rows = read_csv(matches_csv)
    seen = load_events(state_path)
    new_seen = set(seen)
    queue_dir.mkdir(parents=True, exist_ok=True)
    exported: list[Path] = []
    run_stamp = datetime.now().strftime("%Y%m%d%H%M%S")

    for row in rows:
        key = camt_event_key(row)
        legacy_key = legacy_camt_event_key(row)
        if (key in seen or legacy_key in seen) and not force:
            continue

        operation_code = row.get("refer_swift", "")
        reference_loan = row.get("reference_loan", "")
        payload = {
            "event_type": "camt_debited",
            "event_key": key,
            "detected_at": datetime.now().isoformat(timespec="seconds"),
            "target_workbook": "Prestamo Realizados.xlsx",
            "target_sheet": row.get("sheet", "2026"),
            "target_table": "Table2" if row.get("sheet") == "2026" else "Tabla1",
            "key_column": "Refer Swift",
            "operation_code": operation_code,
            "new_status": "Debitado",
            "reference_loan": reference_loan,
            "previous_status": row.get("current_status", ""),
            "camt_amount": row.get("camt_amount", ""),
            "camt_value_date": row.get("camt_value_date", ""),
            "camt_message_id": row.get("camt_message_id", ""),
            "camt_sequence": row.get("camt_sequence", ""),
            "camt_pdf": row.get("camt_pdf", ""),
            "teams_title": f"Prestamo debitado CAMT.054 - {operation_code}",
            "teams_message": (
                f"Ya fue debitado el prestamo {reference_loan} con referencia Swift {operation_code}. "
                f"Monto CAMT: {row.get('camt_amount', '')}. Fecha valor CAMT: {row.get('camt_value_date', '')}. "
                f"Mensaje CAMT: {row.get('camt_message_id', '')}. Secuencia: {row.get('camt_sequence', '')}."
            ),
            "email_to": NOTIFICATION_EMAIL,
            "email_subject": f"Prestamo debitado CAMT.054 - {operation_code}",
            "email_body": (
                f"Ya fue debitado el prestamo {reference_loan} con referencia Swift {operation_code}.\n\n"
                f"Monto CAMT: {row.get('camt_amount', '')}\n"
                f"Fecha valor CAMT: {row.get('camt_value_date', '')}\n"
                f"Mensaje CAMT: {row.get('camt_message_id', '')}\n"
                f"Secuencia: {row.get('camt_sequence', '')}\n"
                f"PDF: {row.get('camt_pdf', '')}"
            ),
        }
        suffix = f"_{run_stamp}" if force else ""
        target = queue_dir / (
            f"camt_debited_{safe_name(operation_code)}_{safe_name(payload['camt_message_id'])}_"
            f"{safe_name(payload['camt_sequence'])}{suffix}.json"
        )
        target.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        exported.append(target)
        new_seen.add(key)

    save_events(state_path, new_seen)
    return exported


def export_due_events(due_csv: Path, queue_dir: Path, state_path: Path, *, force: bool = False) -> list[Path]:
    rows = read_csv(due_csv)
    seen = load_events(state_path)
    new_seen = set(seen)
    queue_dir.mkdir(parents=True, exist_ok=True)
    exported: list[Path] = []
    run_stamp = datetime.now().strftime("%Y%m%d%H%M%S")

    for row in rows:
        key = due_event_key(row)
        if key in seen and not force:
            continue

        operation_code = row.get("refer_swift", "")
        reference_loan = row.get("reference_loan", "")
        amount = row.get("amount_usd") or row.get("amount_other_currencies") or ""
        payload = {
            "event_type": "daily_due",
            "event_key": key,
            "detected_at": datetime.now().isoformat(timespec="seconds"),
            "target_workbook": "Prestamo Realizados.xlsx",
            "target_sheet": row.get("sheet", "2026"),
            "target_table": "Table2" if row.get("sheet") == "2026" else "Tabla1",
            "reference_loan": reference_loan,
            "operation_code": operation_code,
            "operation_local": row.get("operation_local", ""),
            "lender": row.get("lender", ""),
            "borrower": row.get("borrower", ""),
            "amount": amount,
            "currency": row.get("currency", ""),
            "corresponsal": row.get("corresponsal", ""),
            "value_date": row.get("value_date", ""),
            "status": row.get("status", ""),
            "gestor_status": row.get("gestor_status", "No debitado"),
            "match_source": row.get("match_source", ""),
            "camt_message_id": row.get("camt_message_id", ""),
            "camt_sequence": row.get("camt_sequence", ""),
            "camt_pdf": row.get("camt_pdf", ""),
            "teams_title": f"Pago programado hoy - {operation_code or reference_loan}",
            "teams_message": (
                f"Hoy corresponde pago de {reference_loan} con referencia Swift {operation_code}. "
                f"Monto: {amount} {row.get('currency', '')}. Fecha valor: {row.get('value_date', '')}. "
                f"Estado gestor: {row.get('gestor_status', 'No debitado')}. "
                f"Estado actual en matriz: {row.get('status', '')}."
            ),
            "email_to": NOTIFICATION_EMAIL,
            "email_subject": f"Pago programado hoy - {operation_code or reference_loan}",
            "email_body": (
                f"Hoy corresponde pago de {reference_loan} con referencia Swift {operation_code}.\n\n"
                f"Monto: {amount} {row.get('currency', '')}\n"
                f"Fecha valor: {row.get('value_date', '')}\n"
                f"Estado gestor: {row.get('gestor_status', 'No debitado')}\n"
                f"Estado actual en matriz: {row.get('status', '')}\n"
                f"Lender: {row.get('lender', '')}\n"
                f"Borrower: {row.get('borrower', '')}\n"
                f"Corresponsal: {row.get('corresponsal', '')}"
            ),
        }
        suffix = f"_{run_stamp}" if force else ""
        target = queue_dir / f"daily_due_{safe_name(payload['value_date'])}_{safe_name(operation_code or reference_loan)}{suffix}.json"
        target.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        exported.append(target)
        new_seen.add(key)

    save_events(state_path, new_seen)
    return exported


def export_status_update_events(rows: list[dict[str, str]], queue_dir: Path, state_path: Path, *, new_status: str) -> list[Path]:
    seen = load_events(state_path)
    new_seen = set(seen)
    queue_dir.mkdir(parents=True, exist_ok=True)
    exported: list[Path] = []
    run_stamp = datetime.now().strftime("%Y%m%d%H%M%S")

    for row in rows:
        operation_code = row.get("refer_swift", "")
        if not operation_code:
            continue
        key = status_event_key(row, new_status)
        payload = {
            "event_type": "status_update",
            "event_key": key,
            "detected_at": datetime.now().isoformat(timespec="seconds"),
            "target_workbook": "Prestamo Realizados.xlsx",
            "target_sheet": row.get("sheet", "2026"),
            "target_table": "Table2" if row.get("sheet") == "2026" else "Tabla1",
            "key_column": "Refer Swift",
            "operation_code": operation_code,
            "new_status": new_status,
            "reference_loan": row.get("reference_loan", ""),
            "previous_status": row.get("status", ""),
            "value_date": row.get("value_date", ""),
            "borrower": row.get("borrower", ""),
            "lender": row.get("lender", ""),
            "amount": row.get("amount_usd") or row.get("amount_other_currencies") or "",
            "currency": row.get("currency", ""),
            "gestor_status": row.get("gestor_status", ""),
            "teams_title": f"Actualizar Status a {new_status} - {operation_code}",
            "teams_message": (
                f"Solicitud de actualizacion en matriz: cambiar Status a {new_status} para "
                f"{row.get('reference_loan', '')} con referencia Swift {operation_code}. "
                f"Fecha valor: {row.get('value_date', '')}. Estado gestor: {row.get('gestor_status', '')}."
            ),
            "email_to": NOTIFICATION_EMAIL,
            "email_subject": f"Actualizar Status a {new_status} - {operation_code}",
            "email_body": (
                f"Solicitud de actualizacion en matriz.\n\nNuevo Status: {new_status}\n"
                f"Prestamo: {row.get('reference_loan', '')}\nRefer Swift: {operation_code}\n"
                f"Borrower: {row.get('borrower', '')}\nFecha valor: {row.get('value_date', '')}\n"
                f"Estado gestor: {row.get('gestor_status', '')}\nEstado anterior matriz: {row.get('status', '')}"
            ),
        }
        target = queue_dir / f"status_update_{safe_name(new_status)}_{safe_name(operation_code)}_{run_stamp}.json"
        target.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        exported.append(target)
        new_seen.add(key)

    save_events(state_path, new_seen)
    return exported


class CamtPaths:
    def __init__(self, base_dir: str | os.PathLike[str]):
        data_dir = SHARED_DATA_DIR if SHARED_DATA_DIR.exists() else Path(base_dir) / "camt_notificaciones_data"
        self.data_dir = data_dir
        self.scan_state = data_dir / "scan_state.json"
        self.event_state = data_dir / "event_state.json"
        self.due_event_state = data_dir / "due_event_state.json"
        self.status_event_state = data_dir / "status_event_state.json"
        self.debits_csv = data_dir / "camt_ledger.csv"
        self.new_debits_csv = data_dir / "new_debits.csv"
        self.matches_csv = data_dir / "matches.csv"
        self.due_csv = data_dir / "due_payments.csv"


def parse_target_date(raw: str | None = None) -> date:
    raw = (raw or request.values.get("fecha") or "").strip()
    if raw:
        try:
            return datetime.strptime(raw, "%Y-%m-%d").date()
        except ValueError:
            pass
    return date.today()


def state_count(path: Path) -> int:
    return len(load_events(path))


def row_public_status(row: dict[str, Any]) -> str:
    if "JP MORGAN" in normalize(row.get("corresponsal")):
        return "JP Morgan Access"
    return str(row.get("gestor_status") or "No debitado")


def build_summary(paths: CamtPaths, target: date) -> dict[str, Any]:
    if not paths.debits_csv.exists() and CAMT_ROOT.exists():
        scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)

    due_rows = find_due_payments(WORKBOOK_PATH, paths.due_csv, target, paths.debits_csv) if WORKBOOK_PATH.exists() else []
    due_public = [{**row, "display_status": row_public_status(row)} for row in due_rows]
    matches = read_csv(paths.matches_csv)
    debits = read_csv(paths.debits_csv)
    queue_count = len(list(QUEUE_DIR.glob("*.json"))) if QUEUE_DIR.exists() else 0
    due_debited = sum(1 for row in due_rows if normalize(row.get("gestor_status")) == "DEBITADO")

    return {
        "selected_date": target.isoformat(),
        "paths": {
            "camt_root": str(CAMT_ROOT),
            "workbook": str(WORKBOOK_PATH),
            "queue": str(QUEUE_DIR),
            "data_dir": str(paths.data_dir),
        },
        "exists": {
            "camt_root": CAMT_ROOT.exists(),
            "workbook": WORKBOOK_PATH.exists(),
            "queue": QUEUE_DIR.exists(),
        },
        "counts": {
            "debits": len(debits),
            "matches": len(matches),
            "due": len(due_rows),
            "due_debited": due_debited,
            "due_not_debited": len(due_rows) - due_debited,
            "queue": queue_count,
            "exported_camt": state_count(paths.event_state),
            "exported_due": state_count(paths.due_event_state),
        },
        "available_dates": available_value_dates(WORKBOOK_PATH),
        "due": due_public[:200],
        "notification_email": NOTIFICATION_EMAIL,
    }


def register_camt_notificaciones_routes(app, base_dir):
    paths = CamtPaths(base_dir)

    @app.route("/api/camt_notificaciones/summary")
    def camt_summary():
        try:
            target = parse_target_date()
            return jsonify({"ok": True, **build_summary(paths, target)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

    @app.route("/api/camt_notificaciones/scan", methods=["POST"])
    def camt_scan_api():
        try:
            data = request.get_json(silent=True) or {}
            target = parse_target_date(data.get("fecha"))
            include_all = bool(data.get("include_all"))
            notices = scan_camt(
                CAMT_ROOT,
                paths.scan_state,
                paths.debits_csv if include_all else paths.new_debits_csv,
                include_all=include_all,
                commit=not include_all,
            )
            if not include_all:
                scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)
            return jsonify({"ok": True, "message": f"Escaneo completo: {len(notices)} CAMT detectados.", **build_summary(paths, target)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

    @app.route("/api/camt_notificaciones/run_all", methods=["POST"])
    def camt_run_all_api():
        try:
            target = parse_target_date((request.get_json(silent=True) or {}).get("fecha"))
            new_notices = scan_camt(CAMT_ROOT, paths.scan_state, paths.new_debits_csv, include_all=False, commit=False)
            scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)
            due_rows = find_due_payments(WORKBOOK_PATH, paths.due_csv, target, paths.debits_csv)
            due_events = export_due_events(paths.due_csv, QUEUE_DIR, paths.due_event_state, force=False)
            matches = match_workbook(WORKBOOK_PATH, paths.debits_csv, paths.matches_csv)
            camt_events = export_camt_events(paths.matches_csv, QUEUE_DIR, paths.event_state, force=False)
            if camt_events:
                scan_camt(CAMT_ROOT, paths.scan_state, paths.new_debits_csv, include_all=False, commit=True)
            return jsonify(
                {
                    "ok": True,
                    "message": (
                        f"Proceso completo: {len(due_rows)} pagos del dia, {len(due_events)} alertas diarias, "
                        f"{len(new_notices)} CAMT nuevos, {len(matches)} matches, {len(camt_events)} eventos CAMT."
                    ),
                    **build_summary(paths, target),
                }
            )
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

    @app.route("/api/camt_notificaciones/daily_due", methods=["POST"])
    def camt_daily_due_api():
        try:
            target = parse_target_date((request.get_json(silent=True) or {}).get("fecha"))
            scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)
            due_rows = find_due_payments(WORKBOOK_PATH, paths.due_csv, target, paths.debits_csv)
            exported = export_due_events(paths.due_csv, QUEUE_DIR, paths.due_event_state, force=False)
            return jsonify({"ok": True, "message": f"Pagos del dia: {len(due_rows)}; eventos nuevos: {len(exported)}.", **build_summary(paths, target)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

    @app.route("/api/camt_notificaciones/resend_daily_due", methods=["POST"])
    def camt_resend_due_api():
        try:
            target = parse_target_date((request.get_json(silent=True) or {}).get("fecha"))
            scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)
            find_due_payments(WORKBOOK_PATH, paths.due_csv, target, paths.debits_csv)
            exported = export_due_events(paths.due_csv, QUEUE_DIR, paths.due_event_state, force=True)
            return jsonify({"ok": True, "message": f"Reenvio de pagos creado: {len(exported)} eventos.", **build_summary(paths, target)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

    @app.route("/api/camt_notificaciones/resend_camt", methods=["POST"])
    def camt_resend_camt_api():
        try:
            target = parse_target_date((request.get_json(silent=True) or {}).get("fecha"))
            scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)
            match_workbook(WORKBOOK_PATH, paths.debits_csv, paths.matches_csv)
            exported = export_camt_events(paths.matches_csv, QUEUE_DIR, paths.event_state, force=True)
            return jsonify({"ok": True, "message": f"Reenvio CAMT creado: {len(exported)} eventos.", **build_summary(paths, target)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

    @app.route("/api/camt_notificaciones/mark_atendido", methods=["POST"])
    def camt_mark_atendido_api():
        try:
            data = request.get_json(silent=True) or {}
            target = parse_target_date(data.get("fecha"))
            refer_swift = normalize(data.get("refer_swift"))
            scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)
            due_rows = find_due_payments(WORKBOOK_PATH, paths.due_csv, target, paths.debits_csv)
            selected = [row for row in due_rows if normalize(row.get("refer_swift")) == refer_swift]
            exported = export_status_update_events(selected, QUEUE_DIR, paths.status_event_state, new_status="Atendido")
            return jsonify({"ok": True, "message": f"Atendido solicitado para {refer_swift}: {len(exported)} evento(s).", **build_summary(paths, target)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

    @app.route("/api/camt_notificaciones/mark_all_atendido", methods=["POST"])
    def camt_mark_all_atendido_api():
        try:
            target = parse_target_date((request.get_json(silent=True) or {}).get("fecha"))
            scan_camt(CAMT_ROOT, paths.scan_state, paths.debits_csv, include_all=True, commit=False)
            due_rows = find_due_payments(WORKBOOK_PATH, paths.due_csv, target, paths.debits_csv)
            exported = export_status_update_events(due_rows, QUEUE_DIR, paths.status_event_state, new_status="Atendido")
            return jsonify({"ok": True, "message": f"Atendido solicitado para {len(due_rows)} pago(s): {len(exported)} evento(s).", **build_summary(paths, target)})
        except Exception as exc:
            return jsonify({"ok": False, "error": str(exc)})

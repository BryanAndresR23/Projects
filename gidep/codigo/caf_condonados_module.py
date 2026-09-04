# -*- coding: utf-8 -*-
"""Modulo Flask para cargar condonados CAF en el reporte BCE.

Se integra con conciliacion_app.py mediante register_caf_condonados_routes().
"""
import json
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
from copy import copy
from datetime import datetime
from pathlib import Path

from flask import jsonify, request
from openpyxl import load_workbook

try:
    import PyPDF2
except ImportError:  # pragma: no cover - dependencia opcional
    PyPDF2 = None


HEADER_MORA_OM = "INTERES MORA (OTRAS MONEDAS)"
HEADER_MORA_USD = "INTERES MORA (USD)"
HEADER_CONDONADOS_OM = "CONDONADOS (OTRAS MONEDAS)"
HEADER_CONDONADOS_USD = "CONDONADOS (USD)"
HEADER_REFERENCIA = "NO. REFERENCIA PRESTAMO"
NUM_FORMAT = '#,##0.00;\\-#,##0.00;0.00'

MESES = [
    ("Enero", "Ene"),
    ("Febrero", "Feb"),
    ("Marzo", "Mar"),
    ("Abril", "Abr"),
    ("Mayo", "May"),
    ("Junio", "Jun"),
    ("Julio", "Jul"),
    ("Agosto", "Ago"),
    ("Septiembre", "Sep"),
    ("Octubre", "Oct"),
    ("Noviembre", "Nov"),
    ("Diciembre", "Dic"),
]

DEFAULT_YEAR = 2026
DEFAULT_MONTH = "Junio"
DEFAULT_ROOT = (
    "Z:\\GISI\\SSFI\\GESTI\u00d3N PAGOS INTERNACIONALES\\{anio}\\"
    "DEUDA EXTERNA P\u00daBLICA\\Acreedores Internacionales\\CAF\\Pagos"
)
DEFAULT_REPORT = (
    "Z:\\GISI\\SSFI\\GESTI\u00d3N PAGOS INTERNACIONALES\\{anio}\\"
    "DEUDA EXTERNA P\u00daBLICA\\Conciliaciones\\Reporte\\"
    "Reporte Conciliaci\u00f3n BCE - {anio}.xlsx"
)
CONFIG_PATH = Path(__file__).with_name("caf_condonados_config.json")


def quitar_acentos(texto):
    normalizado = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(char for char in normalizado if not unicodedata.combining(char))


def normalizar_texto(valor):
    return re.sub(r"\s+", " ", quitar_acentos(valor)).strip().upper()


def extraer_clave_cfa(texto):
    tokens = re.findall(r"CFA[-\s]?0*(\d+)", normalizar_texto(texto))
    if not tokens:
        return ""
    claves = []
    for token in tokens:
        clave = f"CFA-{int(token)}"
        if clave not in claves:
            claves.append(clave)
    return "|".join(claves)


def clave_set(clave):
    return {token for token in str(clave or "").split("|") if token}


def claves_coinciden(a, b):
    conjunto_a = clave_set(a)
    conjunto_b = clave_set(b)
    return bool(conjunto_a and conjunto_b and conjunto_a.intersection(conjunto_b))


def extraer_sigade(texto):
    match = re.search(r"\b(\d{8,9})\b", str(texto or ""))
    return match.group(1) if match else ""


def parse_decimal(valor):
    if valor in (None, ""):
        return 0.0
    if isinstance(valor, (int, float)):
        return round(float(valor), 2)

    texto = str(valor).strip()
    negativo = texto.startswith("-") or (texto.startswith("(") and texto.endswith(")"))
    texto = texto.strip("()").replace(" ", "").replace("$", "")

    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(".", "").replace(",", ".")

    try:
        numero = float(texto or "0")
    except ValueError:
        numero = 0.0
    return round(-numero if negativo else numero, 2)


def abs_decimal(valor):
    return abs(parse_decimal(valor))


def mes_abreviado(mes):
    for nombre, abreviado in MESES:
        if nombre.lower() == str(mes or "").lower():
            return abreviado
    return str(mes or "")[:3].title()


def construir_pagos_base(anio):
    return DEFAULT_ROOT.format(anio=anio)


def construir_reporte(anio):
    return DEFAULT_REPORT.format(anio=anio)


def construir_pagos_mes(pagos_caf_base, mes):
    return str(Path(pagos_caf_base) / mes)


def construir_hoja(mes, anio):
    return f"Giros al Exterior - {mes_abreviado(mes)}{anio}"


def corregir_ruta_red(ruta):
    if not ruta:
        return ruta

    candidatos = [
        ruta,
        ruta.replace("GESTION PAGOS INTERNACIONALES", "GESTI\u00d3N PAGOS INTERNACIONALES")
        .replace("DEUDA EXTERNA PUBLICA", "DEUDA EXTERNA P\u00daBLICA")
        .replace("Conciliacion BCE", "Conciliaci\u00f3n BCE"),
        ruta.replace("GESTI\u00c3\u0093N PAGOS INTERNACIONALES", "GESTI\u00d3N PAGOS INTERNACIONALES")
        .replace("DEUDA EXTERNA P\u00c3\u009aBLICA", "DEUDA EXTERNA P\u00daBLICA")
        .replace("Conciliaci\u00c3\u00b3n BCE", "Conciliaci\u00f3n BCE"),
    ]

    for candidato in candidatos:
        if os.path.exists(candidato):
            return candidato
    return candidatos[1]


def as_bool(value):
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "si", "yes", "on"}
    return bool(value)


def default_config():
    return {
        "anio": DEFAULT_YEAR,
        "mes": DEFAULT_MONTH,
        "pagos_caf_base": construir_pagos_base(DEFAULT_YEAR),
        "pagos_caf_dir": "",
        "reporte_path": construir_reporte(DEFAULT_YEAR),
        "sheet_name": "",
        "auto_paths": True,
        "crear_backup": True,
    }


def normalizar_configuracion(config):
    resultado = default_config()
    resultado.update({key: value for key, value in (config or {}).items() if value is not None})

    try:
        resultado["anio"] = int(resultado.get("anio") or DEFAULT_YEAR)
    except (TypeError, ValueError):
        resultado["anio"] = DEFAULT_YEAR

    meses_validos = {nombre for nombre, _abbr in MESES}
    if resultado.get("mes") not in meses_validos:
        resultado["mes"] = DEFAULT_MONTH

    resultado["auto_paths"] = as_bool(resultado.get("auto_paths", True))
    resultado["crear_backup"] = as_bool(resultado.get("crear_backup", True))

    if resultado["auto_paths"]:
        resultado["pagos_caf_base"] = construir_pagos_base(resultado["anio"])
        resultado["pagos_caf_dir"] = construir_pagos_mes(resultado["pagos_caf_base"], resultado["mes"])
        resultado["reporte_path"] = construir_reporte(resultado["anio"])
        resultado["sheet_name"] = construir_hoja(resultado["mes"], resultado["anio"])
    else:
        if not resultado.get("pagos_caf_base"):
            resultado["pagos_caf_base"] = construir_pagos_base(resultado["anio"])
        if not resultado.get("pagos_caf_dir"):
            resultado["pagos_caf_dir"] = construir_pagos_mes(resultado["pagos_caf_base"], resultado["mes"])
        if not resultado.get("reporte_path"):
            resultado["reporte_path"] = construir_reporte(resultado["anio"])
        if not resultado.get("sheet_name"):
            resultado["sheet_name"] = construir_hoja(resultado["mes"], resultado["anio"])

    resultado["pagos_caf_base"] = corregir_ruta_red(str(resultado.get("pagos_caf_base") or ""))
    resultado["pagos_caf_dir"] = corregir_ruta_red(str(resultado.get("pagos_caf_dir") or ""))
    resultado["reporte_path"] = corregir_ruta_red(str(resultado.get("reporte_path") or ""))
    resultado["sheet_name"] = str(resultado.get("sheet_name") or "").strip()
    resultado["meses"] = [{"nombre": nombre, "abreviado": abbr} for nombre, abbr in MESES]
    return resultado


def cargar_configuracion():
    if not CONFIG_PATH.exists():
        config = normalizar_configuracion(default_config())
        guardar_configuracion(config)
        return config

    try:
        with CONFIG_PATH.open("r", encoding="utf-8") as file:
            config = json.load(file)
    except (OSError, json.JSONDecodeError):
        return normalizar_configuracion(default_config())
    return normalizar_configuracion(config)


def guardar_configuracion(config):
    config = normalizar_configuracion(config)
    serializable = {key: value for key, value in config.items() if key != "meses"}
    CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with CONFIG_PATH.open("w", encoding="utf-8") as file:
        json.dump(serializable, file, indent=2, ensure_ascii=False)
    return config


def copy_cell_style(source, target):
    target._style = copy(source._style)
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def copy_column_style(ws, source_col, target_col):
    for row in range(1, ws.max_row + 1):
        copy_cell_style(ws.cell(row=row, column=source_col), ws.cell(row=row, column=target_col))


def buscar_encabezado(ws, texto_objetivo, max_rows=50):
    objetivo = normalizar_texto(texto_objetivo)
    for row in range(1, min(ws.max_row, max_rows) + 1):
        for col in range(1, ws.max_column + 1):
            if normalizar_texto(ws.cell(row=row, column=col).value) == objetivo:
                return row, col
    raise ValueError(f"No se encontro el encabezado: {texto_objetivo}")


def asegurar_columnas_condonados(ws):
    header_row, mora_om_col = buscar_encabezado(ws, HEADER_MORA_OM)

    if normalizar_texto(ws.cell(header_row, mora_om_col + 1).value) != HEADER_CONDONADOS_OM:
        ws.insert_cols(mora_om_col + 1)
        copy_column_style(ws, mora_om_col, mora_om_col + 1)
        ws.cell(header_row, mora_om_col + 1).value = "Condonados (Otras Monedas)"
        ws.column_dimensions[ws.cell(row=1, column=mora_om_col + 1).column_letter].width = 18

    header_row, mora_usd_col = buscar_encabezado(ws, HEADER_MORA_USD)
    if normalizar_texto(ws.cell(header_row, mora_usd_col + 1).value) != HEADER_CONDONADOS_USD:
        ws.insert_cols(mora_usd_col + 1)
        copy_column_style(ws, mora_usd_col, mora_usd_col + 1)
        ws.cell(header_row, mora_usd_col + 1).value = "Condonados (USD)"
        ws.column_dimensions[ws.cell(row=1, column=mora_usd_col + 1).column_letter].width = 16

    header_row, cond_om_col = buscar_encabezado(ws, HEADER_CONDONADOS_OM)
    _header_row, cond_usd_col = buscar_encabezado(ws, HEADER_CONDONADOS_USD)
    return header_row, cond_om_col, cond_usd_col


def leer_valores_reporte(config):
    reporte_path = Path(config["reporte_path"])
    if not reporte_path.exists():
        return {}

    try:
        wb = load_workbook(reporte_path, data_only=True, read_only=True)
    except Exception:
        return {}

    if config["sheet_name"] not in wb.sheetnames:
        return {}

    ws = wb[config["sheet_name"]]
    try:
        header_row, cond_om_col = buscar_encabezado(ws, HEADER_CONDONADOS_OM)
        _header_row, ref_col = buscar_encabezado(ws, HEADER_REFERENCIA)
    except ValueError:
        return {}

    valores = {}
    conteo_valores = {}
    for row in range(header_row + 1, ws.max_row + 1):
        fila_txt = " ".join(
            normalizar_texto(ws.cell(row=row, column=col).value)
            for col in range(1, min(ws.max_column, 12) + 1)
        )
        if "TOTAL" in fila_txt or "SUBTOTAL" in fila_txt:
            continue
        clave = extraer_clave_cfa(ws.cell(row=row, column=ref_col).value)
        if not clave:
            continue
        valor = abs_decimal(ws.cell(row=row, column=cond_om_col).value)
        if valor:
            valores[clave] = valor
            conteo_valores[valor] = conteo_valores.get(valor, 0) + 1

    # Si un mismo valor aparece demasiadas veces en el reporte, normalmente es
    # un total copiado por error. No se usa como respaldo para evitar repetirlo
    # en todas las carpetas durante el escaneo.
    sospechosos = {valor for valor, veces in conteo_valores.items() if veces >= 4}
    if sospechosos:
        valores = {clave: valor for clave, valor in valores.items() if valor not in sospechosos}
    return valores


def extraer_condonado_excel(carpeta):
    patrones = ("*.xlsx", "*.xlsm")
    for patron in patrones:
        for archivo in sorted(Path(carpeta).glob(patron)):
            try:
                wb = load_workbook(archivo, data_only=False)
            except Exception:
                continue

            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        formula = cell.value
                        if not isinstance(formula, str) or not formula.startswith("="):
                            continue

                        fila = cell.row
                        col = cell.column
                        etiquetas = []
                        for delta in (1, 2, -1, -2):
                            if col + delta > 0:
                                etiquetas.append(normalizar_texto(ws.cell(row=fila, column=col + delta).value))
                        etiquetas.append(normalizar_texto(ws.cell(row=fila + 1, column=col).value))
                        etiquetas.append(normalizar_texto(ws.cell(row=max(1, fila - 1), column=col).value))

                        if not any("MONTO RETENIDO POR INTERES" in etiqueta for etiqueta in etiquetas):
                            continue

                        restas = re.findall(r"-\s*\(?\s*([0-9]{1,3}(?:[.,][0-9]{3})+(?:[.,][0-9]+)?|[0-9]+(?:[.,][0-9]+)?)", formula)
                        if restas:
                            return abs_decimal(restas[0]), f"excel:{archivo.name}"
    return 0.0, "manual"


def buscar_estado_cuenta(carpeta):
    folder = Path(carpeta)
    candidatos = []
    for archivo in folder.iterdir():
        if not archivo.is_file() or archivo.suffix.lower() != ".pdf":
            continue
        nombre = normalizar_texto(archivo.stem)
        if "ESTADO" in nombre and "CUENTA" in nombre:
            candidatos.append(archivo)
    if candidatos:
        return sorted(candidatos, key=lambda item: len(item.name))[0]
    exacto = folder / "Estado Cuenta.pdf"
    return exacto if exacto.exists() else None


def extraer_texto_pdf(pdf_path, paginas=2):
    if PyPDF2 is None or not pdf_path or not Path(pdf_path).exists():
        return ""

    texto = ""
    try:
        with open(pdf_path, "rb") as file:
            lector = PyPDF2.PdfReader(file)
            for pagina in lector.pages[:paginas]:
                texto += "\n" + (pagina.extract_text() or "")
    except Exception:
        return ""
    return texto


def extraer_valor_financiamiento_compensatorio(texto):
    compacto = normalizar_texto(texto)
    if not compacto:
        return 0.0

    patrones = [
        r"FINANC\.?\s*COMPENSATORIO[:\s]+(?:[0-9.,]+%?\s*)?(\(?[0-9][0-9.,]*\)?)",
        r"FINANCIAMIENTO\s+COMPENSATORIO[:\s]+(?:[0-9.,]+%?\s*)?(\(?[0-9][0-9.,]*\)?)",
    ]
    for patron in patrones:
        match = re.search(patron, compacto)
        if match:
            return abs_decimal(match.group(1))
    return 0.0


def buscar_ejecutable(nombre, candidatos=()):
    encontrado = shutil.which(nombre)
    if encontrado:
        return encontrado
    for candidato in candidatos:
        if candidato and Path(candidato).exists():
            return str(candidato)
    return ""


def extraer_texto_ocr_pdf(pdf_path):
    tesseract = buscar_ejecutable(
        "tesseract",
        (
            r"C:\Program Files\PDF24\tesseract\tesseract.exe",
            r"C:\Users\bromo\Documents\Public External Debt\work\caf_condonados\tesseract\tesseract.exe",
        ),
    )
    pdftoppm = buscar_ejecutable("pdftoppm", ("pdftoppm.cmd",))
    if not tesseract or not pdftoppm:
        return ""

    try:
        with tempfile.TemporaryDirectory(prefix="caf_ocr_") as tmpdir:
            out_prefix = str(Path(tmpdir) / "page")
            subprocess.run(
                [pdftoppm, "-f", "1", "-l", "1", "-r", "220", "-png", str(pdf_path), out_prefix],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=45,
            )
            imagenes = sorted(Path(tmpdir).glob("page*.png"))
            texto = ""
            for imagen in imagenes[:1]:
                resultado = subprocess.run(
                    [tesseract, str(imagen), "stdout", "--psm", "6"],
                    check=True,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.DEVNULL,
                    text=True,
                    timeout=45,
                )
                texto += "\n" + resultado.stdout
            return texto
    except Exception:
        return ""


def extraer_condonado_pdf_detalle(estado_cuenta):
    if not estado_cuenta:
        return 0.0, "manual", ""

    texto = extraer_texto_pdf(estado_cuenta)
    valor = extraer_valor_financiamiento_compensatorio(texto)
    if valor:
        return valor, "pdf_texto", extraer_clave_cfa(texto)

    texto_ocr = extraer_texto_ocr_pdf(estado_cuenta)
    valor = extraer_valor_financiamiento_compensatorio(texto_ocr)
    if valor:
        return valor, "pdf_ocr", extraer_clave_cfa(texto_ocr)

    return 0.0, "manual", extraer_clave_cfa(texto or texto_ocr)


def extraer_condonado_pdf(estado_cuenta):
    valor, fuente, _claves = extraer_condonado_pdf_detalle(estado_cuenta)
    return valor, fuente


def obtener_condonado_carpeta(folder, valores_reporte, grupos_pdf):
    clave = extraer_clave_cfa(folder.name)

    estado_cuenta = buscar_estado_cuenta(folder)
    valor_pdf, fuente_pdf, claves_pdf = extraer_condonado_pdf_detalle(estado_cuenta)
    if valor_pdf:
        claves_pdf_lista = list(clave_set(claves_pdf))
        clave_preferida = claves_pdf.split("|")[0] if claves_pdf else ""
        grupo = claves_pdf or clave
        if grupo in grupos_pdf and grupos_pdf[grupo] != clave:
            return 0.0, f"duplicado_pdf:{grupos_pdf[grupo]}"
        if clave_preferida and clave and clave_preferida not in clave_set(clave):
            if any(claves_coinciden(clave_preferida, existente) for existente in valores_reporte):
                return 0.0, f"duplicado_pdf:{clave_preferida}"
            if clave_preferida in claves_pdf_lista:
                return 0.0, f"duplicado_pdf:{clave_preferida}"
        grupos_pdf[grupo] = clave
        return valor_pdf, fuente_pdf

    valor_excel, fuente_excel = extraer_condonado_excel(folder)
    if valor_excel:
        return valor_excel, fuente_excel

    # No se reutiliza el reporte BCE como fuente de escaneo: si el reporte ya
    # contiene un total mal copiado, ese valor se propagaria a todos los creditos.
    return 0.0, "manual"


def escanear_estado_cuenta(config):
    pagos_dir = Path(config["pagos_caf_dir"])
    if not pagos_dir.exists():
        raise FileNotFoundError(f"No se puede acceder a: {pagos_dir}")

    valores_reporte = leer_valores_reporte(config)
    items = []
    grupos_pdf = {}
    for folder in sorted(pagos_dir.iterdir(), key=lambda item: item.name):
        if not folder.is_dir():
            continue
        estado_cuenta = buscar_estado_cuenta(folder)
        clave = extraer_clave_cfa(folder.name)
        sigade = extraer_sigade(folder.name)
        valor, fuente = obtener_condonado_carpeta(folder, valores_reporte, grupos_pdf)

        items.append(
            {
                "id": f"r{len(items) + 1}",
                "clave": clave,
                "sigade": sigade,
                "carpeta": folder.name,
                "estado_cuenta": str(estado_cuenta) if estado_cuenta else "",
                "valor": valor,
                "fuente": fuente,
                "existe_estado_cuenta": bool(estado_cuenta),
            }
        )
    return items


def normalizar_items(raw_items):
    items = []
    for item in raw_items or []:
        clave = extraer_clave_cfa(item.get("clave") or item.get("carpeta") or item.get("referencia"))
        if not clave:
            continue
        items.append(
            {
                "clave": clave,
                "sigade": str(item.get("sigade") or ""),
                "carpeta": str(item.get("carpeta") or ""),
                "valor": parse_decimal(item.get("valor")),
            }
        )
    return items


def crear_backup(reporte_path):
    reporte = Path(reporte_path)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = reporte.with_name(f"{reporte.stem}.backup-condonados-{stamp}{reporte.suffix}")
    shutil.copy2(reporte, backup)
    return str(backup)


def es_fila_total_caf(ws, row):
    textos = [
        normalizar_texto(ws.cell(row=row, column=col).value)
        for col in range(1, min(ws.max_column, 10) + 1)
    ]
    unido = " ".join(texto for texto in textos if texto)
    es_total = "TOTAL" in unido or "SUBTOTAL" in unido
    es_caf = "CORPORACION ANDINA" in unido or "CAF" in unido
    return es_total and es_caf


def aplicar_condonados(config, raw_items):
    reporte_path = Path(config["reporte_path"])
    if not reporte_path.exists():
        raise FileNotFoundError(f"No se encontro el reporte: {reporte_path}")

    items = normalizar_items(raw_items)
    valores = [(item["clave"], item["valor"]) for item in items]

    wb = load_workbook(reporte_path)
    if config["sheet_name"] not in wb.sheetnames:
        raise ValueError(f"No existe la hoja: {config['sheet_name']}")

    ws = wb[config["sheet_name"]]
    header_row, cond_om_col, cond_usd_col = asegurar_columnas_condonados(ws)
    _header_row, ref_col = buscar_encabezado(ws, HEADER_REFERENCIA)

    matched = []
    missing = {clave for clave, _valor in valores}
    total = 0.0

    for row in range(header_row + 1, ws.max_row + 1):
        for col in (cond_om_col, cond_usd_col):
            cell = ws.cell(row=row, column=col)
            cell.value = 0.0
            cell.number_format = NUM_FORMAT

        clave_reporte = extraer_clave_cfa(ws.cell(row=row, column=ref_col).value)
        if not clave_reporte:
            continue

        elegido = None
        for clave_item, valor in valores:
            if claves_coinciden(clave_reporte, clave_item):
                elegido = (clave_item, valor)
                break

        if elegido:
            clave_item, valor = elegido
            ws.cell(row=row, column=cond_om_col).value = valor
            ws.cell(row=row, column=cond_usd_col).value = valor
            matched.append(
                {
                    "fila": row,
                    "clave": clave_item,
                    "valor": valor,
                    "referencia": ws.cell(row=row, column=ref_col).value,
                }
            )
            missing.discard(clave_item)
            total = round(total + valor, 2)

    for row in range(header_row + 1, ws.max_row + 1):
        if es_fila_total_caf(ws, row):
            ws.cell(row=row, column=cond_om_col).value = total
            ws.cell(row=row, column=cond_usd_col).value = total

    backup_path = crear_backup(reporte_path) if config.get("crear_backup", True) else ""
    wb.save(reporte_path)

    return {
        "reporte_path": str(reporte_path),
        "backup_path": backup_path,
        "total": total,
        "matched": matched,
        "missing": sorted(missing),
        "columnas": {
            "condonados_otras_monedas": ws.cell(row=header_row, column=cond_om_col).coordinate,
            "condonados_usd": ws.cell(row=header_row, column=cond_usd_col).coordinate,
        },
    }


def register_caf_condonados_routes(app, base_dir=None):
    """Registra endpoints JSON del modulo en una app Flask existente."""
    global CONFIG_PATH
    if base_dir:
        CONFIG_PATH = Path(base_dir) / "caf_condonados_config.json"

    @app.get("/api/caf_condonados/config")
    def caf_condonados_config():
        return jsonify({"ok": True, "config": cargar_configuracion()})

    @app.post("/api/caf_condonados/config")
    def caf_condonados_guardar_config():
        data = request.get_json(force=True)
        config = cargar_configuracion()
        for key in (
            "anio",
            "mes",
            "pagos_caf_base",
            "pagos_caf_dir",
            "reporte_path",
            "sheet_name",
            "auto_paths",
            "crear_backup",
        ):
            if key in data:
                config[key] = data[key]
        return jsonify({"ok": True, "config": guardar_configuracion(config)})

    @app.post("/api/caf_condonados/scan")
    def caf_condonados_scan():
        data = request.get_json(silent=True) or {}
        config = cargar_configuracion()
        for key in (
            "anio",
            "mes",
            "pagos_caf_base",
            "pagos_caf_dir",
            "reporte_path",
            "sheet_name",
            "auto_paths",
            "crear_backup",
        ):
            if key in data:
                config[key] = data[key]
        config = normalizar_configuracion(config)
        try:
            items = escanear_estado_cuenta(config)
        except FileNotFoundError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 404
        except Exception as exc:
            return jsonify({"ok": False, "error": f"Error al escanear: {exc}"}), 500
        return jsonify({"ok": True, "items": items, "total": len(items)})

    @app.post("/api/caf_condonados/process")
    def caf_condonados_process():
        data = request.get_json(force=True)
        config = cargar_configuracion()
        for key in (
            "anio",
            "mes",
            "pagos_caf_base",
            "pagos_caf_dir",
            "reporte_path",
            "sheet_name",
            "auto_paths",
            "crear_backup",
        ):
            if key in data:
                config[key] = data[key]
        config = normalizar_configuracion(config)

        try:
            resultado = aplicar_condonados(config, data.get("items") or [])
        except FileNotFoundError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 404
        except ValueError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400
        except Exception as exc:
            return jsonify({"ok": False, "error": f"Error al actualizar reporte: {exc}"}), 500
        return jsonify({"ok": True, **resultado})
